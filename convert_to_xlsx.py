from numbers_parser import Document
import openpyxl

doc = Document("../ChartOfAccount.numbers")
wb = openpyxl.Workbook()
ws = wb.active
ws.title = "COA Latest"

sheet = doc.sheets[0]
table = sheet.tables[0]

for row in table.rows():
    ws.append([cell.value for cell in row])

wb.save("../ChartOfAccount.xlsx")
print("Saved ChartOfAccount.xlsx")
