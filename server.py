import asyncio
from dotenv import load_dotenv
from pathlib import Path
ROOT_DIR = Path(__file__).parent
load_dotenv(ROOT_DIR / '.env')

import os
import re
import uuid
import string
import random
import logging
import bcrypt
import jwt
import aiomysql
import hmac
import hashlib
import requests
try:
    import razorpay
except ImportError:
    razorpay = None
from datetime import datetime, timezone, timedelta, date
from typing import Optional, List
from fastapi import FastAPI, APIRouter, HTTPException, Request, Response, Depends, UploadFile, File
from fastapi.openapi.utils import get_openapi
from starlette.middleware.cors import CORSMiddleware
from pydantic import BaseModel, EmailStr, Field, field_validator
from cryptography.fernet import Fernet
from fastapi import Form, File, UploadFile


ENCRYPTION_KEY = os.getenv("ENCRYPTION_KEY")
cipher_suite = Fernet(ENCRYPTION_KEY.encode()) if ENCRYPTION_KEY else None

def encrypt_data(data: str) -> Optional[str]:
    if not data or not cipher_suite: return data
    return cipher_suite.encrypt(data.encode()).decode()

def decrypt_data(data: str) -> Optional[str]:
    if not data or not cipher_suite: return data
    try:
        return cipher_suite.decrypt(data.encode()).decode()
    except Exception:
        return data  # Return as is if decryption fails (e.g. legacy plain text)

async def log_audit(entity_name: str, entity_id: str, action: str, changes: dict, performed_by: str = None):
    import json
    changes_json = json.dumps(changes) if changes else None
    
    await db_execute(
        "INSERT INTO audit_logs (entity_name, entity_id, action, changes, performed_by) VALUES (%s, %s, %s, %s, %s)",
        (entity_name, entity_id, action, changes_json, performed_by)
    )

class SubscriptionPlanCreate(BaseModel):
    name: str
    code: str
    country: str
    description: Optional[str] = None
    monthly_price: Optional[float] = None
    yearly_price: Optional[float] = None
    trial_days: Optional[int] = 0
    is_active: Optional[bool] = True

class SubscriptionPlanUpdate(BaseModel):
    name: str
    code: str
    country: str
    description: Optional[str] = None
    monthly_price: Optional[float] = None
    yearly_price: Optional[float] = None
    trial_days: Optional[int] = 0
    is_active: Optional[bool] = True

class SubscriptionPlanFeatures(BaseModel):
    feature_ids: List[str]

class AssociationSubscriptionUpdate(BaseModel):
    plan_id: Optional[str] = None
    subscription_start: Optional[str] = None
    subscription_end: Optional[str] = None
    payment_status: Optional[str] = None
    renewal_date: Optional[str] = None
    subscription_status: Optional[str] = None


# ---------- Logging ----------
logging.basicConfig(level=logging.INFO, format='%(asctime)s - %(name)s - %(levelname)s - %(message)s')
logger = logging.getLogger("nestora")


# ---------- Config ----------
JWT_ALGORITHM = "HS256"
JWT_EXP_MIN = 60 * 24 * 7  # 7 days for MVP simplicity


def get_jwt_secret() -> str:
    return os.environ["JWT_SECRET"]


# ---------- MySQL Pool ----------
pool: Optional[aiomysql.Pool] = None


async def init_pool():
    global pool
    pool = await aiomysql.create_pool(
        host=os.environ["MYSQL_HOST"],
        port=int(os.environ["MYSQL_PORT"]),
        user=os.environ["MYSQL_USER"],
        password=os.environ["MYSQL_PASSWORD"],
        db=os.environ["MYSQL_DB"],
        autocommit=True,
        minsize=1,
        maxsize=10,
        charset="utf8mb4",
    )


async def db_execute(sql: str, params: tuple = ()):
    import uuid
    import re
    is_insert = sql.strip().upper().startswith("INSERT INTO")
    new_id = None
    if is_insert:
        new_id = str(uuid.uuid4())
        
        pk_map = {
            'accounts': 'account_id',
            'user_details': 'user_id',
            'employees': 'employee_id',
            'admin_associations': None,
            'event_likes': None,
            'poll_likes': None,
            'poll_votes': None,
            'password_reset_otps': None
        }
        
        match = re.search(r'INSERT INTO\s+([a-zA-Z0-9_]+)\s*\(([^)]+)\)\s*VALUES\s*\(([^)]+)\)', sql, re.IGNORECASE)
        if match:
            table = match.group(1).lower()
            cols = match.group(2)
            vals = match.group(3)
            pk = pk_map.get(table, 'id')
            if pk:
                sql = sql.replace(f"({cols})", f"({pk}, {cols})", 1)
                sql = sql.replace(f"({vals})", f"(%s, {vals})", 1)
                params = (new_id,) + tuple(params)
            else:
                new_id = None
        else:
            match_sel = re.search(r'INSERT INTO\s+([a-zA-Z0-9_]+)\s*\(([^)]+)\)\s*SELECT', sql, re.IGNORECASE)
            if match_sel:
                table = match_sel.group(1).lower()
                cols = match_sel.group(2)
                pk = pk_map.get(table, 'id')
                if pk:
                    sql = sql.replace(f"({cols})", f"({pk}, {cols})", 1)
                    sql = sql.replace("SELECT ", "SELECT %s, ", 1)
                    params = (new_id,) + tuple(params)
                else:
                    new_id = None

    async with pool.acquire() as conn:
        async with conn.cursor(aiomysql.DictCursor) as cur:
            await cur.execute(sql, params)
            if is_insert and new_id:
                return new_id
            return cur.lastrowid


async def db_fetchone(sql: str, params: tuple = ()) -> Optional[dict]:
    async with pool.acquire() as conn:
        async with conn.cursor(aiomysql.DictCursor) as cur:
            await cur.execute(sql, params)
            return await cur.fetchone()


async def db_fetchall(sql: str, params: tuple = ()) -> List[dict]:
    async with pool.acquire() as conn:
        async with conn.cursor(aiomysql.DictCursor) as cur:
            await cur.execute(sql, params)
            return list(await cur.fetchall())


# ---------- Schema Setup ----------
SCHEMA_SQL = [
    """
    CREATE TABLE IF NOT EXISTS password_reset_otps (
        email VARCHAR(255) PRIMARY KEY,
        otp VARCHAR(10),
        expires_at DATETIME
    ) ENGINE=InnoDB DEFAULT CHARSET=utf8mb4;
    """,
    """
    CREATE TABLE IF NOT EXISTS board_committee_chat (
        id CHAR(36) PRIMARY KEY,
        association_id VARCHAR(50),
        pool_type VARCHAR(50),
        pool_id VARCHAR(50),
        sender_id CHAR(36),
        message TEXT,
        created_at DATETIME DEFAULT CURRENT_TIMESTAMP,
        FOREIGN KEY (sender_id) REFERENCES accounts(account_id) ON DELETE CASCADE
    ) ENGINE=InnoDB DEFAULT CHARSET=utf8mb4;
    """,


    """
    CREATE TABLE IF NOT EXISTS family_members (
        id CHAR(36) PRIMARY KEY,
        user_id CHAR(36) NOT NULL,
        name VARCHAR(150) NOT NULL,
        email VARCHAR(150),
        contact_number VARCHAR(20),
        alt_contact_number VARCHAR(20),
        created_at DATETIME DEFAULT CURRENT_TIMESTAMP,
        FOREIGN KEY (user_id) REFERENCES user_details(user_id) ON DELETE CASCADE
    ) ENGINE=InnoDB DEFAULT CHARSET=utf8mb4;
    """,
    """
    CREATE TABLE IF NOT EXISTS pets (
        id CHAR(36) PRIMARY KEY,
        user_id CHAR(36) NOT NULL,
        type VARCHAR(50),
        name VARCHAR(150) NOT NULL,
        breed VARCHAR(100),
        vaccinated BOOLEAN DEFAULT FALSE,
        vaccination_date DATE,
        next_vaccination_reminder BOOLEAN DEFAULT FALSE,
        vaccination_certificate_url VARCHAR(255),
        reminder_date DATE,
        created_at DATETIME DEFAULT CURRENT_TIMESTAMP,
        FOREIGN KEY (user_id) REFERENCES user_details(user_id) ON DELETE CASCADE
    ) ENGINE=InnoDB DEFAULT CHARSET=utf8mb4;
    """,


    """
    CREATE TABLE IF NOT EXISTS vehicles (
        id CHAR(36) PRIMARY KEY,
        user_id CHAR(36) NOT NULL,
        type VARCHAR(50),
        registration_number VARCHAR(50) NOT NULL,
        insurance_url VARCHAR(255),
        puc_url VARCHAR(255),
        created_at DATETIME DEFAULT CURRENT_TIMESTAMP,
        FOREIGN KEY (user_id) REFERENCES user_details(user_id) ON DELETE CASCADE
    ) ENGINE=InnoDB DEFAULT CHARSET=utf8mb4;
    """,


    """
    CREATE TABLE IF NOT EXISTS service_staff (
        id CHAR(36) PRIMARY KEY,
        staff_code VARCHAR(50) UNIQUE NOT NULL,
        name VARCHAR(150) NOT NULL,
        mobile VARCHAR(30) NOT NULL,
        photo_url VARCHAR(255),
        category VARCHAR(100),
        address VARCHAR(255),
        emergency_contact VARCHAR(30),
        police_verified BOOLEAN DEFAULT FALSE,
        id_verified BOOLEAN DEFAULT FALSE,
        status ENUM('Active', 'Suspended', 'Blacklisted') DEFAULT 'Active',
        pass_expiry_date DATE,
        created_at DATETIME DEFAULT CURRENT_TIMESTAMP
    ) ENGINE=InnoDB DEFAULT CHARSET=utf8mb4;
    """,
    """
    CREATE TABLE IF NOT EXISTS service_staff_assignments (
        id CHAR(36) PRIMARY KEY,
        staff_id CHAR(36) NOT NULL,
        unit_id CHAR(36) NOT NULL,
        resident_id CHAR(36),
        start_date DATE,
        end_date DATE,
        status ENUM('Active', 'Inactive') DEFAULT 'Active',
        created_at DATETIME DEFAULT CURRENT_TIMESTAMP,
        FOREIGN KEY (staff_id) REFERENCES service_staff(id) ON DELETE CASCADE,
        FOREIGN KEY (unit_id) REFERENCES units(id) ON DELETE CASCADE
    ) ENGINE=InnoDB DEFAULT CHARSET=utf8mb4;
    """,
    """
    CREATE TABLE IF NOT EXISTS service_staff_logs (
        id CHAR(36) PRIMARY KEY,
        staff_id CHAR(36) NOT NULL,
        check_in DATETIME DEFAULT CURRENT_TIMESTAMP,
        check_out DATETIME,
        gate VARCHAR(50),
        guard_id CHAR(36),
        remarks TEXT,
        FOREIGN KEY (staff_id) REFERENCES service_staff(id) ON DELETE CASCADE,
        FOREIGN KEY (guard_id) REFERENCES accounts(account_id) ON DELETE SET NULL
    ) ENGINE=InnoDB DEFAULT CHARSET=utf8mb4;
    """,


    """
    CREATE TABLE IF NOT EXISTS deliveries (
        id CHAR(36) PRIMARY KEY,
        unit_id CHAR(36) NOT NULL,
        resident_id CHAR(36),
        delivery_type ENUM('Courier', 'Food', 'Parcel', 'Other') NOT NULL,
        company_name VARCHAR(150),
        delivery_person_name VARCHAR(150),
        mobile VARCHAR(30),
        status ENUM('Inside Premises', 'Collected at Gate', 'Completed', 'Cancelled') NOT NULL,
        gate VARCHAR(50),
        guard_id CHAR(36),
        package_photo_url VARCHAR(255),
        check_in DATETIME DEFAULT CURRENT_TIMESTAMP,
        check_out DATETIME,
        created_at DATETIME DEFAULT CURRENT_TIMESTAMP,
        FOREIGN KEY (unit_id) REFERENCES units(id) ON DELETE CASCADE,
        FOREIGN KEY (guard_id) REFERENCES accounts(account_id) ON DELETE SET NULL
    ) ENGINE=InnoDB DEFAULT CHARSET=utf8mb4;
    """,


    """
    CREATE TABLE IF NOT EXISTS pre_approved_visitors (
        id CHAR(36) PRIMARY KEY,
        resident_id CHAR(36) NOT NULL,
        unit_id CHAR(36) NOT NULL,
        visitor_name VARCHAR(150) NOT NULL,
        mobile VARCHAR(20) NOT NULL,
        visitor_type VARCHAR(50) NOT NULL,
        pass_code VARCHAR(50) UNIQUE NOT NULL,
        otp VARCHAR(10),
        visit_date DATE NOT NULL,
        start_time TIME NOT NULL,
        end_time TIME NOT NULL,
        number_of_visitors INT DEFAULT 1,
        vehicle_number VARCHAR(50),
        purpose VARCHAR(255),
        pass_type VARCHAR(50) DEFAULT 'Single Entry',
        status ENUM('Active', 'Used', 'Expired', 'Cancelled') DEFAULT 'Active',
        created_by CHAR(36),
        created_at DATETIME DEFAULT CURRENT_TIMESTAMP,
        FOREIGN KEY (resident_id) REFERENCES user_details(user_id) ON DELETE CASCADE,
        FOREIGN KEY (unit_id) REFERENCES units(id) ON DELETE CASCADE
    ) ENGINE=InnoDB DEFAULT CHARSET=utf8mb4;
    """,
    """
    CREATE TABLE IF NOT EXISTS visitor_logs (
        id CHAR(36) PRIMARY KEY,
        pre_approved_id CHAR(36) NOT NULL,
        check_in DATETIME DEFAULT CURRENT_TIMESTAMP,
        check_out DATETIME,
        gate VARCHAR(50),
        guard_id CHAR(36),
        visitor_photo_url VARCHAR(255),
        remarks TEXT,
        FOREIGN KEY (pre_approved_id) REFERENCES pre_approved_visitors(id) ON DELETE CASCADE,
        FOREIGN KEY (guard_id) REFERENCES accounts(account_id) ON DELETE SET NULL
    ) ENGINE=InnoDB DEFAULT CHARSET=utf8mb4;
    """,

    """
    CREATE TABLE IF NOT EXISTS user_codes (
      id CHAR(36) PRIMARY KEY,
      login_code VARCHAR(20) UNIQUE NOT NULL,
      status ENUM('active','used','expired','revoked') DEFAULT 'active',
      expires_at DATETIME NULL,
      created_at DATETIME DEFAULT CURRENT_TIMESTAMP
    ) ENGINE=InnoDB DEFAULT CHARSET=utf8mb4;
    """,
    """
    CREATE TABLE IF NOT EXISTS user_details (
      user_id CHAR(36) PRIMARY KEY,
      code_id CHAR(36) UNIQUE,
      name VARCHAR(150) NOT NULL,
      address VARCHAR(255) NOT NULL,
      email VARCHAR(150) NOT NULL,
      contact_number VARCHAR(30) NOT NULL,
      first_name VARCHAR(100),
      last_name VARCHAR(100),
      address_line_1 VARCHAR(255),
      address_line_2 VARCHAR(255),
      city VARCHAR(100),
      state VARCHAR(100),
      pincode VARCHAR(20),
      emergency_contact_name VARCHAR(150),
      emergency_contact_number VARCHAR(30),
      id_proof_url VARCHAR(255),
      profile_pic_url VARCHAR(255),
      onboard_date DATE,
      end_date DATE,
      board_member_since DATE,
      created_at DATETIME DEFAULT CURRENT_TIMESTAMP,
      FOREIGN KEY (code_id) REFERENCES user_codes(id) ON DELETE SET NULL
    ) ENGINE=InnoDB DEFAULT CHARSET=utf8mb4;
    """,
    """
    CREATE TABLE IF NOT EXISTS employees (
      employee_id CHAR(36) PRIMARY KEY,
      employee_id_number VARCHAR(50) UNIQUE NOT NULL,
      name VARCHAR(150) NOT NULL,
      address VARCHAR(255) NOT NULL,
      email VARCHAR(150) NOT NULL,
      contact_number VARCHAR(30) NOT NULL,
      first_name VARCHAR(100),
      last_name VARCHAR(100),
      address_line_1 VARCHAR(255),
      address_line_2 VARCHAR(255),
      city VARCHAR(100),
      state VARCHAR(100),
      pincode VARCHAR(20),
      emergency_contact_name VARCHAR(150),
      emergency_contact_number VARCHAR(30),
      id_proof_url VARCHAR(255),
      profile_pic_url VARCHAR(255),
      onboard_date DATE,
      end_date DATE,
      created_at DATETIME DEFAULT CURRENT_TIMESTAMP
    ) ENGINE=InnoDB DEFAULT CHARSET=utf8mb4;
    """,
    """
    CREATE TABLE IF NOT EXISTS entity_types (
      id CHAR(36) PRIMARY KEY,
      name VARCHAR(150) NOT NULL,
      description TEXT,
      created_at DATETIME DEFAULT CURRENT_TIMESTAMP
    ) ENGINE=InnoDB DEFAULT CHARSET=utf8mb4;
    """,
    """
    CREATE TABLE IF NOT EXISTS board_members (
      id CHAR(36) PRIMARY KEY,
      association_id CHAR(36) NOT NULL,
      account_id CHAR(36) NOT NULL,
      term_start_date DATE NOT NULL,
      term_end_date DATE NOT NULL,
      status ENUM('active', 'past') DEFAULT 'active',
      created_by CHAR(36) NULL,
      created_at DATETIME DEFAULT CURRENT_TIMESTAMP,
      FOREIGN KEY (association_id) REFERENCES associations(id) ON DELETE CASCADE,
      FOREIGN KEY (account_id) REFERENCES accounts(account_id) ON DELETE CASCADE,
      FOREIGN KEY (created_by) REFERENCES accounts(account_id) ON DELETE SET NULL
    ) ENGINE=InnoDB DEFAULT CHARSET=utf8mb4;
    """,
    """
    CREATE TABLE IF NOT EXISTS entities (
      id CHAR(36) PRIMARY KEY,
      entity_type_id CHAR(36),
      association_id CHAR(36) NULL,
      name VARCHAR(150) NOT NULL,
      description TEXT,
      created_at DATETIME DEFAULT CURRENT_TIMESTAMP,
      FOREIGN KEY (entity_type_id) REFERENCES entity_types(id) ON DELETE SET NULL
    ) ENGINE=InnoDB DEFAULT CHARSET=utf8mb4;
    """,
    """
    CREATE TABLE IF NOT EXISTS roles (
      id CHAR(36) PRIMARY KEY,
      entity_id CHAR(36) NULL,
      name VARCHAR(150) NOT NULL,
      code VARCHAR(100) UNIQUE NULL,
      description TEXT,
      is_active BOOLEAN DEFAULT TRUE,
      created_at DATETIME DEFAULT CURRENT_TIMESTAMP,
      FOREIGN KEY (entity_id) REFERENCES entities(id) ON DELETE SET NULL
    ) ENGINE=InnoDB DEFAULT CHARSET=utf8mb4;
    """,
    """
    CREATE TABLE IF NOT EXISTS accounts (
      account_id CHAR(36) PRIMARY KEY,
      user_id CHAR(36) NULL,
      employee_id CHAR(36) NULL,
      email VARCHAR(150) UNIQUE NOT NULL,
      password_hash VARCHAR(255) NOT NULL,
      raw_password VARCHAR(255),
      role_id CHAR(36) NOT NULL,
      status ENUM('active', 'inactive', 'suspended') DEFAULT 'active',
      created_at DATETIME DEFAULT CURRENT_TIMESTAMP,
      FOREIGN KEY (user_id) REFERENCES user_details(user_id) ON DELETE CASCADE,
      FOREIGN KEY (employee_id) REFERENCES employees(employee_id) ON DELETE CASCADE,
      FOREIGN KEY (role_id) REFERENCES roles(id) ON DELETE SET NULL
    ) ENGINE=InnoDB DEFAULT CHARSET=utf8mb4;
    """,
    """
    CREATE TABLE IF NOT EXISTS features (
      id CHAR(36) PRIMARY KEY,
      name VARCHAR(150) NOT NULL,
      code VARCHAR(100) UNIQUE NULL,
      description TEXT,
      parent_id CHAR(36) NULL,
      icon LONGTEXT,
      url VARCHAR(255),
      order_index INT DEFAULT 0,
      is_active BOOLEAN DEFAULT TRUE,
      created_at DATETIME DEFAULT CURRENT_TIMESTAMP,
      FOREIGN KEY (parent_id) REFERENCES features(id) ON DELETE SET NULL
    ) ENGINE=InnoDB DEFAULT CHARSET=utf8mb4;
    """,
    """
    CREATE TABLE IF NOT EXISTS role_feature_permissions (
      id CHAR(36) PRIMARY KEY,
      role_id CHAR(36) NOT NULL,
      feature_id CHAR(36) NOT NULL,
      can_create TINYINT(1) DEFAULT 0,
      can_view TINYINT(1) DEFAULT 0,
      can_update TINYINT(1) DEFAULT 0,
      can_delete TINYINT(1) DEFAULT 0,
      sidebar_order INT NOT NULL DEFAULT 0,
      created_at DATETIME DEFAULT CURRENT_TIMESTAMP,
      FOREIGN KEY (role_id) REFERENCES roles(id) ON DELETE CASCADE,
      FOREIGN KEY (feature_id) REFERENCES features(id) ON DELETE CASCADE,
      UNIQUE KEY uniq_role_feature (role_id, feature_id)
    ) ENGINE=InnoDB DEFAULT CHARSET=utf8mb4;
    """,
    """
    CREATE TABLE IF NOT EXISTS associations (
      id CHAR(36) PRIMARY KEY,
      name VARCHAR(255) NOT NULL,
      is_active TINYINT(1) DEFAULT 1,
      country VARCHAR(100) NULL,
      created_at DATETIME DEFAULT CURRENT_TIMESTAMP
    ) ENGINE=InnoDB DEFAULT CHARSET=utf8mb4;
    """,
    """
    CREATE TABLE IF NOT EXISTS admin_associations (
      admin_id CHAR(36) NOT NULL,
      association_id CHAR(36) NOT NULL,
      created_at DATETIME DEFAULT CURRENT_TIMESTAMP,
      PRIMARY KEY (admin_id, association_id),
      FOREIGN KEY (admin_id) REFERENCES accounts(account_id) ON DELETE CASCADE,
      FOREIGN KEY (association_id) REFERENCES associations(id) ON DELETE CASCADE
    ) ENGINE=InnoDB DEFAULT CHARSET=utf8mb4;
    """,
    """
    CREATE TABLE IF NOT EXISTS code_requests (
      id CHAR(36) PRIMARY KEY,
      name VARCHAR(150) NOT NULL,
      email VARCHAR(150) NOT NULL,
      contact_number VARCHAR(30) NOT NULL,
      status ENUM('pending','approved','rejected') DEFAULT 'pending',
      issued_code VARCHAR(20) NULL,
      created_at DATETIME DEFAULT CURRENT_TIMESTAMP
    ) ENGINE=InnoDB DEFAULT CHARSET=utf8mb4;
    """,
    """
    CREATE TABLE IF NOT EXISTS update_requests (
      id CHAR(36) PRIMARY KEY,
      code_id CHAR(36),
      requested_name VARCHAR(150),
      requested_address VARCHAR(255),
      requested_email VARCHAR(150),
      requested_contact VARCHAR(30),
      note TEXT,
      status ENUM('open','resolved') DEFAULT 'open',
      created_at DATETIME DEFAULT CURRENT_TIMESTAMP,
      FOREIGN KEY (code_id) REFERENCES user_codes(id) ON DELETE SET NULL
    ) ENGINE=InnoDB DEFAULT CHARSET=utf8mb4;
    """,
    """
    CREATE TABLE IF NOT EXISTS email_log (
      id CHAR(36) PRIMARY KEY,
      to_email VARCHAR(150) NOT NULL,
      subject VARCHAR(255) NOT NULL,
      body TEXT,
      created_at DATETIME DEFAULT CURRENT_TIMESTAMP
    ) ENGINE=InnoDB DEFAULT CHARSET=utf8mb4;
    """,
    """
    CREATE TABLE IF NOT EXISTS notifications (
      id CHAR(36) PRIMARY KEY,
      account_id CHAR(36) NOT NULL,
      title VARCHAR(255) NOT NULL,
      message TEXT,
      is_read TINYINT(1) DEFAULT 0,
      created_at DATETIME DEFAULT CURRENT_TIMESTAMP,
      FOREIGN KEY (account_id) REFERENCES accounts(account_id) ON DELETE CASCADE
    ) ENGINE=InnoDB DEFAULT CHARSET=utf8mb4;
    """,
    """
    CREATE TABLE IF NOT EXISTS announcements (
      id CHAR(36) PRIMARY KEY,
      association_id CHAR(36) NOT NULL,
      title VARCHAR(200) NOT NULL,
      body TEXT NOT NULL,
      category ENUM('general','maintenance','urgent','celebration') DEFAULT 'general',
      pinned TINYINT(1) DEFAULT 0,
      audience VARCHAR(100) DEFAULT 'Homeowners',
      attachment_url VARCHAR(255),
      created_by CHAR(36) NULL,
      created_at DATETIME DEFAULT CURRENT_TIMESTAMP,
      updated_at DATETIME DEFAULT CURRENT_TIMESTAMP ON UPDATE CURRENT_TIMESTAMP,
      FOREIGN KEY (association_id) REFERENCES associations(id) ON DELETE CASCADE
    ) ENGINE=InnoDB DEFAULT CHARSET=utf8mb4;
    """,
    """
    CREATE TABLE IF NOT EXISTS documents (
      id CHAR(36) PRIMARY KEY,
      association_id CHAR(36) NOT NULL,
      title VARCHAR(100) NOT NULL,
      document_number VARCHAR(100),
      document_type VARCHAR(50),
      category VARCHAR(50),
      file_name VARCHAR(255),
      file_url VARCHAR(500),
      file_type VARCHAR(20),
      file_size_kb INT,
      department VARCHAR(100),
      related_module VARCHAR(50),
      visibility VARCHAR(255),
      allow_download BOOLEAN DEFAULT TRUE,
      issue_date DATE,
      expiry_date DATE,
      reminder_before_expiry_days INT,
      status VARCHAR(20) DEFAULT 'Active',
      keywords VARCHAR(255),
      remarks VARCHAR(1000),
      created_at DATETIME DEFAULT CURRENT_TIMESTAMP,
      FOREIGN KEY (association_id) REFERENCES associations(id) ON DELETE CASCADE
    ) ENGINE=InnoDB DEFAULT CHARSET=utf8mb4;
    """,
    """
    CREATE TABLE IF NOT EXISTS announcement_likes (
      id CHAR(36) PRIMARY KEY,
      announcement_id CHAR(36) NOT NULL,
      account_id CHAR(36) NOT NULL,
      created_at DATETIME DEFAULT CURRENT_TIMESTAMP,
      UNIQUE KEY uniq_ann_like (announcement_id, account_id),
      FOREIGN KEY (announcement_id) REFERENCES announcements(id) ON DELETE CASCADE,
      FOREIGN KEY (account_id) REFERENCES accounts(account_id) ON DELETE CASCADE
    ) ENGINE=InnoDB DEFAULT CHARSET=utf8mb4;
    """,
    """
    CREATE TABLE IF NOT EXISTS announcement_comments (
      id CHAR(36) PRIMARY KEY,
      announcement_id CHAR(36) NOT NULL,
      account_id CHAR(36) NOT NULL,
      comment TEXT NOT NULL,
      created_at DATETIME DEFAULT CURRENT_TIMESTAMP,
      FOREIGN KEY (announcement_id) REFERENCES announcements(id) ON DELETE CASCADE,
      FOREIGN KEY (account_id) REFERENCES accounts(account_id) ON DELETE CASCADE
    ) ENGINE=InnoDB DEFAULT CHARSET=utf8mb4;
    """,
    """
    CREATE TABLE IF NOT EXISTS events (
      id CHAR(36) PRIMARY KEY,
      association_id CHAR(36) NULL,
      title VARCHAR(200) NOT NULL,
      description TEXT,
      category VARCHAR(50),
      banner_url VARCHAR(255),
      location VARCHAR(255),
      starts_at DATETIME NOT NULL,
      ends_at DATETIME NULL,
      is_registration_required TINYINT(1) DEFAULT 0,
      registration_deadline DATETIME,
      max_capacity INT,
      audience VARCHAR(50) DEFAULT 'All',
      is_paid TINYINT(1) DEFAULT 0,
      fee_amount DECIMAL(10,2),
      organizer_name VARCHAR(100),
      organizer_contact VARCHAR(50),
      status VARCHAR(50) DEFAULT 'Published',
      created_by CHAR(36) NULL,
      created_at DATETIME DEFAULT CURRENT_TIMESTAMP,
      updated_at DATETIME DEFAULT CURRENT_TIMESTAMP ON UPDATE CURRENT_TIMESTAMP,
      FOREIGN KEY (association_id) REFERENCES associations(id) ON DELETE CASCADE
    ) ENGINE=InnoDB DEFAULT CHARSET=utf8mb4;
    """,
    """
    CREATE TABLE IF NOT EXISTS event_rsvps (
      id CHAR(36) PRIMARY KEY,
      event_id CHAR(36) NOT NULL,
      account_id CHAR(36) NOT NULL,
      status ENUM('going','maybe','not_going') NOT NULL DEFAULT 'going',
      created_at DATETIME DEFAULT CURRENT_TIMESTAMP,
      updated_at DATETIME DEFAULT CURRENT_TIMESTAMP ON UPDATE CURRENT_TIMESTAMP,
      UNIQUE KEY uniq_event_account (event_id, account_id),
      FOREIGN KEY (event_id) REFERENCES events(id) ON DELETE CASCADE,
      FOREIGN KEY (account_id) REFERENCES accounts(account_id) ON DELETE CASCADE
    ) ENGINE=InnoDB DEFAULT CHARSET=utf8mb4;
    """,
    """
    CREATE TABLE IF NOT EXISTS event_likes (
      event_id CHAR(36) NOT NULL,
      account_id CHAR(36) NOT NULL,
      created_at DATETIME DEFAULT CURRENT_TIMESTAMP,
      PRIMARY KEY (event_id, account_id),
      FOREIGN KEY (event_id) REFERENCES events(id) ON DELETE CASCADE,
      FOREIGN KEY (account_id) REFERENCES accounts(account_id) ON DELETE CASCADE
    ) ENGINE=InnoDB DEFAULT CHARSET=utf8mb4;
    """,
    """
    CREATE TABLE IF NOT EXISTS event_comments (
      id CHAR(36) PRIMARY KEY,
      event_id CHAR(36) NOT NULL,
      account_id CHAR(36) NOT NULL,
      body TEXT NOT NULL,
      created_at DATETIME DEFAULT CURRENT_TIMESTAMP,
      FOREIGN KEY (event_id) REFERENCES events(id) ON DELETE CASCADE,
      FOREIGN KEY (account_id) REFERENCES accounts(account_id) ON DELETE CASCADE
    ) ENGINE=InnoDB DEFAULT CHARSET=utf8mb4;
    """,
    """
    CREATE TABLE IF NOT EXISTS board_tasks (
      id CHAR(36) PRIMARY KEY,
      association_id CHAR(36) NOT NULL,
      created_by CHAR(36) NOT NULL,
      title VARCHAR(100) NOT NULL,
      description TEXT,
      supervised_by CHAR(36) NOT NULL,
      image_url VARCHAR(255),
      status VARCHAR(50) DEFAULT 'New',
      created_at DATETIME DEFAULT CURRENT_TIMESTAMP,
      updated_at DATETIME DEFAULT CURRENT_TIMESTAMP ON UPDATE CURRENT_TIMESTAMP,
      FOREIGN KEY (association_id) REFERENCES associations(id) ON DELETE CASCADE,
      FOREIGN KEY (created_by) REFERENCES accounts(account_id) ON DELETE CASCADE,
      FOREIGN KEY (supervised_by) REFERENCES accounts(account_id) ON DELETE CASCADE
    ) ENGINE=InnoDB DEFAULT CHARSET=utf8mb4;
    CREATE TABLE IF NOT EXISTS board_task_messages (
      id CHAR(36) PRIMARY KEY,
      board_task_id CHAR(36) NOT NULL,
      sender_id CHAR(36) NOT NULL,
      message TEXT,
      attachment_url VARCHAR(255),
      created_at DATETIME DEFAULT CURRENT_TIMESTAMP,
      FOREIGN KEY (board_task_id) REFERENCES board_tasks(id) ON DELETE CASCADE,
      FOREIGN KEY (sender_id) REFERENCES accounts(account_id) ON DELETE CASCADE
    ) ENGINE=InnoDB DEFAULT CHARSET=utf8mb4;
    """,
    """
    CREATE TABLE IF NOT EXISTS meetings (
      id CHAR(36) PRIMARY KEY,
      association_id CHAR(36) NOT NULL,
      created_by CHAR(36) NOT NULL,
      title VARCHAR(100) NOT NULL,
      meeting_type VARCHAR(100),
      priority VARCHAR(50),
      audience VARCHAR(100),
      agenda VARCHAR(100),
      description TEXT,
      meeting_date DATE,
      meeting_time TIME,
      duration VARCHAR(50),
      venue VARCHAR(100),
      meeting_link VARCHAR(255),
      organizer CHAR(36),
      attachment_url VARCHAR(255),
      status VARCHAR(50) DEFAULT 'Scheduled',
      meeting_minutes TEXT,
      discussed_topic TEXT,
      created_at DATETIME DEFAULT CURRENT_TIMESTAMP,
      updated_at DATETIME DEFAULT CURRENT_TIMESTAMP ON UPDATE CURRENT_TIMESTAMP,
      FOREIGN KEY (association_id) REFERENCES associations(id) ON DELETE CASCADE,
      FOREIGN KEY (created_by) REFERENCES accounts(account_id) ON DELETE CASCADE,
      FOREIGN KEY (organizer) REFERENCES accounts(account_id) ON DELETE CASCADE
    ) ENGINE=InnoDB DEFAULT CHARSET=utf8mb4;
    """,
    """
    CREATE TABLE IF NOT EXISTS polls (
      id CHAR(36) PRIMARY KEY,
      association_id CHAR(36) NULL,
      question VARCHAR(255) NOT NULL,
      description TEXT,
      visibility VARCHAR(50) DEFAULT 'All',
      status VARCHAR(50) DEFAULT 'Published',
      is_multiple_choice TINYINT(1) DEFAULT 0,
      end_date DATETIME,
      created_by CHAR(36) NULL,
      created_at DATETIME DEFAULT CURRENT_TIMESTAMP,
      updated_at DATETIME DEFAULT CURRENT_TIMESTAMP ON UPDATE CURRENT_TIMESTAMP,
      FOREIGN KEY (association_id) REFERENCES associations(id) ON DELETE CASCADE,
      FOREIGN KEY (created_by) REFERENCES accounts(account_id) ON DELETE SET NULL
    ) ENGINE=InnoDB DEFAULT CHARSET=utf8mb4;
    """,
    """
    CREATE TABLE IF NOT EXISTS poll_options (
      id CHAR(36) PRIMARY KEY,
      poll_id CHAR(36) NOT NULL,
      option_text VARCHAR(255) NOT NULL,
      created_at DATETIME DEFAULT CURRENT_TIMESTAMP,
      FOREIGN KEY (poll_id) REFERENCES polls(id) ON DELETE CASCADE
    ) ENGINE=InnoDB DEFAULT CHARSET=utf8mb4;
    """,
    """
    CREATE TABLE IF NOT EXISTS poll_votes (
      poll_id CHAR(36) NOT NULL,
      option_id CHAR(36) NOT NULL,
      account_id CHAR(36) NOT NULL,
      created_at DATETIME DEFAULT CURRENT_TIMESTAMP,
      PRIMARY KEY (option_id, account_id),
      FOREIGN KEY (poll_id) REFERENCES polls(id) ON DELETE CASCADE,
      FOREIGN KEY (option_id) REFERENCES poll_options(id) ON DELETE CASCADE,
      FOREIGN KEY (account_id) REFERENCES accounts(account_id) ON DELETE CASCADE
    ) ENGINE=InnoDB DEFAULT CHARSET=utf8mb4;
    """,
    """
    CREATE TABLE IF NOT EXISTS poll_likes (
      poll_id CHAR(36) NOT NULL,
      account_id CHAR(36) NOT NULL,
      created_at DATETIME DEFAULT CURRENT_TIMESTAMP,
      PRIMARY KEY (poll_id, account_id),
      FOREIGN KEY (poll_id) REFERENCES polls(id) ON DELETE CASCADE,
      FOREIGN KEY (account_id) REFERENCES accounts(account_id) ON DELETE CASCADE
    ) ENGINE=InnoDB DEFAULT CHARSET=utf8mb4;
    """,
    """
    CREATE TABLE IF NOT EXISTS poll_comments (
      id CHAR(36) PRIMARY KEY,
      poll_id CHAR(36) NOT NULL,
      account_id CHAR(36) NOT NULL,
      content TEXT NOT NULL,
      created_at DATETIME DEFAULT CURRENT_TIMESTAMP,
      FOREIGN KEY (poll_id) REFERENCES polls(id) ON DELETE CASCADE,
      FOREIGN KEY (account_id) REFERENCES accounts(account_id) ON DELETE CASCADE
    ) ENGINE=InnoDB DEFAULT CHARSET=utf8mb4;
    """,
    """
    CREATE TABLE IF NOT EXISTS unit_documents (
      id CHAR(36) PRIMARY KEY,
      association_id CHAR(36) NOT NULL,
      unit_id CHAR(36),
      user_id CHAR(36) NOT NULL,
      name VARCHAR(255) NOT NULL,
      type VARCHAR(50) NOT NULL,
      description TEXT,
      file_url VARCHAR(500),
      file_name VARCHAR(255),
      file_type VARCHAR(20),
      file_size_kb INT,
      created_at DATETIME DEFAULT CURRENT_TIMESTAMP,
      FOREIGN KEY (association_id) REFERENCES associations(id) ON DELETE CASCADE,
      FOREIGN KEY (unit_id) REFERENCES units(id) ON DELETE CASCADE,
      FOREIGN KEY (user_id) REFERENCES accounts(account_id) ON DELETE CASCADE
    ) ENGINE=InnoDB DEFAULT CHARSET=utf8mb4;
    """,
    """
    CREATE TABLE IF NOT EXISTS wallets (
      id CHAR(36) PRIMARY KEY,
      account_id CHAR(36) NOT NULL UNIQUE,
      balance DECIMAL(12,2) DEFAULT 0.00,
      reward_points INT DEFAULT 0,
      security_pin VARCHAR(255),
      status VARCHAR(50) DEFAULT 'active',
      created_at DATETIME DEFAULT CURRENT_TIMESTAMP,
      FOREIGN KEY (account_id) REFERENCES accounts(account_id) ON DELETE CASCADE
    ) ENGINE=InnoDB DEFAULT CHARSET=utf8mb4;
    """,
    """
    CREATE TABLE IF NOT EXISTS wallet_transactions (
      id CHAR(36) PRIMARY KEY,
      wallet_id CHAR(36) NOT NULL,
      type VARCHAR(50) NOT NULL,
      amount DECIMAL(12,2) NOT NULL,
      status VARCHAR(50) DEFAULT 'Completed',
      reference_number VARCHAR(100),
      description TEXT,
      created_at DATETIME DEFAULT CURRENT_TIMESTAMP,
      FOREIGN KEY (wallet_id) REFERENCES wallets(id) ON DELETE CASCADE
    ) ENGINE=InnoDB DEFAULT CHARSET=utf8mb4;
    """,
    """
    CREATE TABLE IF NOT EXISTS rewards (
      id CHAR(36) PRIMARY KEY,
      wallet_id CHAR(36) NOT NULL,
      points INT NOT NULL,
      source VARCHAR(100),
      expiry_date DATETIME,
      created_at DATETIME DEFAULT CURRENT_TIMESTAMP,
      FOREIGN KEY (wallet_id) REFERENCES wallets(id) ON DELETE CASCADE
    ) ENGINE=InnoDB DEFAULT CHARSET=utf8mb4;
    """,
    """
    CREATE TABLE IF NOT EXISTS bank_integrations (
      id CHAR(36) PRIMARY KEY,
      association_id CHAR(36),
      provider_name VARCHAR(100) NOT NULL,
      client_id VARCHAR(255),
      client_secret VARCHAR(255),
      api_key VARCHAR(255),
      api_secret VARCHAR(255),
      webhook_url VARCHAR(255),
      webhook_secret VARCHAR(255),
      base_url VARCHAR(255),
      status VARCHAR(50) DEFAULT 'active',
      created_at DATETIME DEFAULT CURRENT_TIMESTAMP
    ) ENGINE=InnoDB DEFAULT CHARSET=utf8mb4;
    """,
    """
    CREATE TABLE IF NOT EXISTS collection_accounts (
      id CHAR(36) PRIMARY KEY,
      association_id CHAR(36),
      bank_name VARCHAR(100) NOT NULL,
      account_number VARCHAR(100) NOT NULL,
      ifsc VARCHAR(50) NOT NULL,
      account_holder_name VARCHAR(150),
      account_type VARCHAR(50) DEFAULT 'Current',
      vpa VARCHAR(100),
      status VARCHAR(50) DEFAULT 'active',
      created_at DATETIME DEFAULT CURRENT_TIMESTAMP
    ) ENGINE=InnoDB DEFAULT CHARSET=utf8mb4;
    """,
    """
    CREATE TABLE IF NOT EXISTS virtual_accounts (
      id CHAR(36) PRIMARY KEY,
      account_id CHAR(36) NOT NULL,
      provider_id CHAR(36),
      virtual_account_id VARCHAR(100),
      virtual_account_number VARCHAR(100) NOT NULL UNIQUE,
      virtual_ifsc VARCHAR(50) NOT NULL,
      virtual_vpa VARCHAR(100),
      beneficiary_name VARCHAR(150),
      reference_number VARCHAR(100),
      status VARCHAR(50) DEFAULT 'active',
      expires_at DATETIME,
      created_at DATETIME DEFAULT CURRENT_TIMESTAMP,
      FOREIGN KEY (account_id) REFERENCES accounts(account_id) ON DELETE CASCADE
    ) ENGINE=InnoDB DEFAULT CHARSET=utf8mb4;
    """,
    """
    CREATE TABLE IF NOT EXISTS wallet_topups (
      id CHAR(36) PRIMARY KEY,
      account_id CHAR(36) NOT NULL,
      amount DECIMAL(12,2) NOT NULL,
      reference_number VARCHAR(100) NOT NULL UNIQUE,
      payment_method VARCHAR(50) DEFAULT 'Virtual Account',
      status VARCHAR(50) DEFAULT 'PENDING',
      created_at DATETIME DEFAULT CURRENT_TIMESTAMP,
      FOREIGN KEY (account_id) REFERENCES accounts(account_id) ON DELETE CASCADE
    ) ENGINE=InnoDB DEFAULT CHARSET=utf8mb4;
    """,
    """
    CREATE TABLE IF NOT EXISTS payment_webhooks (
      id CHAR(36) PRIMARY KEY,
      provider VARCHAR(100) NOT NULL,
      event_type VARCHAR(100),
      transaction_id VARCHAR(100),
      virtual_account_id VARCHAR(100),
      payment_id VARCHAR(100),
      utr VARCHAR(100),
      amount DECIMAL(12,2),
      status VARCHAR(50),
      signature_verified TINYINT(1) DEFAULT 0,
      payload LONGTEXT,
      processed TINYINT(1) DEFAULT 0,
      error_log TEXT,
      received_at DATETIME DEFAULT CURRENT_TIMESTAMP
    ) ENGINE=InnoDB DEFAULT CHARSET=utf8mb4;
    """,
    """
    CREATE TABLE IF NOT EXISTS wallet_ledger (
      id CHAR(36) PRIMARY KEY,
      wallet_id CHAR(36) NOT NULL,
      account_id CHAR(36) NOT NULL,
      transaction_type VARCHAR(50) NOT NULL,
      amount DECIMAL(12,2) NOT NULL,
      balance_after DECIMAL(12,2) NOT NULL,
      reference_id VARCHAR(100),
      payment_id VARCHAR(100),
      utr VARCHAR(100),
      description TEXT,
      status VARCHAR(50) DEFAULT 'SUCCESS',
      created_at DATETIME DEFAULT CURRENT_TIMESTAMP,
      FOREIGN KEY (wallet_id) REFERENCES wallets(id) ON DELETE CASCADE,
      FOREIGN KEY (account_id) REFERENCES accounts(account_id) ON DELETE CASCADE
    ) ENGINE=InnoDB DEFAULT CHARSET=utf8mb4;
    """,
    """
    CREATE TABLE IF NOT EXISTS upi_collections (
      id CHAR(36) PRIMARY KEY,
      account_id CHAR(36) NOT NULL,
      amount DECIMAL(12,2) NOT NULL,
      vpa VARCHAR(100) NOT NULL,
      customer_name VARCHAR(150),
      customer_email VARCHAR(150),
      customer_phone VARCHAR(50),
      request_id VARCHAR(100) NOT NULL UNIQUE,
      txnid VARCHAR(100),
      status VARCHAR(50) DEFAULT 'PENDING',
      expiry_time DATETIME,
      created_at DATETIME DEFAULT CURRENT_TIMESTAMP,
      FOREIGN KEY (account_id) REFERENCES accounts(account_id) ON DELETE CASCADE
    ) ENGINE=InnoDB DEFAULT CHARSET=utf8mb4;
    """,
    """
    CREATE TABLE IF NOT EXISTS bank_reconciliations (
      id CHAR(36) PRIMARY KEY,
      payment_id VARCHAR(100),
      utr VARCHAR(100),
      virtual_account_id VARCHAR(100),
      account_id CHAR(36),
      amount DECIMAL(12,2) NOT NULL,
      bank_amount DECIMAL(12,2),
      wallet_amount DECIMAL(12,2),
      status VARCHAR(50) DEFAULT 'PENDING',
      discrepancy_reason TEXT,
      supporting_evidence TEXT,
      reconciled_at DATETIME,
      created_at DATETIME DEFAULT CURRENT_TIMESTAMP
    ) ENGINE=InnoDB DEFAULT CHARSET=utf8mb4;
    """,
    """
    CREATE TABLE IF NOT EXISTS reconciliation_audits (
      id CHAR(36) PRIMARY KEY,
      reconciliation_id CHAR(36) NOT NULL,
      action VARCHAR(100) NOT NULL,
      old_status VARCHAR(50),
      new_status VARCHAR(50),
      reason TEXT NOT NULL,
      comments TEXT NOT NULL,
      supporting_ref VARCHAR(100),
      performed_by VARCHAR(150) NOT NULL,
      performed_at DATETIME DEFAULT CURRENT_TIMESTAMP,
      FOREIGN KEY (reconciliation_id) REFERENCES bank_reconciliations(id) ON DELETE CASCADE
    ) ENGINE=InnoDB DEFAULT CHARSET=utf8mb4;
    """
]


async def setup_schema():
    for stmt in SCHEMA_SQL:
        await db_execute(stmt)
    
    # Safely alter table to add new columns if they don't exist
    alter_statements = [
        "ALTER TABLE events ADD COLUMN association_id CHAR(36) NULL",
        "ALTER TABLE events ADD COLUMN category VARCHAR(50)",
        "ALTER TABLE events ADD COLUMN banner_url VARCHAR(255)",
        "ALTER TABLE events ADD COLUMN is_registration_required TINYINT(1) DEFAULT 0",
        "ALTER TABLE events ADD COLUMN registration_deadline DATETIME",
        "ALTER TABLE events ADD COLUMN max_capacity INT",
        "ALTER TABLE events ADD COLUMN audience VARCHAR(50) DEFAULT 'All'",
        "ALTER TABLE events ADD COLUMN is_paid TINYINT(1) DEFAULT 0",
        "ALTER TABLE events ADD COLUMN fee_amount DECIMAL(10,2)",
        "ALTER TABLE events ADD COLUMN organizer_name VARCHAR(100)",
        "ALTER TABLE events ADD COLUMN organizer_contact VARCHAR(50)",
        "ALTER TABLE events ADD COLUMN status VARCHAR(50) DEFAULT 'Published'",
        "ALTER TABLE events ADD CONSTRAINT fk_event_assoc FOREIGN KEY (association_id) REFERENCES associations(id) ON DELETE CASCADE"
    ]
    for alt in alter_statements:
        try:
            await db_execute(alt)
        except Exception:
            pass # Ignore duplicate column errors


# ---------- Utilities ----------
def hash_password(password: str) -> str:
    return bcrypt.hashpw(password.encode("utf-8"), bcrypt.gensalt()).decode("utf-8")


def verify_password(plain: str, hashed: str) -> bool:
    return bcrypt.checkpw(plain.encode("utf-8"), hashed.encode("utf-8"))


def create_access_token(account_id: str, email: str, role: str) -> str:
    payload = {
        "sub": str(account_id),
        "email": email,
        "role": role,
        "exp": datetime.now(timezone.utc) + timedelta(minutes=JWT_EXP_MIN),
        "type": "access",
    }
    return jwt.encode(payload, get_jwt_secret(), algorithm=JWT_ALGORITHM)


def generate_code(length: str = 8) -> str:
    alphabet = string.ascii_uppercase + string.digits
    # Avoid ambiguous chars
    alphabet = alphabet.replace("O", "").replace("0", "").replace("I", "").replace("1", "")
    return "".join(random.choices(alphabet, k=length))


def strong_password(pw: str) -> Optional[str]:
    if len(pw) < 8:
        return "Password must be at least 8 characters."
    if not re.search(r"[A-Z]", pw):
        return "Password must contain an uppercase letter."
    if not re.search(r"\d", pw):
        return "Password must contain a number."
    if not re.search(r"[!@#$%^&*()_+\-=\[\]{};:'\",.<>/?`~\\|]", pw):
        return "Password must contain a special character."
    return None


async def log_email(to_email: str, subject: str, body: str):
    """Dev mode: log email to DB + console. Replace with SendGrid when key is provided."""
    await db_execute(
        "INSERT INTO email_log (to_email, subject, body) VALUES (%s, %s, %s)",
        (to_email, subject, body),
    )
    logger.info(f"[EMAIL] to={to_email} subject={subject}")
    logger.info(f"[EMAIL BODY] {body}")


# ---------- Auth Dependencies ----------
async def get_current_account(request: Request) -> dict:
    token = request.cookies.get("access_token")
    if not token:
        auth_header = request.headers.get("Authorization", "")
        if auth_header.startswith("Bearer "):
            token = auth_header[7:]
    if not token:
        raise HTTPException(status_code=401, detail="Not authenticated")
    try:
        payload = jwt.decode(token, get_jwt_secret(), algorithms=[JWT_ALGORITHM])
    except jwt.ExpiredSignatureError:
        raise HTTPException(status_code=401, detail="Token expired")
    except jwt.InvalidTokenError:
        raise HTTPException(status_code=401, detail="Invalid token")
    account_id = str(payload["sub"])
    account = await db_fetchone(
        "SELECT a.account_id, a.user_id, a.employee_id, a.email, a.role_id, r.name as role, r.code as role_code, a.created_at FROM accounts a LEFT JOIN roles r ON a.role_id = r.id WHERE a.account_id=%s",
        (account_id,),
    )
    if not account:
        raise HTTPException(status_code=401, detail="Account not found")
    return account


async def require_admin(account: dict = Depends(get_current_account)) -> dict:
    if account["role"] not in ["Super admin", "Admin"]:
        raise HTTPException(status_code=403, detail="Admin access required")
    return account

def require_role(allowed_roles: list[str]):
    async def role_checker(account: dict = Depends(get_current_account)):
        if not account or not account.get("role"):
            raise HTTPException(status_code=403, detail="Role not assigned")
        
        normalized_account_role = account["role"].lower().replace(" ", "_")
        normalized_allowed = [r.lower().replace(" ", "_") for r in allowed_roles]
        
        if normalized_account_role not in normalized_allowed:
            raise HTTPException(status_code=403, detail=f"Requires one of roles: {', '.join(allowed_roles)}")
        return account
    return role_checker

async def get_board_member_association(account_id: str):
    res = await db_fetchall('''
        SELECT b.association_id 
        FROM accounts ac
        JOIN user_details ud ON ac.user_id = ud.user_id
        JOIN units u ON ud.unit_id = u.id
        JOIN blocks b ON u.block_id = b.id
        WHERE ac.account_id = %s
    ''', (account_id,))
    if res:
        return res[0]['association_id']
    return None

async def check_permission(account: dict, feature_name: str, action: str = "can_view"):
    """Check if the user has permission for a specific feature and action."""
    if account["role"] == "Super admin":
        return True
    
    query = f"""
        SELECT p.{action} FROM role_feature_permissions p
        JOIN roles r ON p.role_id = r.id
        JOIN features f ON p.feature_id = f.id
        WHERE r.name = %s AND f.name = %s
    """
    row = await db_fetchone(query, (account["role"], feature_name))
    if not row or not row[action]:
        raise HTTPException(status_code=403, detail="Permission denied")
    return True


def require_permission(feature_name: str, action: str = "can_view"):
    async def dependency(account: dict = Depends(get_current_account)):
        await check_permission(account, feature_name, action)
        return account
    return dependency


def set_auth_cookie(response: Response, token: str):
    response.set_cookie(
        key="access_token",
        value=token,
        httponly=True,
        secure=False,
        samesite="lax",
        max_age=JWT_EXP_MIN * 60,
        path="/",
    )


# ---------- Pydantic Models ----------
class ValidateCodeIn(BaseModel):
    code: str

    @field_validator("code")
    @classmethod
    def clean(cls, v: str) -> str:
        return v.strip().upper()


class RequestCodeIn(BaseModel):
    name: str = Field(min_length=2, max_length=150)
    email: EmailStr
    contact_number: str = Field(min_length=6, max_length=30)


class CreateAccountIn(BaseModel):
    code: str
    email: EmailStr
    password: str
    confirm_password: str


class UpdateDetailsIn(BaseModel):
    code: str
    requested_name: Optional[str] = None
    requested_address: Optional[str] = None
    requested_email: Optional[str] = None
    requested_contact: Optional[str] = None
    note: Optional[str] = None

class ProfileUpdateIn(BaseModel):
    name: Optional[str] = None
    address: Optional[str] = None
    email: Optional[EmailStr] = None
    contact_number: Optional[str] = None
    alt_contact_number: Optional[str] = None
    profile_pic_url: Optional[str] = None

class FamilyMemberIn(BaseModel):
    name: str
    email: Optional[EmailStr] = None
    contact_number: Optional[str] = None
    alt_contact_number: Optional[str] = None

class EmployeeEducationIn(BaseModel):
    education_level: str
    degree: str
    field_of_study: Optional[str] = None
    institution: str
    board_university: Optional[str] = None
    start_date: Optional[str] = None
    end_date: Optional[str] = None
    currently_studying: Optional[bool] = False
    grade: Optional[str] = None
    location: Optional[str] = None
    description: Optional[str] = None
    certificate_url: Optional[str] = None

class EmployeeExperienceIn(BaseModel):
    job_title: str
    employment_type: str
    company: str
    industry: Optional[str] = None
    location: Optional[str] = None
    work_mode: Optional[str] = None
    start_date: str
    end_date: Optional[str] = None
    currently_working: Optional[bool] = False
    description: Optional[str] = None
    skills: Optional[str] = None
    website_url: Optional[str] = None
    certificate_url: Optional[str] = None

class VehicleIn(BaseModel):
    type: str
    registration_number: str
    insurance_url: Optional[str] = None
    puc_url: Optional[str] = None

class PetIn(BaseModel):
    type: str
    name: str
    breed: str
    vaccinated: bool = False
    vaccination_date: Optional[str] = None
    next_vaccination_reminder: bool = False
    vaccination_certificate_url: Optional[str] = None
    reminder_date: Optional[str] = None


class ForgotPasswordIn(BaseModel):
    email: str

class VerifyOtpIn(BaseModel):
    email: str
    otp: str

class ResetPasswordIn(BaseModel):
    email: str
    otp: str
    new_password: str

class LoginIn(BaseModel):
    email: EmailStr
    password: str


class AdminIssueCodeIn(BaseModel):
    request_id: str


class AdminCreateUserIn(BaseModel):
    first_name: str
    last_name: str
    email: str
    contact_number: str
    role_id: str
    association_id: Optional[str] = None

class AdminCreateMemberIn(BaseModel):
    name: str
    address: str
    email: EmailStr
    contact_number: str

class EmployeeIn(BaseModel):
    first_name: str
    last_name: str
    address_line_1: str
    address_line_2: Optional[str] = None
    city: str
    state: str
    pincode: str
    email: EmailStr
    contact_number: str
    emergency_contact_name: str
    emergency_contact_number: str
    id_proof_url: Optional[str] = None
    role_id: str
    association_ids: List[str]
    onboard_date: Optional[str] = None
    end_date: Optional[str] = None


class AnnouncementIn(BaseModel):
    association_id: str | None = None
    title: str = Field(min_length=2, max_length=100)
    body: str = Field(min_length=2, max_length=5000)
    category: str = Field(default="general")
    pinned: bool = False
    audience: str = Field(default="Homeowners")
    attachment_url: Optional[str] = None
    audience: str = Field(default="Homeowners")
    attachment_url: Optional[str] = None

    @field_validator("category")
    @classmethod
    def _cat(cls, v: str) -> str:
        v = (v or "general").lower()
        if v not in ("general", "maintenance", "urgent", "celebration"):
            raise ValueError("Invalid category")
        return v

class CommentIn(BaseModel):
    comment: str = Field(min_length=1, max_length=2000)


class EventIn(BaseModel):
    title: str = Field(min_length=2, max_length=200)
    description: Optional[str] = None
    category: Optional[str] = None
    banner_url: Optional[str] = None
    association_id: Optional[str] = None
    location: Optional[str] = None
    starts_at: datetime
    ends_at: Optional[datetime] = None
    is_registration_required: bool = False
    registration_deadline: Optional[datetime] = None
    max_capacity: Optional[int] = None
    audience: Optional[str] = "All"
    is_paid: bool = False
    fee_amount: Optional[float] = None
    organizer_name: Optional[str] = None
    organizer_contact: Optional[str] = None
    status: Optional[str] = "Published"

class PollOptionIn(BaseModel):
    text: str

class PollIn(BaseModel):
    question: str
    description: Optional[str] = None
    visibility: Optional[str] = 'All'
    status: Optional[str] = 'Published'
    is_multiple_choice: Optional[bool] = False
    end_date: Optional[str] = None
    association_id: Optional[str] = None
    options: List[PollOptionIn]

class PollVoteIn(BaseModel):
    option_ids: List[str]

class RsvpIn(BaseModel):
    status: str = Field(default="going")

    @field_validator("status")
    @classmethod
    def _st(cls, v: str) -> str:
        v = (v or "going").lower()
        if v not in ("going", "maybe", "not_going"):
            raise ValueError("status must be going, maybe, or not_going")
        return v


class EntityTypeIn(BaseModel):
    name: str
    description: Optional[str] = None


class EntityIn(BaseModel):
    entity_type_id: str
    association_id: Optional[str] = None
    name: str
    description: Optional[str] = None


class RoleIn(BaseModel):
    entity_id: Optional[str] = None
    name: str
    code: str
    description: Optional[str] = None
    is_active: bool = True


class FeatureIn(BaseModel):
    name: str
    code: str
    description: Optional[str] = None
    parent_id: Optional[str] = None
    icon: Optional[str] = None
    url: Optional[str] = None
    is_active: bool = True


class RoleFeaturePermissionIn(BaseModel):
    role_id: str
    feature_id: str
    can_create: bool = False
    can_view: bool = False
    can_update: bool = False
    can_delete: bool = False
    sidebar_order: int = Field(default=0, ge=0)


class FeaturePermissionItem(BaseModel):
    feature_id: str
    can_create: bool = False
    can_view: bool = False
    can_update: bool = False
    can_delete: bool = False
    sidebar_order: int = Field(default=0, ge=0)


class BulkRolePermissionsIn(BaseModel):
    permissions: List[FeaturePermissionItem]


# ---------- App ----------
app = FastAPI(title="Nestora API")
api_router = APIRouter(prefix="/api")


async def check_expired_board_members():
    while True:
        try:
            # fetch expired board members with active status
            expired = await db_fetchall('''
                SELECT bm.id, bm.account_id 
                FROM board_members bm
                WHERE bm.status = 'active' AND bm.term_end_date < CURRENT_DATE()
            ''')
            if expired:
                role = await db_fetchone("SELECT id FROM roles WHERE name='Homeowner'")
                homeowner_role_id = role["id"]
                for b in expired:
                    # Update status to past
                    await db_execute("UPDATE board_members SET status = 'past' WHERE id = %s", (b["id"],))
                    # Update account role to homeowner
                    await db_execute("UPDATE accounts SET role_id = %s WHERE account_id = %s", (homeowner_role_id, b["account_id"]))
                    logger.info(f"Auto-reverted board member {b['account_id']} to Homeowner")
        except Exception as e:
            logger.error(f"Error checking expired board members: {e}")
            
        await asyncio.sleep(3600)  # Check every hour

@app.on_event("startup")
async def on_startup():
    await init_pool()
    # Database setup and seeding should now be run manually via `python init_db.py`
    # await setup_schema()
    # await seed_database()
    # await seed_demo_codes()
    # await seed_demo_content()
    asyncio.create_task(check_expired_board_members())
    logger.info("Nestora API started (MySQL connected).")


@app.on_event("shutdown")
async def on_shutdown():
    if pool:
        pool.close()
        await pool.wait_closed()


async def seed_database():
    # 1. Seed Entity Types & Entities
    await db_execute("INSERT IGNORE INTO entity_types (name) VALUES ('HOA')")
    entity_type = await db_fetchone("SELECT id FROM entity_types WHERE name='HOA'")
    entity_type_id = entity_type['id'] if entity_type else 1
        
    await db_execute("INSERT IGNORE INTO entities (entity_type_id, name) VALUES (%s, 'Nestora Default HOA')", (entity_type_id,))
    entity = await db_fetchone("SELECT id FROM entities WHERE name='Nestora Default HOA'")
    entity_id = entity['id'] if entity else 1

    # 2. Seed Roles
    roles = ['Super admin', 'Admin', 'CSR', 'Accountant', 'Security', 'Board member', 'Homeowner', 'Tenant']
    for role in roles:
        await db_execute("INSERT INTO roles (entity_id, name) SELECT %s, %s WHERE NOT EXISTS (SELECT 1 FROM roles WHERE name=%s)", (entity_id, role, role))

    # 3. Seed Admin and other role test users
    default_password = hash_password("Password123!")
    for role in roles:
        role_record = await db_fetchone("SELECT id FROM roles WHERE name=%s", (role,))
        if not role_record:
            continue
            
        role_id = role_record['id']
        if role == 'Super admin':
            email = os.environ["ADMIN_EMAIL"].lower()
            password = os.environ["ADMIN_PASSWORD"]
            hashed = hash_password(password)
        else:
            email = f"{role.lower().replace(' ', '')}@nestora.io"
            hashed = default_password
            
        row = await db_fetchone("SELECT account_id, password_hash FROM accounts WHERE email=%s", (email,))
        if not row:
            await db_execute(
                "INSERT INTO accounts (email, password_hash, role_id) VALUES (%s, %s, %s)",
                (email, hashed, role_id),
            )
            logger.info(f"Seeded test user for {role}: {email}")
        elif role == 'Super admin':
            if not verify_password(password, row["password_hash"]):
                await db_execute(
                    "UPDATE accounts SET password_hash=%s WHERE account_id=%s",
                    (hashed, row["account_id"]),
                )
                logger.info("Admin password re-hashed from env.")


async def seed_demo_codes():
    """Seed 3 demo codes with sample member details for a smooth first-run demo."""
    count = await db_fetchone("SELECT COUNT(*) AS c FROM user_codes")
    if count and count["c"] > 0:
        return
    demo = [
        ("NST-DEMO1", "Aarav Sharma", "12 Marigold Lane, Bengaluru, KA 560001", "aarav.demo@nestora.io", "+91-9876543210"),
        ("NST-DEMO2", "Priya Iyer", "Villa 7, Ashwin Heights, Mumbai, MH 400050", "priya.demo@nestora.io", "+91-9812345678"),
        ("NST-DEMO3", "Rohan Kapoor", "Flat 302, Willow Court, Pune, MH 411006", "rohan.demo@nestora.io", "+91-9900112233"),
    ]
    for code, name, addr, email, phone in demo:
        code_id = await db_execute(
            "INSERT INTO user_codes (login_code, status) VALUES (%s, 'active')", (code,)
        )
        await db_execute(
            "INSERT INTO user_details (code_id, name, address, email, contact_number) VALUES (%s,%s,%s,%s,%s)",
            (code_id, name, addr, email, phone),
        )
    logger.info("Seeded demo access codes: NST-DEMO1, NST-DEMO2, NST-DEMO3")


async def seed_demo_content():
    """Seed 2 sample announcements + 2 upcoming events for a lively first-run experience."""
    admin = await db_fetchone("SELECT a.account_id FROM accounts a JOIN roles r ON a.role_id = r.id WHERE r.name='Super admin' LIMIT 1")
    admin_id = admin["account_id"] if admin else None

    ann_count = await db_fetchone("SELECT COUNT(*) AS c FROM announcements")
    if ann_count and ann_count["c"] == 0:
        await db_execute(
            "INSERT INTO announcements (title, body, category, pinned, created_by) VALUES (%s,%s,%s,%s,%s)",
            (
                "Community garden refresh — plans dropping this Sunday",
                "We're kicking off the seasonal landscape refresh this weekend. Join us at the clubhouse lawn on Sunday at 10 AM for the walkthrough. Coffee & croissants on us.",
                "general",
                1,
                admin_id,
            ),
        )
        await db_execute(
            "INSERT INTO announcements (title, body, category, pinned, created_by) VALUES (%s,%s,%s,%s,%s)",
            (
                "Scheduled water tank cleaning — Tuesday, 8 AM",
                "Water supply will be interrupted for approximately two hours starting Tuesday at 8 AM. Please store water accordingly. Thank you for your patience.",
                "maintenance",
                0,
                admin_id,
            ),
        )
        logger.info("Seeded demo announcements.")

    ev_count = await db_fetchone("SELECT COUNT(*) AS c FROM events")
    if ev_count and ev_count["c"] == 0:
        now = datetime.now(timezone.utc)
        await db_execute(
            "INSERT INTO events (title, description, location, starts_at, ends_at, created_by) VALUES (%s,%s,%s,%s,%s,%s)",
            (
                "Monthly community meet",
                "Our monthly get-together — light refreshments, quick updates, and a chance to meet new neighbours.",
                "Clubhouse · East Wing",
                now + timedelta(days=6, hours=18),
                now + timedelta(days=6, hours=20),
                admin_id,
            ),
        )
        await db_execute(
            "INSERT INTO events (title, description, location, starts_at, ends_at, created_by) VALUES (%s,%s,%s,%s,%s,%s)",
            (
                "Yoga in the courtyard",
                "Bring your own mat. Beginner-friendly session led by a certified instructor.",
                "Central courtyard",
                now + timedelta(days=2, hours=6, minutes=30),
                now + timedelta(days=2, hours=7, minutes=30),
                admin_id,
            ),
        )
        logger.info("Seeded demo events.")


# ---------- Routes: Public ----------
@api_router.get("/")
async def root():
    return {"message": "Nestora API", "status": "ok"}


@api_router.post("/login-code")
async def validate_login_code(payload: ValidateCodeIn):
    code = payload.code
    if not re.fullmatch(r"[A-Z0-9\-!@#$%^&*]{4,20}", code):
        raise HTTPException(status_code=400, detail="Invalid code format.")
    row = await db_fetchone(
        "SELECT id, status, expires_at FROM user_codes WHERE login_code=%s", (code,)
    )
    if not row:
        raise HTTPException(status_code=404, detail="Invalid Code. Please try again.")
    if row["status"] != "active":
        raise HTTPException(status_code=400, detail="This code is no longer active.")
    if row["expires_at"] and row["expires_at"] < datetime.now():
        raise HTTPException(status_code=400, detail="This code has expired.")

    # Check if this code is already used by an account
    linked = await db_fetchone(
        "SELECT a.account_id FROM accounts a JOIN user_details ud ON ud.user_id=a.user_id WHERE ud.code_id=%s",
        (row["id"],),
    )
    return {"valid": True, "code_id": row["id"], "already_registered": bool(linked)}


@api_router.post("/request-code")
async def request_code(payload: RequestCodeIn):
    existing = await db_fetchone(
        "SELECT id FROM code_requests WHERE email=%s AND status='pending'", (payload.email.lower(),)
    )
    if existing:
        raise HTTPException(status_code=409, detail="A request with this email is already pending.")
    await db_execute(
        "INSERT INTO code_requests (name, email, contact_number) VALUES (%s,%s,%s)",
        (payload.name.strip(), payload.email.lower(), payload.contact_number.strip()),
    )
    
    # Also generate a public service request
    count_res = await db_fetchone('SELECT COUNT(id) as cnt FROM service_requests')
    sr_count = (count_res['cnt'] + 1) if count_res else 1
    sr_display_id = f"SR-PUB-#{sr_count}"
    
    desc = f"Name: {payload.name.strip()}\nEmail: {payload.email.lower()}\nContact Number: {payload.contact_number.strip()}"
    
    await db_execute('''
        INSERT INTO service_requests (user_id, association_id, unit_id, sr_display_id, service_type, custom_title, description, status)
        VALUES (NULL, NULL, NULL, %s, 'Access Code Request', %s, %s, 'New')
    ''', (sr_display_id, f"Code Request: {payload.name.strip()}", desc))

    await log_email(
        to_email=payload.email,
        subject="Nestora — We've received your access request",
        body=f"Hi {payload.name},\n\nThanks for requesting access to Nestora. An administrator will review your request shortly and email you an access code.\n\n— Nestora Community Team",
    )
    return {"ok": True, "message": "Request submitted. You'll receive your access code via email once approved."}


@api_router.get("/user-details")
async def get_user_details_by_code(code: str):
    code = code.strip().upper()
    row = await db_fetchone(
        """
        SELECT ud.user_id, ud.code_id, ud.name, ud.address, ud.email, ud.contact_number, uc.login_code,
               assoc.name as association_name,
               assoc.address_line_1, assoc.address_line_2,
               assoc.city, assoc.state, assoc.pincode,
               b.name as block_name, un.unit_number
        FROM user_details ud 
        JOIN user_codes uc ON uc.id=ud.code_id
        LEFT JOIN units un ON ud.unit_id = un.id
        LEFT JOIN blocks b ON un.block_id = b.id
        LEFT JOIN associations assoc ON b.association_id = assoc.id OR ud.association_id = assoc.id
        WHERE uc.login_code=%s
        """,
        (code,),
    )
    if not row:
        raise HTTPException(status_code=404, detail="Details not found for this code.")
        
    if not row["address"]:
        parts = []
        if row.get("block_name") and row.get("unit_number"):
            parts.append(f"Block {row['block_name']} - Unit {row['unit_number']}")
        elif row.get("association_name"):
            parts.append(row["association_name"])
            
        if row.get("address_line_1"):
            parts.append(row["address_line_1"])
        if row.get("address_line_2"):
            parts.append(row["address_line_2"])
            
        location_parts = []
        if row.get("city"):
            location_parts.append(row["city"])
        if row.get("state"):
            location_parts.append(row["state"])
            
        location_str = ", ".join(location_parts)
        if location_str and row.get("pincode"):
            location_str += f" - {row['pincode']}"
        elif row.get("pincode"):
            location_str = row["pincode"]
            
        if location_str:
            parts.append(location_str)
        
        if parts:
            row["address"] = ", ".join(parts)
        
    return row


@api_router.post("/create-account")
async def create_account(payload: CreateAccountIn, response: Response):
    code = payload.code.strip().upper()
    if payload.password != payload.confirm_password:
        raise HTTPException(status_code=400, detail="Passwords do not match.")
    err = strong_password(payload.password)
    if err:
        raise HTTPException(status_code=400, detail=err)

    code_row = await db_fetchone(
        "SELECT id, status FROM user_codes WHERE login_code=%s", (code,)
    )
    if not code_row or code_row["status"] != "active":
        raise HTTPException(status_code=400, detail="Access code is invalid or inactive.")

    details = await db_fetchone(
        "SELECT user_id, role_id FROM user_details WHERE code_id=%s", (code_row["id"],)
    )
    if not details:
        raise HTTPException(status_code=400, detail="No member details linked to this code.")

    exists = await db_fetchone(
        "SELECT account_id FROM accounts WHERE email=%s OR user_id=%s",
        (payload.email.lower(), details["user_id"]),
    )
    if exists:
        raise HTTPException(status_code=409, detail="An account already exists for this email or member.")

    role_id = details.get("role_id")
    if not role_id:
        role = await db_fetchone("SELECT id FROM roles WHERE name='Homeowner'")
        role_id = role["id"] if role else None
        
    account_id = await db_execute(
        "INSERT INTO accounts (user_id, email, password_hash, role_id) VALUES (%s,%s,%s,%s)",
        (details["user_id"], payload.email.lower(), hash_password(payload.password), role_id),
    )
    await db_execute("UPDATE user_codes SET status='used' WHERE id=%s", (code_row["id"],))

    # Create a wallet for the newly verified user
    await db_execute(
        "INSERT INTO wallets (id, account_id, balance, reward_points) VALUES (UUID(), %s, 0, 0)",
        (account_id,)
    )
    try:
        await ensure_virtual_account(account_id)
    except Exception as e:
        logger.warning(f"Auto VA generation warning: {e}")

    token = create_access_token(account_id, payload.email.lower(), "Homeowner")
    set_auth_cookie(response, token)
    return {
        "ok": True,
        "account": {"account_id": account_id, "email": payload.email.lower(), "role": "Homeowner"},
        "token": token,
    }


@api_router.post("/update-details-request")
async def update_details_request(payload: UpdateDetailsIn):
    code = payload.code.strip().upper()
    row = await db_fetchone("SELECT id FROM user_codes WHERE login_code=%s", (code,))
    if not row:
        raise HTTPException(status_code=404, detail="Invalid code.")
    await db_execute(
        """INSERT INTO update_requests
           (code_id, requested_name, requested_address, requested_email, requested_contact, note)
           VALUES (%s,%s,%s,%s,%s,%s)""",
        (
            row["id"],
            payload.requested_name,
            payload.requested_address,
            payload.requested_email,
            payload.requested_contact,
            payload.note,
        ),
    )
    
    # Generate service request
    ud_row = await db_fetchone('''
        SELECT ud.user_id, ud.unit_id, b.association_id 
        FROM user_details ud
        LEFT JOIN units u ON ud.unit_id = u.id
        LEFT JOIN blocks b ON u.block_id = b.id
        WHERE ud.code_id = %s
    ''', (row["id"],))
    
    count_res = await db_fetchone('SELECT COUNT(id) as cnt FROM service_requests')
    sr_count = (count_res['cnt'] + 1) if count_res else 1
    sr_display_id = f"SR-PUB-#{sr_count}"
    
    desc = f"Name: {payload.requested_name}\nEmail: {payload.requested_email}\nContact Number: {payload.requested_contact}\nAddress: {payload.requested_address}\nNote: {payload.note}"
    
    user_id = ud_row["user_id"] if ud_row else None
    unit_id = ud_row["unit_id"] if ud_row else None
    association_id = ud_row["association_id"] if ud_row else None
    
    await db_execute('''
        INSERT INTO service_requests (user_id, association_id, unit_id, sr_display_id, service_type, custom_title, description, status)
        VALUES (%s, %s, %s, %s, 'Detail Correction Request', %s, %s, 'New')
    ''', (user_id, association_id, unit_id, sr_display_id, f"Correction: {payload.requested_name}", desc))

    await log_email(
        to_email=os.environ.get("SUPPORT_EMAIL", "support@nestora.io"),
        subject=f"Nestora — Detail correction request for code {code}",
        body=(
            f"A member reported incorrect details.\n\n"
            f"Code: {code}\n"
            f"Requested Name: {payload.requested_name}\n"
            f"Requested Address: {payload.requested_address}\n"
            f"Requested Email: {payload.requested_email}\n"
            f"Requested Contact: {payload.requested_contact}\n"
            f"Note: {payload.note or '-'}\n"
        ),
    )
    return {"ok": True, "message": "Your correction request has been sent to our support team."}


@api_router.post("/auth/forgot-password")
async def forgot_password(payload: ForgotPasswordIn):
    account = await db_fetchone("SELECT account_id FROM accounts WHERE email = %s", (payload.email.lower(),))
    if not account:
        # For security, you might want to just return success even if email not found, 
        # but returning error makes testing easier.
        raise HTTPException(status_code=404, detail="Email not registered.")
    
    otp = "".join(random.choices(string.digits, k=6))
    expires_at = datetime.now() + timedelta(minutes=5)
    
    # Store OTP (upsert)
    await db_execute(
        "INSERT INTO password_reset_otps (email, otp, expires_at) VALUES (%s, %s, %s) ON DUPLICATE KEY UPDATE otp=VALUES(otp), expires_at=VALUES(expires_at)",
        (payload.email.lower(), otp, expires_at)
    )
    
    await log_email(
        payload.email.lower(),
        "Password Reset OTP",
        f"Your password reset OTP is {otp}. It will expire in 5 minutes."
    )
    return {"ok": True, "message": "OTP sent successfully"}

@api_router.post("/auth/verify-otp")
async def verify_otp(payload: VerifyOtpIn):
    record = await db_fetchone("SELECT otp, expires_at FROM password_reset_otps WHERE email = %s", (payload.email.lower(),))
    if not record:
        raise HTTPException(status_code=400, detail="No OTP requested for this email.")
    
    if record["expires_at"] < datetime.now():
        raise HTTPException(status_code=400, detail="OTP has expired.")
        
    if record["otp"] != payload.otp:
        raise HTTPException(status_code=400, detail="Invalid OTP.")
        
    return {"ok": True, "message": "OTP verified successfully"}

@api_router.post("/auth/reset-password")
async def reset_password(payload: ResetPasswordIn):
    record = await db_fetchone("SELECT otp, expires_at FROM password_reset_otps WHERE email = %s", (payload.email.lower(),))
    if not record or record["otp"] != payload.otp or record["expires_at"] < datetime.now():
        raise HTTPException(status_code=400, detail="Invalid or expired OTP.")
        
    pw_err = strong_password(payload.new_password)
    if pw_err:
        raise HTTPException(status_code=400, detail=pw_err)
        
    hashed = hash_password(payload.new_password)
    
    await db_execute("UPDATE accounts SET password_hash = %s WHERE email = %s", (hashed, payload.email.lower()))
    await db_execute("DELETE FROM password_reset_otps WHERE email = %s", (payload.email.lower(),))
    
    return {"ok": True, "message": "Password reset successfully"}

@api_router.post("/auth/login")
async def login(payload: LoginIn, response: Response):
    row = await db_fetchone(
        "SELECT a.account_id, a.email, a.password_hash, r.name as role, r.code as role_code, a.user_id, a.employee_id FROM accounts a LEFT JOIN roles r ON a.role_id = r.id WHERE a.email=%s",
        (payload.email.lower(),),
    )
    if not row or not verify_password(payload.password, row["password_hash"]):
        raise HTTPException(status_code=401, detail="Invalid email or password.")

    # Check if association ended
    if row["role"] not in ["Super admin", "Admin"]:
        assoc = None
        if row["user_id"]:
            assoc = await db_fetchone("""
                SELECT a.end_date FROM user_details ud
                JOIN units u ON ud.unit_id = u.id
                JOIN blocks b ON u.block_id = b.id
                JOIN associations a ON b.association_id = a.id
                WHERE ud.user_id = %s
            """, (row["user_id"],))
        elif row["employee_id"]:
            assoc = await db_fetchone("""
                SELECT end_date FROM employees
                WHERE employee_id = %s
            """, (row["employee_id"],))
        if assoc and assoc.get("end_date"):
            from datetime import datetime
            if str(assoc["end_date"])[:10] <= str(datetime.now().date()):
                raise HTTPException(status_code=403, detail="Association access has ended.")
    # Auto-create wallet if it doesn't exist
    wallet = await db_fetchone("SELECT id FROM wallets WHERE account_id=%s", (row["account_id"],))
    if not wallet:
        await db_execute("INSERT INTO wallets (account_id) VALUES (%s)", (row["account_id"],))

    token = create_access_token(row["account_id"], row["email"], row["role"])
    set_auth_cookie(response, token)
    return {
        "ok": True,
        "account": {"account_id": row["account_id"], "email": row["email"], "role": row["role"]},
        "token": token,
    }


@api_router.post("/auth/logout")
async def logout(response: Response):
    response.delete_cookie("access_token", path="/")
    return {"ok": True}


@api_router.get("/auth/me")
async def me(account: dict = Depends(get_current_account)):
    profile = None
    
    # Fetch all feature icons for the sidebar
    all_features = await db_fetchall("SELECT name, icon FROM features WHERE icon IS NOT NULL")
    account["feature_icons"] = {f["name"]: f["icon"] for f in all_features}

    if account.get("user_id"):
        profile = await db_fetchone(
            "SELECT name, address, email, contact_number FROM user_details WHERE user_id=%s",
            (account["user_id"],),
        )
        # Fetch association info for Homeowner/Board member
        assoc_info = await db_fetchone("""
            SELECT a.id as association_id, a.country as association_country, a.subscription_status, a.current_plan_id
            FROM user_details ud
            JOIN units u ON ud.unit_id = u.id
            JOIN blocks b ON u.block_id = b.id
            JOIN associations a ON b.association_id = a.id
            WHERE ud.user_id = %s
        """, (account["user_id"],))

        if not assoc_info and account["role"] in ["Committee member", "Board member", "Committee Member", "Board Member"]:
            assoc_info = await db_fetchone("""
                SELECT a.id as association_id, a.country as association_country, a.subscription_status, a.current_plan_id
                FROM committee_members cm
                JOIN committees c ON cm.committee_id = c.id
                JOIN associations a ON c.association_id = a.id
                WHERE cm.user_id = %s
                LIMIT 1
            """, (account["user_id"],))

        if not assoc_info and account["role"] in ["Board member", "Board Member"]:
            assoc_info = await db_fetchone("""
                SELECT a.id as association_id, a.country as association_country, a.subscription_status, a.current_plan_id
                FROM board_members bm
                JOIN associations a ON bm.association_id = a.id
                WHERE bm.account_id = %s
                LIMIT 1
            """, (account["account_id"],))
        
        if assoc_info:
            account["association_id"] = assoc_info["association_id"]
            account["association_country"] = assoc_info["association_country"]
            account["subscription_status"] = assoc_info["subscription_status"]
            
            # Fetch assessment rules
            assessment = await db_fetchone("SELECT default_amount, frequency, due_day_of_month FROM assessment_rules WHERE association_id = %s", (assoc_info["association_id"],))
            base_amount = float(assessment["default_amount"]) if assessment and assessment["default_amount"] else 0.0
            due_day = int(assessment["due_day_of_month"]) if assessment and assessment["due_day_of_month"] else 1
            
            account["assessment_amount"] = base_amount
            account["assessment_frequency"] = assessment["frequency"] if assessment else "Monthly"
            account["assessment_due_day"] = due_day
            
            # Dynamic dues calculation
            account["assessment_fine"] = 0.0
            account["assessment_total_due"] = base_amount
            account["assessment_paid_this_month"] = False
            
            if base_amount > 0:
                # Get the user's unit_id
                user_unit = await db_fetchone("SELECT unit_id FROM user_details WHERE user_id = %s", (account["user_id"],))
                unit_id = user_unit["unit_id"] if user_unit else None
                
                payment = None
                now = datetime.now()
                current_month_start = datetime(now.year, now.month, 1).strftime('%Y-%m-%d 00:00:00')

                if unit_id:
                    # Check if ANY user in this unit paid this month
                    payment = await db_fetchone("""
                        SELECT wt.id 
                        FROM wallet_transactions wt
                        JOIN wallets w ON wt.wallet_id = w.id
                        JOIN accounts a ON w.account_id = a.account_id
                        JOIN user_details ud ON a.user_id = ud.user_id
                        WHERE ud.unit_id = %s
                          AND wt.type IN ('Debit', 'UPI') 
                          AND wt.description = 'Maintenance Dues Paid' 
                          AND wt.created_at >= %s
                        LIMIT 1
                    """, (unit_id, current_month_start))
                else:
                    # Fallback to current user's wallet
                    wallet = await db_fetchone("SELECT id FROM wallets WHERE account_id = %s", (account["account_id"],))
                    if wallet:
                        payment = await db_fetchone(
                            "SELECT id FROM wallet_transactions WHERE wallet_id = %s AND type IN ('Debit', 'UPI') AND description = 'Maintenance Dues Paid' AND created_at >= %s",
                            (wallet["id"], current_month_start)
                        )
                
                if payment:
                    account["assessment_paid_this_month"] = True
                    account["assessment_total_due"] = 0.0
                else:
                    # Calculate late fee
                    today = now.day
                    if today > due_day:
                        fine_rule = await db_fetchone(
                            "SELECT amount FROM fine_rules WHERE association_id = %s AND fine_type = 'Late Payment Fee'",
                            (assoc_info["association_id"],)
                        )
                        if fine_rule and fine_rule["amount"]:
                            days_late = today - due_day
                            fine_amount = days_late * float(fine_rule["amount"])
                            account["assessment_fine"] = fine_amount
                            account["assessment_total_due"] = base_amount + fine_amount
            
            features = []
            if assoc_info["current_plan_id"]:
                f_rows = await db_fetchall("""
                    SELECT f.name FROM subscription_plan_features spf
                    JOIN features f ON spf.feature_id = f.id
                    WHERE spf.plan_id = %s
                """, (assoc_info["current_plan_id"],))
                features = [r["name"] for r in f_rows]
            account["allowed_features"] = features
    elif account.get("employee_id"):
        profile = await db_fetchone(
            "SELECT name, first_name, last_name, email, contact_number FROM employees WHERE employee_id=%s",
            (account["employee_id"],),
        )

    # Fetch role-based feature permissions
    if account.get("role_id"):
        perms = await db_fetchall("""
            SELECT f.id as feature_id, f.name as feature_name, f.code as feature_code, 
                   f.parent_id, f.icon, f.url, f.order_index,
                   p.can_create, p.can_view, p.can_update, p.can_delete, p.sidebar_order
            FROM role_feature_permissions p
            JOIN features f ON p.feature_id = f.id
            WHERE p.role_id = %s 
              AND f.is_active = 1
              AND (p.can_create = 1 OR p.can_view = 1 OR p.can_update = 1 OR p.can_delete = 1)
            ORDER BY p.sidebar_order ASC, f.name ASC
        """, (account["role_id"],))
        
        account["role_permissions"] = [
            {
                "feature_id": p["feature_id"],
                "feature_name": p["feature_name"],
                "feature_code": p.get("feature_code"),
                "parent_id": p.get("parent_id"),
                "icon": p.get("icon"),
                "url": p.get("url"),
                "order_index": p.get("order_index", 0),
                "sidebar_order": p.get("sidebar_order", 0),
                "can_create": bool(p["can_create"]),
                "can_view": bool(p["can_view"]),
                "can_update": bool(p["can_update"]),
                "can_delete": bool(p["can_delete"]),
            }
            for p in perms
        ]
    elif account.get("role") == "Super admin":
        all_feats = await db_fetchall("""
            SELECT f.id as feature_id, f.name as feature_name, f.code as feature_code, 
                   f.parent_id, f.icon, f.url, f.order_index
            FROM features f
            WHERE f.is_active = 1
            ORDER BY f.order_index ASC, f.name ASC
        """)
        account["role_permissions"] = [
            {
                "feature_id": f["feature_id"],
                "feature_name": f["feature_name"],
                "feature_code": f.get("feature_code"),
                "parent_id": f.get("parent_id"),
                "icon": f.get("icon"),
                "url": f.get("url"),
                "order_index": f.get("order_index", 0),
                "sidebar_order": f.get("order_index", 0),
                "can_create": True,
                "can_view": True,
                "can_update": True,
                "can_delete": True,
            }
            for f in all_feats
        ]
    else:
        account["role_permissions"] = []
            
    return {"account": account, "profile": profile}


# ---------- Routes: RBAC Admin ----------
@api_router.get("/admin/entity-types")
async def get_entity_types(_: dict = Depends(require_admin)):
    return await db_fetchall("SELECT * FROM entity_types ORDER BY created_at DESC")

@api_router.post("/admin/entity-types")
async def create_entity_type(payload: EntityTypeIn, _: dict = Depends(require_admin)):
    existing = await db_fetchone("SELECT id FROM entity_types WHERE name=%s", (payload.name,))
    if existing: raise HTTPException(400, "Entity Type name already exists.")
    type_id = await db_execute("INSERT INTO entity_types (name, description) VALUES (%s, %s)", (payload.name, payload.description))
    return {"id": type_id, **payload.dict()}

@api_router.put("/admin/entity-types/{item_id}")
async def update_entity_type(item_id: str, payload: EntityTypeIn, _: dict = Depends(require_admin)):
    existing = await db_fetchone("SELECT id FROM entity_types WHERE name=%s AND id!=%s", (payload.name, item_id))
    if existing: raise HTTPException(400, "Entity Type name already exists.")
    await db_execute("UPDATE entity_types SET name=%s, description=%s WHERE id=%s", (payload.name, payload.description, item_id))
    return {"ok": True}

@api_router.delete("/admin/entity-types/{item_id}")
async def delete_entity_type(item_id: str, _: dict = Depends(require_admin)):
    await db_execute("DELETE FROM entity_types WHERE id=%s", (item_id,))
    return {"ok": True}

@api_router.get("/admin/entities")
async def get_entities(_: dict = Depends(require_admin)):
    return await db_fetchall("""
        SELECT e.*, et.name as entity_type_name 
        FROM entities e 
        LEFT JOIN entity_types et ON e.entity_type_id = et.id 
        ORDER BY e.created_at DESC
    """)

@api_router.post("/admin/entities")
async def create_entity(payload: EntityIn, _: dict = Depends(require_admin)):
    existing = await db_fetchone("SELECT id FROM entities WHERE name=%s AND entity_type_id=%s", (payload.name, payload.entity_type_id))
    if existing: raise HTTPException(400, "Entity name already exists for this type.")
    entity_id = await db_execute("INSERT INTO entities (entity_type_id, association_id, name, description) VALUES (%s, %s, %s, %s)", 
                                 (payload.entity_type_id, payload.association_id, payload.name, payload.description))
    return {"id": entity_id, **payload.dict()}

@api_router.put("/admin/entities/{item_id}")
async def update_entity(item_id: str, payload: EntityIn, _: dict = Depends(require_admin)):
    existing = await db_fetchone("SELECT id FROM entities WHERE name=%s AND entity_type_id=%s AND id!=%s", (payload.name, payload.entity_type_id, item_id))
    if existing: raise HTTPException(400, "Entity name already exists for this type.")
    await db_execute("UPDATE entities SET entity_type_id=%s, association_id=%s, name=%s, description=%s WHERE id=%s", 
                     (payload.entity_type_id, payload.association_id, payload.name, payload.description, item_id))
    return {"ok": True}

@api_router.delete("/admin/entities/{item_id}")
async def delete_entity(item_id: str, _: dict = Depends(require_admin)):
    await db_execute("DELETE FROM entities WHERE id=%s", (item_id,))
    return {"ok": True}

@api_router.get("/admin/roles")
async def get_roles(_: dict = Depends(require_admin)):
    return await db_fetchall("""
        SELECT r.*, e.name AS entity_name 
        FROM roles r 
        LEFT JOIN entities e ON r.entity_id = e.id 
        ORDER BY r.created_at DESC
    """)

@api_router.post("/admin/roles")
async def create_role(payload: RoleIn, _: dict = Depends(require_admin)):
    clean_name = payload.name.strip()
    clean_code = payload.code.strip().lower()

    if not clean_code:
        raise HTTPException(400, "Role code is required.")
    if not re.match(r"^[a-zA-Z0-9_-]+$", clean_code):
        raise HTTPException(400, "Code can only contain letters, numbers, underscores (_), and hyphens (-). Special characters and spaces are not allowed.")

    existing_name = await db_fetchone("SELECT id FROM roles WHERE LOWER(name)=LOWER(%s)", (clean_name,))
    if existing_name:
        raise HTTPException(400, "Role name already exists.")

    existing_code = await db_fetchone("SELECT id FROM roles WHERE LOWER(code)=LOWER(%s)", (clean_code,))
    if existing_code:
        raise HTTPException(400, "Role code already exists.")

    role_id = await db_execute(
        "INSERT INTO roles (entity_id, name, code, description, is_active) VALUES (%s, %s, %s, %s, %s)", 
        (payload.entity_id, clean_name, clean_code, payload.description, payload.is_active)
    )
    return {"id": role_id, **payload.dict(), "code": clean_code}

@api_router.put("/admin/roles/{item_id}")
async def update_role(item_id: str, payload: RoleIn, _: dict = Depends(require_admin)):
    clean_name = payload.name.strip()
    clean_code = payload.code.strip().lower()

    if not clean_code:
        raise HTTPException(400, "Role code is required.")
    if not re.match(r"^[a-zA-Z0-9_-]+$", clean_code):
        raise HTTPException(400, "Code can only contain letters, numbers, underscores (_), and hyphens (-). Special characters and spaces are not allowed.")

    existing_name = await db_fetchone("SELECT id FROM roles WHERE LOWER(name)=LOWER(%s) AND id!=%s", (clean_name, item_id))
    if existing_name:
        raise HTTPException(400, "Role name already exists.")

    existing_code = await db_fetchone("SELECT id FROM roles WHERE LOWER(code)=LOWER(%s) AND id!=%s", (clean_code, item_id))
    if existing_code:
        raise HTTPException(400, "Role code already exists.")

    await db_execute(
        "UPDATE roles SET entity_id=%s, name=%s, code=%s, description=%s, is_active=%s WHERE id=%s", 
        (payload.entity_id, clean_name, clean_code, payload.description, payload.is_active, item_id)
    )
    return {"ok": True}

@api_router.delete("/admin/roles/{item_id}")
async def delete_role(item_id: str, _: dict = Depends(require_admin)):
    acct = await db_fetchone("SELECT account_id FROM accounts WHERE role_id=%s LIMIT 1", (item_id,))
    if acct: raise HTTPException(400, "Cannot delete Role because it is assigned to an account.")
    await db_execute("DELETE FROM role_feature_permissions WHERE role_id=%s", (item_id,))
    await db_execute("DELETE FROM roles WHERE id=%s", (item_id,))
    return {"ok": True}

@api_router.get("/admin/stats")
async def get_admin_stats(account: dict = Depends(require_admin)):
    pass # placeholder for existing stats logic (already below)

import secrets
import string

def generate_temp_password(length=8):
    alphabet = string.ascii_letters + string.digits
    return "".join(secrets.choice(alphabet) for i in range(length))

def mock_send_email(to_email, subject, body):
    # Log to console
    print(f"--- EMAIL SENT ---\nTo: {to_email}\nSubject: {subject}\nBody:\n{body}\n------------------")

@api_router.post("/admin/employees")
async def add_employee(payload: EmployeeIn, account: dict = Depends(require_admin)):
    # Check if email exists
    existing = await db_fetchone("SELECT account_id FROM accounts WHERE email=%s", (payload.email,))
    if existing:
        raise HTTPException(status_code=400, detail="Email already exists.")

    temp_password = generate_temp_password()
    password_hash = hash_password(temp_password)
    
    # Create employee details
    role = await db_fetchone("SELECT name FROM roles WHERE id=%s", (payload.role_id,))
    if not role:
        raise HTTPException(status_code=400, detail="Invalid role ID.")
        
    prefix = "NT" + role["name"][:2].upper() + "#"
    max_id_row = await db_fetchone(
        "SELECT employee_id_number FROM employees WHERE employee_id_number LIKE %s ORDER BY LENGTH(employee_id_number) DESC, employee_id_number DESC LIMIT 1",
        (prefix + "%",)
    )
    if max_id_row:
        try:
            last_num = int(max_id_row["employee_id_number"].split("#")[1])
            next_num = last_num + 1
        except Exception:
            next_num = 1
    else:
        next_num = 1
    emp_id_num = f"{prefix}{next_num}"

    full_name = f"{payload.first_name} {payload.last_name}"
    full_address = f"{payload.address_line_1}, {payload.address_line_2 or ''}, {payload.city}, {payload.state} - {payload.pincode}"
    
    onboard = payload.onboard_date if payload.onboard_date else None
    end = payload.end_date if payload.end_date else None
    
    emp_id = await db_execute(
        """INSERT INTO employees 
           (employee_id_number, name, address, email, contact_number, first_name, last_name, address_line_1, address_line_2, city, state, pincode, emergency_contact_name, emergency_contact_number, id_proof_url, onboard_date, end_date) 
           VALUES (%s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s)""",
        (emp_id_num, full_name, full_address, payload.email, payload.contact_number, payload.first_name, payload.last_name, payload.address_line_1, payload.address_line_2, payload.city, payload.state, payload.pincode, payload.emergency_contact_name, payload.emergency_contact_number, payload.id_proof_url, onboard, end)
    )

    # Create account
    account_id = await db_execute(
        "INSERT INTO accounts (employee_id, email, password_hash, role_id, raw_password) VALUES (%s, %s, %s, %s, %s)",
        (emp_id, payload.email, password_hash, payload.role_id, temp_password)
    )

    # Insert associations
    for assoc_id in payload.association_ids:
        await db_execute(
            "INSERT INTO admin_associations (admin_id, association_id) VALUES (%s, %s)",
            (account_id, assoc_id)
        )

    # Log email
    subject = "Welcome to Nestora! Your Login Details"
    body = f"Hello {payload.first_name},\n\nYou have been onboarded. Your temporary password is: {temp_password}\nPlease login and change it.\n\nThanks,\nNestora Admin"
    await db_execute("INSERT INTO email_log (to_email, subject, body) VALUES (%s, %s, %s)", (payload.email, subject, body))
    mock_send_email(payload.email, subject, body)

    return {"ok": True, "account_id": account_id, "temp_password": temp_password}

@api_router.get("/admin/employees")
async def list_employees(account: dict = Depends(require_admin)):
    # Fetch all employees (users who are not just homeowners/board members, we'll just join on accounts with role != homeowner)
    # Actually, we can just fetch all accounts with their user details
    rows = await db_fetchall(
        """SELECT a.account_id, a.email, a.raw_password, a.role_id, r.name as role_name, 
                  e.employee_id_number, e.first_name, e.last_name, e.contact_number, e.city, e.state, e.address, e.address_line_1, e.address_line_2, e.pincode, e.onboard_date, e.end_date, e.emergency_contact_name, e.emergency_contact_number, e.id_proof_url
           FROM accounts a
           LEFT JOIN roles r ON a.role_id = r.id
           LEFT JOIN employees e ON a.employee_id = e.employee_id
           WHERE r.name NOT IN ('Homeowner', 'Board member', 'Committee member', 'Tenant', 'Security')
           ORDER BY a.created_at DESC"""
    )
    
    # Attach associations and serialize dates
    for row in rows:
        assocs = await db_fetchall(
            """SELECT a.id, a.name 
               FROM admin_associations aa 
               JOIN associations a ON aa.association_id = a.id 
               WHERE aa.admin_id=%s""", 
            (row["account_id"],)
        )
        row["associations"] = assocs
        for k, v in list(row.items()):
            if isinstance(v, (date, datetime)):
                row[k] = str(v)
        
    return {"ok": True, "employees": rows}


@api_router.post("/admin/users")
async def admin_create_user(payload: AdminCreateUserIn, _: dict = Depends(require_admin)):
    code = None
    for _try in range(6):
        candidate = f"NST-{generate_code(6)}"
        exists = await db_fetchone("SELECT id FROM user_codes WHERE login_code=%s", (candidate,))
        if not exists:
            code = candidate
            break
    if code is None:
        raise HTTPException(status_code=500, detail="Could not generate a unique code.")
    code_id = await db_execute("INSERT INTO user_codes (login_code) VALUES (%s)", (code,))
    
    full_name = f"{payload.first_name} {payload.last_name}".strip()
    
    await db_execute(
        "INSERT INTO user_details (code_id, name, first_name, last_name, email, contact_number, role_id, association_id, address) VALUES (%s,%s,%s,%s,%s,%s,%s,%s,%s)",
        (code_id, full_name, payload.first_name, payload.last_name, payload.email.lower(), payload.contact_number, payload.role_id, payload.association_id, ""),
    )
    
    # Optional: send email here if we want to mimic admin_create_member
    
    return {"ok": True, "message": "User created successfully", "activation_code": code}

@api_router.get("/admin/users")
async def get_all_users(_: dict = Depends(require_admin)):
    rows = await db_fetchall("""
        SELECT 
            u.user_id, u.first_name, u.last_name, u.contact_number, u.email,
            a.account_id, uc.login_code as activation_code,
            COALESCE(r.name, 'Homeowner') as role_name,
            un.id as unit_id,
            b.id as block_id,
            COALESCE(assoc.id, dir_assoc.id) as association_id,
            COALESCE(assoc.name, dir_assoc.name) as association_name,
            b.name as block_name,
            un.unit_number,
            COALESCE(assoc.address_line_1, dir_assoc.address_line_1) as assoc_addr1,
            COALESCE(assoc.address_line_2, dir_assoc.address_line_2) as assoc_addr2,
            COALESCE(assoc.city, dir_assoc.city) as assoc_city,
            COALESCE(assoc.state, dir_assoc.state) as assoc_state,
            COALESCE(assoc.pincode, dir_assoc.pincode) as assoc_pincode
        FROM user_details u
        LEFT JOIN accounts a ON u.user_id = a.user_id
        LEFT JOIN roles r ON COALESCE(a.role_id, u.role_id) = r.id
        LEFT JOIN user_codes uc ON u.code_id = uc.id
        LEFT JOIN units un ON u.unit_id = un.id
        LEFT JOIN blocks b ON un.block_id = b.id
        LEFT JOIN associations assoc ON b.association_id = assoc.id
        LEFT JOIN associations dir_assoc ON u.association_id = dir_assoc.id
        WHERE b.association_id IS NOT NULL OR u.association_id IS NOT NULL
        ORDER BY u.created_at DESC
    """)
    return {"ok": True, "users": rows}

@api_router.post("/admin/users/{user_id}/send-code")
async def send_activation_code(user_id: str, _: dict = Depends(require_admin)):
    user = await db_fetchone(
        "SELECT u.name, u.first_name, u.last_name, u.email, uc.login_code FROM user_details u LEFT JOIN user_codes uc ON u.code_id = uc.id WHERE u.user_id=%s", 
        (user_id,)
    )
    if not user:
        raise HTTPException(status_code=404, detail="User not found.")
    if not user["login_code"]:
        raise HTTPException(status_code=400, detail="User does not have an activation code.")
    if not user["email"]:
        raise HTTPException(status_code=400, detail="User does not have an email.")

    name = user["first_name"] or user["name"] or "User"
    await log_email(
        to_email=user["email"],
        subject="Nestora — Your access code",
        body=(
            f"Hi {name},\n\n"
            f"Your Nestora access code is:\n\n    {user['login_code']}\n\n"
            f"Use this code at the Nestora portal to set up or access your account.\n\n— Nestora Community Team"
        ),
    )
    return {"ok": True, "message": "Activation code sent successfully."}

class AdminUserUpdateIn(BaseModel):
    email: Optional[EmailStr] = None
    contact_number: Optional[str] = None
    association_id: Optional[str] = None
    block_name: Optional[str] = None
    unit_number: Optional[str] = None
    role_name: Optional[str] = None

@api_router.put("/admin/users/{user_id}")
async def update_user(user_id: str, payload: AdminUserUpdateIn, _: dict = Depends(require_admin)):
    user = await db_fetchone("SELECT * FROM user_details WHERE user_id=%s", (user_id,))
    if not user:
        raise HTTPException(status_code=404, detail="User not found")

    new_unit_id = user["unit_id"]
    if payload.association_id and payload.block_name and payload.unit_number:
        block = await db_fetchone("SELECT id FROM blocks WHERE association_id=%s AND name=%s", (payload.association_id, payload.block_name))
        if not block:
            raise HTTPException(status_code=400, detail=f"Block '{payload.block_name}' not found in the selected association.")
        
        unit = await db_fetchone("SELECT id FROM units WHERE block_id=%s AND unit_number=%s", (block["id"], payload.unit_number))
        if not unit:
            raise HTTPException(status_code=400, detail=f"Unit '{payload.unit_number}' not found in block '{payload.block_name}'.")
        
        new_unit_id = unit["id"]
        
    email = payload.email if payload.email else user["email"]
    contact = payload.contact_number if payload.contact_number else user["contact_number"]
    
    await db_execute(
        "UPDATE user_details SET email=%s, contact_number=%s, unit_id=%s WHERE user_id=%s",
        (email, contact, new_unit_id, user_id)
    )
    
    await db_execute(
        "UPDATE accounts SET email=%s WHERE user_id=%s",
        (email, user_id)
    )
    
    if payload.role_name:
        role = await db_fetchone("SELECT id FROM roles WHERE name=%s", (payload.role_name,))
        if role:
            await db_execute("UPDATE user_details SET role_id=%s WHERE user_id=%s", (role["id"], user_id))
            await db_execute("UPDATE accounts SET role_id=%s WHERE user_id=%s", (role["id"], user_id))
    
    return {"ok": True, "message": "Role assigned successfully"}



# ------------------------------------------------------------------------------
# Vendors API
# ------------------------------------------------------------------------------

class VendorIn(BaseModel):
    name: str
    service_type: Optional[str] = None
    contact_person: Optional[str] = None
    email: Optional[str] = None
    contact_number: Optional[str] = None
    address: Optional[str] = None
    gst_number: Optional[str] = None
    licence_url: Optional[str] = None
    certificate_url: Optional[str] = None
    status: Optional[str] = "Active"
    
    vendor_code: Optional[str] = None
    vendor_category: Optional[str] = None
    business_name: Optional[str] = None
    mobile_number: Optional[str] = None
    alternate_mobile_number: Optional[str] = None
    website: Optional[str] = None
    whatsapp_number: Optional[str] = None
    address_line_1: Optional[str] = None
    address_line_2: Optional[str] = None
    city: Optional[str] = None
    state: Optional[str] = None
    country: Optional[str] = None
    zip_code: Optional[str] = None
    pan_number: Optional[str] = None
    registration_number: Optional[str] = None
    license_number: Optional[str] = None
    trade_license_number: Optional[str] = None
    years_of_experience: Optional[int] = None
    available_days: Optional[str] = None
    working_hours: Optional[str] = None
    emergency_service: Optional[bool] = None
    support_24_7: Optional[bool] = None
    contract_start_date: Optional[str] = None
    contract_end_date: Optional[str] = None
    contract_value: Optional[float] = None
    payment_terms: Optional[str] = None
    contract_document_url: Optional[str] = None
    renewal_reminder: Optional[bool] = None
    bank_account_name: Optional[str] = None
    bank_name: Optional[str] = None
    bank_account_number: Optional[str] = None
    bank_ifsc_code: Optional[str] = None
    bank_upi_id: Optional[str] = None
    gst_certificate_url: Optional[str] = None
    pan_card_url: Optional[str] = None
    business_license_url: Optional[str] = None
    insurance_certificate_url: Optional[str] = None
    agreement_copy_url: Optional[str] = None
    identity_proof_url: Optional[str] = None
    address_proof_url: Optional[str] = None
    other_documents_url: Optional[str] = None
    vendor_rating: Optional[float] = None
    preferred_vendor: Optional[bool] = None
    verified_vendor: Optional[bool] = None
    remarks: Optional[str] = None
    association_name: Optional[str] = None
    assigned_blocks: Optional[str] = None
    assigned_services: Optional[str] = None
    assigned_manager: Optional[str] = None
    contact_details_json: Optional[str] = None
    services_offered_json: Optional[str] = None

@api_router.get("/admin/vendors")
async def get_vendors(_: dict = Depends(require_admin)):
    rows = await db_fetchall("SELECT * FROM vendors ORDER BY created_at DESC")
    return {"ok": True, "vendors": rows}

@api_router.post("/admin/vendors")
async def create_vendor(payload: VendorIn, _: dict = Depends(require_admin)):
    cols = []
    vals = []
    
    # Extract fields from payload that are not None
    payload_dict = payload.dict()
    if not payload_dict.get('service_type'):
        payload_dict['service_type'] = 'General'
        
    for k, v in payload_dict.items():
        if v is not None:
            cols.append(k)
            vals.append(v)
            
    query = f"INSERT INTO vendors ({', '.join(cols)}) VALUES ({', '.join(['%s']*len(cols))})"
    
    v_id = await db_execute(query, tuple(vals))
    return {"ok": True, "message": "Vendor created", "id": v_id}

@api_router.put("/admin/vendors/{id}")
async def update_vendor(id: str, payload: VendorIn, _: dict = Depends(require_admin)):
    cols = []
    vals = []
    
    payload_dict = payload.dict(exclude_unset=True) # Only update fields that are set
    for k, v in payload_dict.items():
        cols.append(f"{k}=%s")
        vals.append(v)
        
    if not cols:
        return {"ok": True, "message": "Nothing to update"}
        
    vals.append(id)
    query = f"UPDATE vendors SET {', '.join(cols)} WHERE id=%s"
    
    await db_execute(query, tuple(vals))
    return {"ok": True, "message": "Vendor updated"}

@api_router.delete("/admin/vendors/{id}")
async def delete_vendor(id: str, _: dict = Depends(require_admin)):
    await db_execute("DELETE FROM vendors WHERE id=%s", (id,))
    return {"ok": True, "message": "Vendor deleted"}

# ------------------------------------------------------------------------------
# Documents API
# ------------------------------------------------------------------------------

class DocumentIn(BaseModel):
    association_id: str
    title: str
    document_number: Optional[str] = None
    document_type: Optional[str] = None
    category: Optional[str] = None
    
    file_name: Optional[str] = None
    file_url: Optional[str] = None
    file_type: Optional[str] = None
    file_size_kb: Optional[int] = None
    
    department: Optional[str] = None
    related_module: Optional[str] = None
    
    visibility: Optional[str] = None
    allow_download: Optional[bool] = True
    
    issue_date: Optional[str] = None
    expiry_date: Optional[str] = None
    reminder_before_expiry_days: Optional[int] = None
    
    status: Optional[str] = "Active"
    keywords: Optional[str] = None
    remarks: Optional[str] = None

@api_router.get("/admin/documents")
async def get_documents(assoc_id: Optional[str] = None, account: dict = Depends(get_current_account)):
    # Non-admins can only see documents if they have 'Documents' feature permission
    # The UI handles the feature lock, but we should also enforce it
    
    query = """
        SELECT d.*, a.name as association_name
        FROM documents d
        JOIN associations a ON d.association_id = a.id
    """
    params = []
    
    if account["role"] not in ["Super admin", "Admin"]:
        # If not admin, they can only fetch for their own association if they are a board member, etc.
        # but let's just let the WHERE clause handle visibility first.
        query += " WHERE 1=1"
    else:
        query += " WHERE 1=1"
    if assoc_id:
        query += " AND d.association_id = %s"
        params.append(assoc_id)
        
    query += " ORDER BY d.created_at DESC"
    
    rows = await db_fetchall(query, tuple(params))
    
    # Filter in Python for simplicity based on visibility
    filtered_rows = []
    for r in rows:
        if account["role"] in ["Super admin", "Admin"]:
            filtered_rows.append(r)
            continue
            
        visibility = r.get("visibility", "")
        if "Public" in visibility:
            filtered_rows.append(r)
            continue
            
        if account["role"] in ["Homeowner", "Tenant"] and "Residents" in visibility:
            filtered_rows.append(r)
            continue
            
        if account["role"] == "Committee member" and ("Committee" in visibility or "Residents" in visibility):
            filtered_rows.append(r)
            continue
            
        if account["role"] == "Board member" and ("Board Member" in visibility or "Committee" in visibility or "Residents" in visibility):
            filtered_rows.append(r)
            continue
            
    rows = filtered_rows
    for r in rows:
        if r.get('issue_date'): r['issue_date'] = str(r['issue_date'])
        if r.get('expiry_date'): r['expiry_date'] = str(r['expiry_date'])
        if r.get('created_at'): r['created_at'] = str(r['created_at'])
    return {"ok": True, "data": rows}

@api_router.post("/admin/documents")
async def create_document(payload: DocumentIn, account: dict = Depends(require_permission("Documents", "can_create"))):
    issue = payload.issue_date if payload.issue_date else None
    expiry = payload.expiry_date if payload.expiry_date else None
    await db_execute(
        """
        INSERT INTO documents (
            association_id, title, document_number, document_type, category,
            file_name, file_url, file_type, file_size_kb,
            department, related_module, visibility, allow_download,
            issue_date, expiry_date, reminder_before_expiry_days,
            status, keywords, remarks
        ) VALUES (
            %s, %s, %s, %s, %s,
            %s, %s, %s, %s,
            %s, %s, %s, %s,
            %s, %s, %s,
            %s, %s, %s
        )
        """,
        (
            payload.association_id, payload.title, payload.document_number, payload.document_type, payload.category,
            payload.file_name, payload.file_url, payload.file_type, payload.file_size_kb,
            payload.department, payload.related_module, payload.visibility, payload.allow_download,
            issue, expiry, payload.reminder_before_expiry_days,
            payload.status, payload.keywords, payload.remarks
        )
    )
    return {"ok": True}

@api_router.delete("/admin/documents/{doc_id}")
async def delete_document(doc_id: str, account: dict = Depends(require_permission("Documents", "can_delete"))):
    await db_execute("DELETE FROM documents WHERE id=%s", (doc_id,))
    return {"ok": True}

class DocumentUpdateIn(BaseModel):
    title: str
    document_number: Optional[str] = None
    document_type: Optional[str] = None
    category: Optional[str] = None
    department: Optional[str] = None
    related_module: Optional[str] = None
    visibility: Optional[str] = None
    allow_download: Optional[bool] = True
    issue_date: Optional[str] = None
    expiry_date: Optional[str] = None
    reminder_before_expiry_days: Optional[int] = None
    status: Optional[str] = "Active"
    keywords: Optional[str] = None
    remarks: Optional[str] = None

@api_router.put("/admin/documents/{doc_id}")
async def update_document(doc_id: str, payload: DocumentUpdateIn, account: dict = Depends(require_permission("Documents", "can_update"))):
    issue = payload.issue_date if payload.issue_date else None
    expiry = payload.expiry_date if payload.expiry_date else None
    await db_execute(
        """
        UPDATE documents SET
            title=%s, document_number=%s, document_type=%s, category=%s,
            department=%s, related_module=%s, visibility=%s, allow_download=%s,
            issue_date=%s, expiry_date=%s, reminder_before_expiry_days=%s,
            status=%s, keywords=%s, remarks=%s
        WHERE id=%s
        """,
        (
            payload.title, payload.document_number, payload.document_type, payload.category,
            payload.department, payload.related_module, payload.visibility, payload.allow_download,
            issue, expiry, payload.reminder_before_expiry_days,
            payload.status, payload.keywords, payload.remarks,
            doc_id
        )
    )
    return {"ok": True}



class EmployeeUpdateIn(BaseModel):
    role_id: str
    association_ids: List[str]
    onboard_date: Optional[str] = None
    end_date: Optional[str] = None

@api_router.put("/admin/employees/{account_id}")
async def update_employee(account_id: str, payload: EmployeeUpdateIn, account: dict = Depends(require_admin)):
    # Update role in accounts
    await db_execute("UPDATE accounts SET role_id=%s WHERE account_id=%s", (payload.role_id, account_id))
    
    onboard = payload.onboard_date if payload.onboard_date else None
    end = payload.end_date if payload.end_date else None
    
    # Update dates in employees
    await db_execute(
        "UPDATE employees SET onboard_date=%s, end_date=%s WHERE employee_id=(SELECT employee_id FROM accounts WHERE account_id=%s)",
        (onboard, end, account_id)
    )
    
    # Update associations (clear and re-insert)
    await db_execute("DELETE FROM admin_associations WHERE admin_id=%s", (account_id,))
    for assoc_id in payload.association_ids:
        await db_execute(
            "INSERT INTO admin_associations (admin_id, association_id) VALUES (%s, %s)",
            (account_id, assoc_id)
        )
    return {"ok": True}

@api_router.get("/admin/features")
async def get_features(_: dict = Depends(require_admin)):
    return await db_fetchall("""
        SELECT f.*, p.name AS parent_name 
        FROM features f 
        LEFT JOIN features p ON f.parent_id = p.id 
        ORDER BY f.order_index ASC, f.created_at ASC
    """)

@api_router.post("/admin/features")
async def create_feature(payload: FeatureIn, _: dict = Depends(require_admin)):
    clean_name = payload.name.strip()
    clean_code = payload.code.strip().lower()
    
    if not clean_code:
        raise HTTPException(400, "Feature code is required.")
    if not re.match(r"^[a-zA-Z0-9_-]+$", clean_code):
        raise HTTPException(400, "Code can only contain letters, numbers, underscores (_), and hyphens (-). Special characters and spaces are not allowed.")

    existing_name = await db_fetchone("SELECT id FROM features WHERE LOWER(name)=LOWER(%s)", (clean_name,))
    if existing_name:
        raise HTTPException(400, "Feature name already exists.")

    existing_code = await db_fetchone("SELECT id FROM features WHERE LOWER(code)=LOWER(%s)", (clean_code,))
    if existing_code:
        raise HTTPException(400, "Feature code already exists.")

    feature_id = await db_execute(
        "INSERT INTO features (name, code, description, parent_id, icon, url, is_active) VALUES (%s,%s,%s,%s,%s,%s,%s)",
        (clean_name, clean_code, payload.description, payload.parent_id, payload.icon, payload.url, payload.is_active)
    )
    return {"id": feature_id, **payload.dict(), "code": clean_code}

@api_router.put("/admin/features/{item_id}")
async def update_feature(item_id: str, payload: FeatureIn, _: dict = Depends(require_admin)):
    clean_name = payload.name.strip()
    clean_code = payload.code.strip().lower()
    
    if not clean_code:
        raise HTTPException(400, "Feature code is required.")
    if not re.match(r"^[a-zA-Z0-9_-]+$", clean_code):
        raise HTTPException(400, "Code can only contain letters, numbers, underscores (_), and hyphens (-). Special characters and spaces are not allowed.")

    existing_name = await db_fetchone("SELECT id FROM features WHERE LOWER(name)=LOWER(%s) AND id!=%s", (clean_name, item_id))
    if existing_name:
        raise HTTPException(400, "Feature name already exists.")

    existing_code = await db_fetchone("SELECT id FROM features WHERE LOWER(code)=LOWER(%s) AND id!=%s", (clean_code, item_id))
    if existing_code:
        raise HTTPException(400, "Feature code already exists.")

    await db_execute(
        "UPDATE features SET name=%s, code=%s, description=%s, parent_id=%s, icon=%s, url=%s, is_active=%s WHERE id=%s",
        (clean_name, clean_code, payload.description, payload.parent_id, payload.icon, payload.url, payload.is_active, item_id)
    )
    if not payload.is_active:
        await db_execute("DELETE FROM role_feature_permissions WHERE feature_id=%s", (item_id,))
    return {"ok": True}

@api_router.delete("/admin/features/{item_id}")
async def delete_feature(item_id: str, _: dict = Depends(require_admin)):
    await db_execute("DELETE FROM role_feature_permissions WHERE feature_id=%s", (item_id,))
    await db_execute("DELETE FROM features WHERE id=%s", (item_id,))
    return {"ok": True}

@api_router.get("/admin/features/bulk/template")
async def download_feature_template(_: dict = Depends(require_admin)):
    import pandas as pd
    import io
    from fastapi.responses import StreamingResponse
    
    columns = ["Feature Name", "Description", "Parent Feature", "Feature Number", "Feature Icon", "URL"]
    df = pd.DataFrame(columns=columns)
    
    output = io.BytesIO()
    with pd.ExcelWriter(output, engine='openpyxl') as writer:
        df.to_excel(writer, index=False, sheet_name='Features')
    output.seek(0)
    
    return StreamingResponse(
        output,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": "attachment; filename=Feature_Bulk_Upload_Template.xlsx"}
    )

@api_router.post("/admin/features/bulk")
async def bulk_upload_features(file: UploadFile = File(...), _: dict = Depends(require_admin)):
    import pandas as pd
    import io
    import math
    from fastapi import HTTPException
    
    try:
        contents = await file.read()
        df = pd.read_excel(io.BytesIO(contents))
    except Exception as e:
        raise HTTPException(status_code=400, detail=f"Failed to read Excel file: {str(e)}")
        
    required_cols = ["Feature Name", "Description", "Parent Feature", "Feature Number", "Feature Icon", "URL"]
    for col in required_cols:
        if col not in df.columns:
            raise HTTPException(status_code=400, detail=f"Missing required column: {col}")
            
    success_count = 0
    errors = []
    
    # Pre-fetch all existing features to resolve parent IDs
    existing_features = await db_fetchall("SELECT id, name FROM features")
    feature_map = {f['name'].lower().strip(): f['id'] for f in existing_features}
    
    for index, row in df.iterrows():
        try:
            name = str(row.get("Feature Name", "")).strip()
            if not name or name == "nan":
                errors.append(f"Row {index + 2}: Feature Name is required")
                continue
                
            desc = str(row.get("Description", "")).strip()
            if desc == "nan": desc = None
            
            parent_name = str(row.get("Parent Feature", "")).strip()
            parent_id = None
            if parent_name and parent_name != "nan":
                parent_id = feature_map.get(parent_name.lower())
                if not parent_id:
                    errors.append(f"Row {index + 2}: Parent Feature '{parent_name}' not found in database")
                    continue
                    
            feat_num = str(row.get("Feature Number", "")).strip()
            if feat_num == "nan": feat_num = None
            
            icon_url = str(row.get("Feature Icon", "")).strip()
            if icon_url == "nan": icon_url = None
            
            url = str(row.get("URL", "")).strip()
            if url == "nan": url = None
            
            # Insert logic
            await db_execute(
                """INSERT INTO features (name, description, parent_id, icon, url, feature_number, is_active)
                   VALUES (%s, %s, %s, %s, %s, %s, %s)""",
                (name, desc, parent_id, icon_url, url, feat_num, True)
            )
            success_count += 1
            
            # Update feature map just in case a subsequent row refers to this newly inserted feature
            newly_inserted = await db_fetchone("SELECT id FROM features WHERE name=%s ORDER BY created_at DESC LIMIT 1", (name,))
            if newly_inserted:
                feature_map[name.lower()] = newly_inserted['id']
                
        except Exception as e:
            errors.append(f"Row {index + 2}: {str(e)}")
            
    return {
        "ok": True,
        "message": f"Successfully imported {success_count} features.",
        "errors": errors
    }

@api_router.post("/admin/permissions")
async def create_permission(payload: RoleFeaturePermissionIn, _: dict = Depends(require_admin)):
    existing = await db_fetchone("SELECT id FROM role_feature_permissions WHERE role_id=%s AND feature_id=%s", (payload.role_id, payload.feature_id))
    if existing:
        raise HTTPException(status_code=400, detail="Permission configuration for this Role and Feature already exists.")
    perm_id = await db_execute(
        """INSERT INTO role_feature_permissions
           (role_id, feature_id, can_create, can_view, can_update, can_delete, sidebar_order)
           VALUES (%s, %s, %s, %s, %s, %s, %s)""",
        (payload.role_id, payload.feature_id, int(payload.can_create), int(payload.can_view), int(payload.can_update), int(payload.can_delete), payload.sidebar_order)
    )
    return {"id": perm_id, **payload.dict()}

@api_router.put("/admin/permissions/{item_id}")
async def update_permission(item_id: str, payload: RoleFeaturePermissionIn, _: dict = Depends(require_admin)):
    existing = await db_fetchone("SELECT id FROM role_feature_permissions WHERE role_id=%s AND feature_id=%s AND id != %s", (payload.role_id, payload.feature_id, item_id))
    if existing:
        raise HTTPException(status_code=400, detail="Permission configuration for this Role and Feature already exists.")
    await db_execute(
        """UPDATE role_feature_permissions 
           SET role_id=%s, feature_id=%s, can_create=%s, can_view=%s, can_update=%s, can_delete=%s, sidebar_order=%s
           WHERE id=%s""",
        (payload.role_id, payload.feature_id, int(payload.can_create), int(payload.can_view), int(payload.can_update), int(payload.can_delete), payload.sidebar_order, item_id)
    )
    return {"id": item_id, **payload.dict()}


@api_router.get("/admin/permissions")
async def get_permissions(_: dict = Depends(require_admin)):
    return await db_fetchall("""
        SELECT p.*, r.name as role_name, f.name as feature_name 
        FROM role_feature_permissions p
        LEFT JOIN roles r ON p.role_id = r.id
        LEFT JOIN features f ON p.feature_id = f.id
        ORDER BY p.created_at DESC
    """)

@api_router.get("/admin/roles-permissions-summary")
async def get_roles_permissions_summary(_: dict = Depends(require_admin)):
    return await db_fetchall("""
        SELECT 
            r.id, 
            r.name, 
            r.code, 
            r.description, 
            r.entity_id, 
            r.is_active, 
            r.created_at,
            e.name AS entity_name,
            (SELECT COUNT(DISTINCT feature_id) FROM role_feature_permissions WHERE role_id = r.id AND (can_create=1 OR can_view=1 OR can_update=1 OR can_delete=1)) AS configured_features_count,
            (SELECT COUNT(*) FROM accounts WHERE role_id = r.id) AS accounts_count,
            (SELECT COUNT(*) FROM features WHERE is_active = 1) AS total_features_count
        FROM roles r
        LEFT JOIN entities e ON r.entity_id = e.id
        ORDER BY r.created_at DESC
    """)

@api_router.get("/admin/roles/{role_id}/permissions-matrix")
async def get_role_permissions_matrix(role_id: str, _: dict = Depends(require_admin)):
    role = await db_fetchone("SELECT r.*, e.name as entity_name FROM roles r LEFT JOIN entities e ON r.entity_id = e.id WHERE r.id=%s", (role_id,))
    if not role:
        raise HTTPException(404, "Role not found")
    
    features_matrix = await db_fetchall("""
        SELECT 
            f.id AS feature_id,
            f.name AS feature_name,
            f.code AS feature_code,
            f.description AS feature_description,
            f.parent_id,
            p.name AS parent_name,
            f.url,
            f.order_index,
            f.icon,
            COALESCE(rfp.can_create, 0) AS can_create,
            COALESCE(rfp.can_view, 0) AS can_view,
            COALESCE(rfp.can_update, 0) AS can_update,
            COALESCE(rfp.can_delete, 0) AS can_delete,
            COALESCE(rfp.sidebar_order, 0) AS sidebar_order
        FROM features f
        LEFT JOIN features p ON f.parent_id = p.id
        LEFT JOIN role_feature_permissions rfp ON f.id = rfp.feature_id AND rfp.role_id = %s
        WHERE f.is_active = 1
        ORDER BY f.name ASC
    """, (role_id,))
    
    return {
        "role": role,
        "features": features_matrix
    }

@api_router.post("/admin/roles/{role_id}/permissions-bulk")
async def update_role_permissions_bulk(role_id: str, payload: BulkRolePermissionsIn, _: dict = Depends(require_admin)):
    role = await db_fetchone("SELECT id FROM roles WHERE id=%s", (role_id,))
    if not role:
        raise HTTPException(404, "Role not found")
    
    # Delete existing permissions for this role
    await db_execute("DELETE FROM role_feature_permissions WHERE role_id=%s", (role_id,))
    
    # Insert granted permissions
    for p in payload.permissions:
        c = 1 if p.can_create else 0
        v = 1 if p.can_view else 0
        u = 1 if p.can_update else 0
        d = 1 if p.can_delete else 0
        
        if c or v or u or d:
            await db_execute("""
                INSERT INTO role_feature_permissions
                    (role_id, feature_id, can_create, can_view, can_update, can_delete, sidebar_order)
                VALUES (%s, %s, %s, %s, %s, %s, %s)
            """, (role_id, p.feature_id, c, v, u, d, p.sidebar_order))
            
    return {"ok": True, "message": "Permissions updated successfully"}

@api_router.delete("/admin/permissions/{item_id}")
async def delete_permission(item_id: str, _: dict = Depends(require_admin)):
    await db_execute("DELETE FROM role_feature_permissions WHERE id=%s", (item_id,))
    return {"ok": True}

@api_router.get("/features")
async def get_user_features(account: dict = Depends(get_current_account)):
    if account["role"] == "Super admin":
        return await db_fetchall("SELECT * FROM features")
    query = """
        SELECT f.* FROM features f
        JOIN role_feature_permissions p ON f.id = p.feature_id
        JOIN roles r ON p.role_id = r.id
        WHERE r.name = %s AND p.can_view = 1
    """
    return await db_fetchall(query, (account["role"],))

# ---------- Routes: Admin ----------
@api_router.get("/admin/code-requests")
async def admin_list_code_requests(_: dict = Depends(require_admin)):
    rows = await db_fetchall(
        "SELECT id, name, email, contact_number, status, issued_code, created_at FROM code_requests ORDER BY created_at DESC"
    )
    for r in rows:
        if r.get("created_at"):
            r["created_at"] = r["created_at"].isoformat()
    return rows


@api_router.post("/admin/code-requests/approve")
async def admin_approve_request(payload: AdminIssueCodeIn, _: dict = Depends(require_admin)):
    req = await db_fetchone("SELECT * FROM code_requests WHERE id=%s", (payload.request_id,))
    if not req:
        raise HTTPException(status_code=404, detail="Request not found.")
    if req["status"] != "pending":
        raise HTTPException(status_code=400, detail="Request already processed.")

    # Generate a unique code
    for _ in range(6):
        code = f"NST-{generate_code(6)}"
        exists = await db_fetchone("SELECT id FROM user_codes WHERE login_code=%s", (code,))
        if not exists:
            break
    else:
        raise HTTPException(status_code=500, detail="Could not generate a unique code.")

    code_id = await db_execute("INSERT INTO user_codes (login_code) VALUES (%s)", (code,))
    await db_execute(
        "INSERT INTO user_details (code_id, name, address, email, contact_number) VALUES (%s,%s,%s,%s,%s)",
        (code_id, req["name"], "Address pending update", req["email"], req["contact_number"]),
    )
    await db_execute(
        "UPDATE code_requests SET status='approved', issued_code=%s WHERE id=%s", (code, payload.request_id)
    )
    await log_email(
        to_email=req["email"],
        subject="Nestora — Your access code is ready",
        body=(
            f"Hi {req['name']},\n\n"
            f"Welcome to Nestora! Your access code is:\n\n    {code}\n\n"
            f"Use this code at the Nestora portal to complete your account setup.\n\n— Nestora Community Team"
        ),
    )
    return {"ok": True, "issued_code": code}


@api_router.post("/admin/code-requests/reject")
async def admin_reject_request(payload: AdminIssueCodeIn, _: dict = Depends(require_admin)):
    req = await db_fetchone("SELECT * FROM code_requests WHERE id=%s", (payload.request_id,))
    if not req:
        raise HTTPException(status_code=404, detail="Request not found.")
    await db_execute("UPDATE code_requests SET status='rejected' WHERE id=%s", (payload.request_id,))
    return {"ok": True}


@api_router.get("/admin/codes")
async def admin_list_codes(_: dict = Depends(require_admin)):
    rows = await db_fetchall(
        """
        SELECT uc.id, uc.login_code, uc.status, uc.created_at,
               ud.name AS member_name, ud.email AS member_email, ud.address
        FROM user_codes uc
        LEFT JOIN user_details ud ON ud.code_id=uc.id
        ORDER BY uc.created_at DESC
        """
    )
    for r in rows:
        if r.get("created_at"):
            r["created_at"] = r["created_at"].isoformat()
    return rows


@api_router.post("/admin/codes/{code_id}/reset")
async def admin_reset_code(code_id: str, _: dict = Depends(require_admin)):
    """Reset a specific access code — unlinks any account created with it and marks it active again."""
    row = await db_fetchone("SELECT id, login_code, status FROM user_codes WHERE id=%s", (code_id,))
    if not row:
        raise HTTPException(status_code=404, detail="Code not found.")
    # Remove any member account attached via user_details for this code
    detail = await db_fetchone("SELECT user_id FROM user_details WHERE code_id=%s", (code_id,))
    if detail:
        await db_execute(
            "DELETE a FROM accounts a JOIN roles r ON a.role_id = r.id WHERE a.user_id=%s AND r.name='Homeowner'", (detail["user_id"],)
        )
    await db_execute("UPDATE user_codes SET status='active' WHERE id=%s", (code_id,))
    return {"ok": True, "login_code": row["login_code"], "status": "active"}


@api_router.post("/admin/codes/reset-demo")
async def admin_reset_demo_codes(_: dict = Depends(require_admin)):
    """Convenience: reset all NST-DEMO* codes to active and remove any member accounts using them."""
    demo_ids = await db_fetchall("SELECT id FROM user_codes WHERE login_code LIKE %s", ("NST-DEMO%",))
    reset_count = 0
    for r in demo_ids:
        det = await db_fetchone("SELECT user_id FROM user_details WHERE code_id=%s", (r["id"],))
        if det:
            await db_execute(
                "DELETE a FROM accounts a JOIN roles r ON a.role_id = r.id WHERE a.user_id=%s AND r.name='Homeowner'", (det["user_id"],)
            )
        await db_execute("UPDATE user_codes SET status='active' WHERE id=%s", (r["id"],))
        reset_count += 1
    return {"ok": True, "reset_count": reset_count}


@api_router.post("/admin/members")
async def admin_create_member(payload: AdminCreateMemberIn, _: dict = Depends(require_admin)):
    code = None
    for _try in range(6):
        candidate = f"NST-{generate_code(6)}"
        exists = await db_fetchone("SELECT id FROM user_codes WHERE login_code=%s", (candidate,))
        if not exists:
            code = candidate
            break
    if code is None:
        raise HTTPException(status_code=500, detail="Could not generate a unique code.")
    code_id = await db_execute("INSERT INTO user_codes (login_code) VALUES (%s)", (code,))
    await db_execute(
        "INSERT INTO user_details (code_id, name, address, email, contact_number) VALUES (%s,%s,%s,%s,%s)",
        (code_id, payload.name, payload.address, payload.email.lower(), payload.contact_number),
    )
    await log_email(
        to_email=payload.email,
        subject="Nestora — Your access code",
        body=f"Hi {payload.name},\n\nYour access code is: {code}\n\nUse it at the Nestora portal to complete your account.",
    )
    return {"ok": True, "issued_code": code}


@api_router.get("/admin/update-requests")
async def admin_list_update_requests(_: dict = Depends(require_admin)):
    rows = await db_fetchall(
        """
        SELECT ur.id, ur.status, ur.requested_name, ur.requested_address,
               ur.requested_email, ur.requested_contact, ur.note, ur.created_at,
               uc.login_code
        FROM update_requests ur LEFT JOIN user_codes uc ON uc.id=ur.code_id
        ORDER BY ur.created_at DESC
        """
    )
    for r in rows:
        if r.get("created_at"):
            r["created_at"] = r["created_at"].isoformat()
    return rows


@api_router.post("/admin/update-requests/{req_id}/resolve")
async def admin_resolve_update_request(req_id: str, _: dict = Depends(require_admin)):
    await db_execute("UPDATE update_requests SET status='resolved' WHERE id=%s", (req_id,))
    return {"ok": True}


@api_router.get("/admin/emails")
async def admin_list_emails(_: dict = Depends(require_admin)):
    rows = await db_fetchall("SELECT * FROM email_log ORDER BY created_at DESC LIMIT 100")
    for r in rows:
        if r.get("created_at"):
            r["created_at"] = r["created_at"].isoformat()
    return rows


@api_router.get("/admin/stats")
async def admin_get_dashboard_stats(_: dict = Depends(require_admin)):
    codes = await db_fetchone("SELECT COUNT(*) AS c FROM user_codes")
    used = await db_fetchone("SELECT COUNT(*) AS c FROM user_codes WHERE status='used'")
    reqs = await db_fetchone("SELECT COUNT(*) AS c FROM code_requests WHERE status='pending'")
    members = await db_fetchone("SELECT COUNT(*) AS c FROM accounts a JOIN roles r ON a.role_id = r.id WHERE r.name='Homeowner'")
    updates = await db_fetchone("SELECT COUNT(*) AS c FROM update_requests WHERE status='open'")
    return {
        "total_codes": codes["c"] if codes else 0,
        "used_codes": used["c"] if used else 0,
        "pending_requests": reqs["c"] if reqs else 0,
        "members": members["c"] if members else 0,
        "open_update_requests": updates["c"] if updates else 0,
    }


# ---------- Announcements ----------
def _serialize_announcement(r: dict) -> dict:
    if r.get("created_at"):
        r["created_at"] = r["created_at"].isoformat()
    if r.get("updated_at"):
        r["updated_at"] = r["updated_at"].isoformat()
    r["pinned"] = bool(r.get("pinned"))
    return r


@api_router.get("/announcements")
async def list_announcements(association_id: Optional[str] = None, account: dict = Depends(get_current_account)):
    is_homeowner = account["role"] == "Homeowner"
    is_board = account["role"] == "Board member"
    is_admin = account["role"] in ("Admin", "Super admin")
    
    where_clause = "WHERE 1=1"
    params = []
    
    if is_homeowner or is_board:
        res = await db_fetchall('''
            SELECT b.association_id 
            FROM accounts ac
            JOIN user_details ud ON ac.user_id = ud.user_id
            JOIN units u ON ud.unit_id = u.id
            JOIN blocks b ON u.block_id = b.id
            WHERE ac.account_id = %s
        ''', (account["account_id"],))
        if res:
            where_clause += " AND a.association_id = %s"
            params.append(res[0]['association_id'])
        else:
            return []
    elif is_admin:
        if association_id:
            res = await db_fetchall("SELECT association_id FROM admin_associations WHERE admin_id = %s AND association_id = %s", (account["account_id"], association_id))
            if res:
                where_clause += " AND a.association_id = %s"
                params.append(association_id)
            else:
                return []
        else:
            res = await db_fetchall("SELECT association_id FROM admin_associations WHERE admin_id = %s", (account["account_id"],))
            if res:
                assoc_ids = [r['association_id'] for r in res]
                placeholders = ','.join(['%s'] * len(assoc_ids))
                where_clause += f" AND a.association_id IN ({placeholders})"
                params.extend(assoc_ids)
            else:
                return []
    
    if is_homeowner:
        where_clause += " AND a.audience IN ('Homeowner', 'Homeowners')"
        
    query = f"""
        SELECT a.id, a.title, a.body, a.category, a.pinned, a.created_at, a.updated_at,
               a.audience, a.attachment_url,
               COALESCE(ud.name, ac.email) AS author_name
        FROM announcements a
        LEFT JOIN accounts ac ON ac.account_id=a.created_by
        LEFT JOIN user_details ud ON ud.user_id=ac.user_id
        {where_clause}
        ORDER BY a.pinned DESC, a.created_at DESC
    """
    
    rows = await db_fetchall(query, tuple(params))
    return [_serialize_announcement(r) for r in rows]


@api_router.post("/announcements")
async def create_announcement(payload: AnnouncementIn, account: dict = Depends(require_role(["Admin", "Super admin", "Board member"]))):
    assoc_id = payload.association_id
    if not assoc_id or assoc_id == "me":
        if account["role"] == "Board member":
            res = await db_fetchall('''
                SELECT b.association_id 
                FROM accounts ac
                JOIN user_details ud ON ac.user_id = ud.user_id
                JOIN units u ON ud.unit_id = u.id
                JOIN blocks b ON u.block_id = b.id
                WHERE ac.account_id = %s
            ''', (account["account_id"],))
            if not res:
                raise HTTPException(400, "Board member has no association assigned")
            assoc_id = res[0]['association_id']
        else:
            raise HTTPException(400, "association_id is required for admins")
            
    aid = await db_execute(
        "INSERT INTO announcements (association_id, title, body, category, pinned, audience, attachment_url, created_by) VALUES (%s,%s,%s,%s,%s,%s,%s,%s)",
        (assoc_id, payload.title, payload.body, payload.category, 1 if payload.pinned else 0, payload.audience, payload.attachment_url, account["account_id"]),
    )
    return {"ok": True, "id": aid}

@api_router.post("/announcements/{aid}/like")
async def toggle_announcement_like(aid: str, account: dict = Depends(get_current_account)):
    row = await db_fetchone("SELECT id FROM announcements WHERE id=%s", (aid,))
    if not row:
        raise HTTPException(status_code=404, detail="Announcement not found")
        
    existing = await db_fetchone("SELECT id FROM announcement_likes WHERE announcement_id=%s AND account_id=%s", (aid, account["account_id"]))
    if existing:
        await db_execute("DELETE FROM announcement_likes WHERE id=%s", (existing["id"],))
        return {"ok": True, "liked": False}
    else:
        await db_execute("INSERT INTO announcement_likes (announcement_id, account_id) VALUES (%s, %s)", (aid, account["account_id"]))
        return {"ok": True, "liked": True}

@api_router.post("/announcements/{aid}/comment")
async def add_announcement_comment(aid: str, payload: CommentIn, account: dict = Depends(get_current_account)):
    row = await db_fetchone("SELECT id FROM announcements WHERE id=%s", (aid,))
    if not row:
        raise HTTPException(status_code=404, detail="Announcement not found")
        
    cid = await db_execute("INSERT INTO announcement_comments (announcement_id, account_id, comment) VALUES (%s, %s, %s)", (aid, account["account_id"], payload.comment))
    return {"ok": True, "id": cid}

@api_router.get("/announcements/{aid}/comments")
async def get_announcement_comments(aid: str, account: dict = Depends(get_current_account)):
    rows = await db_fetchall('''
        SELECT c.id, c.comment, c.created_at, COALESCE(ud.name, ac.email) as author_name
        FROM announcement_comments c
        JOIN accounts ac ON c.account_id = ac.account_id
        LEFT JOIN user_details ud ON ac.user_id = ud.user_id
        WHERE c.announcement_id = %s
        ORDER BY c.created_at ASC
    ''', (aid,))
    for r in rows:
        if r.get("created_at"):
            r["created_at"] = r["created_at"].isoformat()
    return rows

@api_router.put("/admin/announcements/{aid}")
async def admin_update_announcement(aid: str, payload: AnnouncementIn, account: dict = Depends(require_role(["Admin", "Super admin", "Board member"]))):
    row = await db_fetchone("SELECT id, created_at FROM announcements WHERE id=%s", (aid,))
    if not row:
        raise HTTPException(status_code=404, detail="Announcement not found.")
    
    if account["role"] == "Board member":
        if row["created_at"]:
            from datetime import datetime, timedelta
            if datetime.utcnow() - row["created_at"] > timedelta(hours=2):
                raise HTTPException(status_code=403, detail="Board members cannot edit announcements after 2 hours.")

    await db_execute(
        "UPDATE announcements SET title=%s, body=%s, category=%s, pinned=%s, audience=%s, attachment_url=%s WHERE id=%s",
        (payload.title, payload.body, payload.category, 1 if payload.pinned else 0, payload.audience, payload.attachment_url, aid),
    )
    return {"ok": True}

@api_router.delete("/admin/announcements/{aid}")
async def admin_delete_announcement(aid: str, account: dict = Depends(require_role(["Admin", "Super admin", "Board member"]))):
    row = await db_fetchone("SELECT id, created_at FROM announcements WHERE id=%s", (aid,))
    if not row:
        raise HTTPException(status_code=404, detail="Announcement not found.")
    
    if account["role"] == "Board member":
        if row["created_at"]:
            from datetime import datetime, timedelta
            if datetime.utcnow() - row["created_at"] > timedelta(hours=2):
                raise HTTPException(status_code=403, detail="Board members cannot delete announcements after 2 hours.")

    await db_execute("DELETE FROM announcements WHERE id=%s", (aid,))
    return {"ok": True}


# ---------- Events ----------
def _serialize_event(r: dict) -> dict:
    for k in ("starts_at", "ends_at", "registration_deadline", "created_at", "updated_at"):
        if r.get(k):
            r[k] = r[k].isoformat()
    if r.get("fee_amount") is not None:
        r["fee_amount"] = float(r["fee_amount"])
    return r


async def _attach_rsvp_data(events_list: List[dict], account_id: str) -> List[dict]:
    """Enrich event dicts with rsvp_counts, my_rsvp_status, attendees (top 6 'going')."""
    if not events_list:
        return events_list
    event_ids = [e["id"] for e in events_list]
    placeholders = ",".join(["%s"] * len(event_ids))

    # Counts by status
    counts_rows = await db_fetchall(
        f"SELECT event_id, status, COUNT(*) AS c FROM event_rsvps WHERE event_id IN ({placeholders}) GROUP BY event_id, status",
        tuple(event_ids),
    )
    counts_map: dict = {}
    for r in counts_rows:
        counts_map.setdefault(r["event_id"], {"going": 0, "maybe": 0, "not_going": 0})
        counts_map[r["event_id"]][r["status"]] = r["c"]

    # My RSVP status
    my_rows = await db_fetchall(
        f"SELECT event_id, status FROM event_rsvps WHERE account_id=%s AND event_id IN ({placeholders})",
        (account_id, *event_ids),
    )
    my_map = {r["event_id"]: r["status"] for r in my_rows}

    # Top-6 'going' attendees per event (name + email fallback)
    attendee_rows = await db_fetchall(
        f"""
        SELECT r.event_id, COALESCE(ud.name, ac.email) AS display_name
        FROM event_rsvps r
        JOIN accounts ac ON ac.account_id=r.account_id
        LEFT JOIN user_details ud ON ud.user_id=ac.user_id
        WHERE r.status='going' AND r.event_id IN ({placeholders})
        ORDER BY r.updated_at DESC
        """,
        tuple(event_ids),
    )
    attendees_map: dict = {}
    for r in attendee_rows:
        attendees_map.setdefault(r["event_id"], []).append(r["display_name"])

    for ev in events_list:
        ev["rsvp_counts"] = counts_map.get(ev["id"], {"going": 0, "maybe": 0, "not_going": 0})
        ev["my_rsvp_status"] = my_map.get(ev["id"])
        ev["attendees_preview"] = attendees_map.get(ev["id"], [])[:6]
    return events_list


@api_router.get("/events")
async def list_events(
    scope: str = "upcoming",
    association_id: Optional[str] = None,
    account: dict = Depends(get_current_account),
):
    scope = (scope or "upcoming").lower()
    
    is_homeowner = account["role"] == "Homeowner"
    is_board = account["role"] == "Board member"
    is_admin = account["role"] in ("Admin", "Super admin")
    
    where_clause = "WHERE 1=1"
    params = []
    
    if is_homeowner or is_board:
        res = await db_fetchall('''
            SELECT b.association_id 
            FROM accounts ac
            JOIN user_details ud ON ac.user_id = ud.user_id
            JOIN units u ON ud.unit_id = u.id
            JOIN blocks b ON u.block_id = b.id
            WHERE ac.account_id = %s
        ''', (account["account_id"],))
        if res:
            where_clause += " AND e.association_id = %s"
            params.append(res[0]['association_id'])
        else:
            return []
    elif is_admin:
        if association_id:
            res = await db_fetchall("SELECT association_id FROM admin_associations WHERE admin_id = %s AND association_id = %s", (account["account_id"], association_id))
            if res:
                where_clause += " AND e.association_id = %s"
                params.append(association_id)
            else:
                return []
        else:
            res = await db_fetchall("SELECT association_id FROM admin_associations WHERE admin_id = %s", (account["account_id"],))
            if res:
                assoc_ids = [r['association_id'] for r in res]
                placeholders = ','.join(['%s'] * len(assoc_ids))
                where_clause += f" AND e.association_id IN ({placeholders})"
                params.extend(assoc_ids)
            else:
                return []
                
    if is_homeowner:
        where_clause += " AND e.audience IN ('All', 'Homeowner', 'Homeowners') AND e.status = 'Published'"
        
    if scope == "past":
        where_clause += " AND e.starts_at < NOW()"
        order_clause = "ORDER BY e.starts_at DESC"
    elif scope == "all":
        order_clause = "ORDER BY e.starts_at DESC"
    else:  # upcoming
        where_clause += " AND e.starts_at >= NOW()"
        order_clause = "ORDER BY e.starts_at ASC"

    query = f"""
        SELECT e.*, COALESCE(ud.name, ac.email) AS author_name,
               (SELECT COUNT(*) FROM event_likes WHERE event_id = e.id) as like_count,
               (SELECT COUNT(*) FROM event_comments WHERE event_id = e.id) as comment_count,
               EXISTS(SELECT 1 FROM event_likes WHERE event_id = e.id AND account_id = %s) as user_has_liked
        FROM events e
        LEFT JOIN accounts ac ON ac.account_id=e.created_by
        LEFT JOIN user_details ud ON ud.user_id=ac.user_id
        {where_clause}
        {order_clause}
    """
    
    final_params = (account["account_id"],) + tuple(params)
    rows = await db_fetchall(query, final_params)
    for row in rows:
        row['user_has_liked'] = bool(row['user_has_liked'])
    serialized = [_serialize_event(r) for r in rows]
    return await _attach_rsvp_data(serialized, account["account_id"])


@api_router.post("/events")
async def create_event(payload: EventIn, account: dict = Depends(require_role(["Admin", "Super admin", "Board member"]))):
    assoc_id = payload.association_id
    if not assoc_id or assoc_id == "me":
        if account["role"] == "Board member":
            res = await db_fetchall('''
                SELECT b.association_id 
                FROM accounts ac
                JOIN user_details ud ON ac.user_id = ud.user_id
                JOIN units u ON ud.unit_id = u.id
                JOIN blocks b ON u.block_id = b.id
                WHERE ac.account_id = %s
            ''', (account["account_id"],))
            if not res:
                raise HTTPException(400, "Board member has no association assigned")
            assoc_id = res[0]['association_id']
        else:
            raise HTTPException(400, "association_id is required for admins")
            
    eid = await db_execute(
        """INSERT INTO events (
            association_id, title, description, category, banner_url, location, 
            starts_at, ends_at, is_registration_required, registration_deadline, max_capacity, 
            audience, is_paid, fee_amount, organizer_name, organizer_contact, status, created_by
        ) VALUES (%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s)""",
        (
            assoc_id, payload.title, payload.description, payload.category, payload.banner_url, payload.location,
            payload.starts_at, payload.ends_at, 1 if payload.is_registration_required else 0, payload.registration_deadline, payload.max_capacity,
            payload.audience, 1 if payload.is_paid else 0, payload.fee_amount, payload.organizer_name, payload.organizer_contact, payload.status, account["account_id"]
        ),
    )
    return {"ok": True, "id": eid}


@api_router.put("/admin/events/{eid}")
async def admin_update_event(eid: str, payload: EventIn, account: dict = Depends(require_role(["Admin", "Super admin", "Board member"]))):
    row = await db_fetchone("SELECT id, created_at FROM events WHERE id=%s", (eid,))
    if not row:
        raise HTTPException(status_code=404, detail="Event not found.")
        
    if account["role"] == "Board member":
        if row["created_at"]:
            from datetime import datetime, timedelta
            if datetime.utcnow() - row["created_at"] > timedelta(hours=2):
                raise HTTPException(status_code=403, detail="Board members cannot edit events after 2 hours.")

    await db_execute(
        """UPDATE events SET 
            title=%s, description=%s, category=%s, banner_url=%s, location=%s, 
            starts_at=%s, ends_at=%s, is_registration_required=%s, registration_deadline=%s, max_capacity=%s, 
            audience=%s, is_paid=%s, fee_amount=%s, organizer_name=%s, organizer_contact=%s, status=%s 
        WHERE id=%s""",
        (
            payload.title, payload.description, payload.category, payload.banner_url, payload.location,
            payload.starts_at, payload.ends_at, 1 if payload.is_registration_required else 0, payload.registration_deadline, payload.max_capacity,
            payload.audience, 1 if payload.is_paid else 0, payload.fee_amount, payload.organizer_name, payload.organizer_contact, payload.status, eid
        ),
    )
    return {"ok": True}


@api_router.delete("/admin/events/{eid}")
async def admin_delete_event(eid: str, account: dict = Depends(require_role(["Admin", "Super admin", "Board member"]))):
    row = await db_fetchone("SELECT id, created_at FROM events WHERE id=%s", (eid,))
    if not row:
        raise HTTPException(status_code=404, detail="Event not found.")
        
    if account["role"] == "Board member":
        if row["created_at"]:
            from datetime import datetime, timedelta
            if datetime.utcnow() - row["created_at"] > timedelta(hours=2):
                raise HTTPException(status_code=403, detail="Board members cannot delete events after 2 hours.")

    await db_execute("DELETE FROM events WHERE id=%s", (eid,))
    return {"ok": True}


@api_router.post("/events/{eid}/rsvp")
async def rsvp_event(eid: str, payload: RsvpIn, account: dict = Depends(get_current_account)):
    ev = await db_fetchone("SELECT id FROM events WHERE id=%s", (eid,))
    if not ev:
        raise HTTPException(status_code=404, detail="Event not found.")
    await db_execute(
        """
        INSERT INTO event_rsvps (event_id, account_id, status) VALUES (%s,%s,%s)
        ON DUPLICATE KEY UPDATE status=VALUES(status), updated_at=CURRENT_TIMESTAMP
        """,
        (eid, account["account_id"], payload.status),
    )
    return {"ok": True, "status": payload.status}


@api_router.delete("/events/{eid}/rsvp")
async def clear_rsvp(eid: str, account: dict = Depends(get_current_account)):
    await db_execute(
        "DELETE FROM event_rsvps WHERE event_id=%s AND account_id=%s",
        (eid, account["account_id"]),
    )
    return {"ok": True}


@api_router.get("/admin/events/{eid}/rsvps")
async def admin_event_rsvps(eid: str, _: dict = Depends(require_admin)):
    ev = await db_fetchone("SELECT id, title FROM events WHERE id=%s", (eid,))
    if not ev:
        raise HTTPException(status_code=404, detail="Event not found.")
    rows = await db_fetchall(
        """
        SELECT r.status, r.updated_at, ac.email,
               COALESCE(ud.name, ac.email) AS name, ud.contact_number
        FROM event_rsvps r
        JOIN accounts ac ON ac.account_id=r.account_id
        LEFT JOIN user_details ud ON ud.user_id=ac.user_id
        WHERE r.event_id=%s
        ORDER BY FIELD(r.status,'going','maybe','not_going'), r.updated_at DESC
        """,
        (eid,),
    )
    for r in rows:
        if r.get("updated_at"):
            r["updated_at"] = r["updated_at"].isoformat()
    counts = {"going": 0, "maybe": 0, "not_going": 0}
    for r in rows:
        counts[r["status"]] += 1
    return {"event": ev, "counts": counts, "rsvps": rows}


@api_router.post("/admin/events")
async def admin_create_event(payload: EventIn, admin: dict = Depends(require_admin)):
    if payload.ends_at and payload.ends_at < payload.starts_at:
        raise HTTPException(status_code=400, detail="End time must be after start time.")
    eid = await db_execute(
        "INSERT INTO events (title, description, location, starts_at, ends_at, created_by) VALUES (%s,%s,%s,%s,%s,%s)",
        (payload.title, payload.description, payload.location, payload.starts_at, payload.ends_at, admin["account_id"]),
    )
    return {"ok": True, "id": eid}


@api_router.put("/admin/events/{eid}")
async def admin_update_event(eid: str, payload: EventIn, _: dict = Depends(require_admin)):
    row = await db_fetchone("SELECT id FROM events WHERE id=%s", (eid,))
    if not row:
        raise HTTPException(status_code=404, detail="Event not found.")
    if payload.ends_at and payload.ends_at < payload.starts_at:
        raise HTTPException(status_code=400, detail="End time must be after start time.")
    await db_execute(
        "UPDATE events SET title=%s, description=%s, location=%s, starts_at=%s, ends_at=%s WHERE id=%s",
        (payload.title, payload.description, payload.location, payload.starts_at, payload.ends_at, eid),
    )
    return {"ok": True}


@api_router.delete("/admin/events/{eid}")
async def admin_delete_event(eid: str, _: dict = Depends(require_admin)):
    row = await db_fetchone("SELECT id FROM events WHERE id=%s", (eid,))
    if not row:
        raise HTTPException(status_code=404, detail="Event not found.")
    await db_execute("DELETE FROM events WHERE id=%s", (eid,))
    return {"ok": True}

# ---------- Polls ----------

@api_router.get("/polls")
async def list_polls(
    association_id: Optional[str] = None,
    account: dict = Depends(get_current_account),
):
    is_homeowner = account["role"] == "Homeowner"
    is_board = account["role"] == "Board member"
    is_admin = account["role"] in ("Admin", "Super admin")
    
    where_clause = "WHERE p.status = 'Published'"
    params = []
    
    if is_homeowner or is_board:
        res = await db_fetchall('''
            SELECT b.association_id 
            FROM accounts ac
            JOIN user_details ud ON ac.user_id = ud.user_id
            JOIN units u ON ud.unit_id = u.id
            JOIN blocks b ON u.block_id = b.id
            WHERE ac.account_id = %s
        ''', (account["account_id"],))
        if res:
            where_clause += " AND p.association_id = %s"
            params.append(res[0]['association_id'])
            if is_homeowner:
                where_clause += " AND p.visibility IN ('All', 'Homeowner', 'Homeowners')"
            elif is_board:
                where_clause += " AND p.visibility IN ('All', 'Board Members')"
        else:
            return []
    elif is_admin:
        where_clause = "WHERE 1=1"
        if association_id:
            res = await db_fetchall("SELECT association_id FROM admin_associations WHERE admin_id = %s AND association_id = %s", (account["account_id"], association_id))
            if res:
                where_clause += " AND p.association_id = %s"
                params.append(association_id)
            else:
                return []
        else:
            res = await db_fetchall("SELECT association_id FROM admin_associations WHERE admin_id = %s", (account["account_id"],))
            if res:
                assoc_ids = [r['association_id'] for r in res]
                placeholders = ','.join(['%s'] * len(assoc_ids))
                where_clause += f" AND p.association_id IN ({placeholders})"
                params.extend(assoc_ids)
            else:
                return []
                
    polls = await db_fetchall(f"""
        SELECT p.*, COALESCE(ud.name, e.name, a.email) as author_name
        FROM polls p
        LEFT JOIN accounts a ON p.created_by = a.account_id
        LEFT JOIN user_details ud ON a.user_id = ud.user_id
        LEFT JOIN employees e ON a.employee_id = e.employee_id
        {where_clause}
        ORDER BY p.created_at DESC
    """, tuple(params))
    
    for p in polls:
        p["created_at"] = p["created_at"].isoformat() if p.get("created_at") else None
        p["updated_at"] = p["updated_at"].isoformat() if p.get("updated_at") else None
        p["end_date"] = p["end_date"].isoformat() if p.get("end_date") else None
        p["is_multiple_choice"] = bool(p.get("is_multiple_choice"))
        
        # fetch options
        options = await db_fetchall("SELECT * FROM poll_options WHERE poll_id = %s ORDER BY created_at ASC", (p["id"],))
        p["options"] = options
        
        # fetch votes
        votes = await db_fetchall("SELECT option_id, account_id FROM poll_votes WHERE poll_id = %s", (p["id"],))
        p["total_votes"] = len(set(v["account_id"] for v in votes))
        
        my_votes = [v["option_id"] for v in votes if v["account_id"] == account["account_id"]]
        p["my_votes"] = my_votes
        
        for opt in options:
            opt_votes = len([v for v in votes if v["option_id"] == opt["id"]])
            opt["vote_count"] = opt_votes
            opt["vote_percentage"] = round((opt_votes / p["total_votes"] * 100) if p["total_votes"] > 0 else 0)
            
        # fetch likes
        likes = await db_fetchall("SELECT account_id FROM poll_likes WHERE poll_id = %s", (p["id"],))
        p["like_count"] = len(likes)
        p["user_has_liked"] = any(l["account_id"] == account["account_id"] for l in likes)
        
        # fetch comments
        comments = await db_fetchall("SELECT id FROM poll_comments WHERE poll_id = %s", (p["id"],))
        p["comment_count"] = len(comments)
            
    return polls



@api_router.post("/polls/{poll_id}/like")
async def toggle_poll_like(poll_id: str, account: dict = Depends(get_current_account)):
    row = await db_fetchone("SELECT * FROM poll_likes WHERE poll_id = %s AND account_id = %s", (poll_id, account["account_id"]))
    if row:
        await db_execute("DELETE FROM poll_likes WHERE poll_id = %s AND account_id = %s", (poll_id, account["account_id"]))
        return {"ok": True, "liked": False}
    else:
        await db_execute("INSERT INTO poll_likes (poll_id, account_id) VALUES (%s, %s)", (poll_id, account["account_id"]))
        return {"ok": True, "liked": True}

@api_router.get("/polls/{poll_id}/comments")
async def get_poll_comments(poll_id: str, account: dict = Depends(get_current_account)):
    comments = await db_fetchall('''
        SELECT c.*, COALESCE(ud.name, e.name, a.email) as author_name
        FROM poll_comments c
        JOIN accounts a ON c.account_id = a.account_id
        LEFT JOIN user_details ud ON a.user_id = ud.user_id
        LEFT JOIN employees e ON a.employee_id = e.employee_id
        WHERE c.poll_id = %s
        ORDER BY c.created_at ASC
    ''', (poll_id,))
    for c in comments:
        c["created_at"] = c["created_at"].isoformat() if c.get("created_at") else None
    return {"ok": True, "data": comments}

class CommentIn(BaseModel):
    comment: str

@api_router.post("/polls/{poll_id}/comments")
async def post_poll_comment(poll_id: str, payload: CommentIn, account: dict = Depends(get_current_account)):
    await db_execute("INSERT INTO poll_comments (poll_id, account_id, content) VALUES (%s, %s, %s)", (poll_id, account["account_id"], payload.comment))
    return {"ok": True}

@api_router.post("/polls")
@api_router.post("/admin/polls")
async def create_poll(payload: PollIn, account: dict = Depends(get_current_account)):
    try:
        if account["role"] not in ["Super admin", "Admin", "Board member"]:
            raise HTTPException(status_code=403, detail="Not authorized")
            
        assoc_id = payload.association_id
        if assoc_id and (str(assoc_id).strip() == "" or str(assoc_id).upper() == "ALL"):
            assoc_id = None

        if account["role"] == "Board member":
            res = await db_fetchall('''
                SELECT b.association_id 
                FROM accounts ac
                JOIN user_details ud ON ac.user_id = ud.user_id
                JOIN units u ON ud.unit_id = u.id
                JOIN blocks b ON u.block_id = b.id
                WHERE ac.account_id = %s
            ''', (account["account_id"],))
            if res:
                assoc_id = res[0]["association_id"]
            else:
                raise HTTPException(status_code=403, detail="Board member has no association.")
                
        end_date_val = payload.end_date if payload.end_date else None
        poll_id = await db_execute("""
            INSERT INTO polls (association_id, question, description, visibility, status, is_multiple_choice, end_date, created_by)
            VALUES (%s, %s, %s, %s, %s, %s, %s, %s)
        """, (assoc_id, payload.question, payload.description, payload.visibility, payload.status, int(payload.is_multiple_choice), end_date_val, account["account_id"]))
        
        for opt in payload.options:
            await db_execute("INSERT INTO poll_options (poll_id, option_text) VALUES (%s, %s)", (poll_id, opt.text))
            
        return {"ok": True, "id": poll_id}
    except Exception as e:
        import traceback
        traceback.print_exc()
        raise HTTPException(status_code=500, detail=str(e))

@api_router.put("/polls/{poll_id}")
@api_router.put("/admin/polls/{poll_id}")
async def update_poll(poll_id: str, payload: PollIn, account: dict = Depends(get_current_account)):
    if account["role"] not in ["Super admin", "Admin", "Board member"]:
        raise HTTPException(status_code=403, detail="Not authorized")
        
    row = await db_fetchone("SELECT id FROM polls WHERE id = %s", (poll_id,))
    if not row:
        raise HTTPException(status_code=404, detail="Poll not found")
        
    end_date_val = payload.end_date if payload.end_date else None
    await db_execute("""
        UPDATE polls 
        SET question=%s, description=%s, visibility=%s, status=%s, is_multiple_choice=%s, end_date=%s
        WHERE id=%s
    """, (payload.question, payload.description, payload.visibility, payload.status, int(payload.is_multiple_choice), end_date_val, poll_id))
    
    await db_execute("DELETE FROM poll_options WHERE poll_id = %s", (poll_id,))
    for opt in payload.options:
        await db_execute("INSERT INTO poll_options (poll_id, option_text) VALUES (%s, %s)", (poll_id, opt.text))
        
    return {"ok": True}

@api_router.delete("/polls/{poll_id}")
@api_router.delete("/admin/polls/{poll_id}")
async def delete_poll(poll_id: str, account: dict = Depends(get_current_account)):
    if account["role"] not in ["Super admin", "Admin", "Board member"]:
        raise HTTPException(status_code=403, detail="Not authorized")
    await db_execute("DELETE FROM polls WHERE id = %s", (poll_id,))
    return {"ok": True}
    
@api_router.post("/polls/{poll_id}/vote")
async def vote_poll(poll_id: str, payload: PollVoteIn, account: dict = Depends(get_current_account)):
    poll = await db_fetchone("SELECT id, is_multiple_choice, status FROM polls WHERE id = %s", (poll_id,))
    if not poll:
        raise HTTPException(status_code=404, detail="Poll not found")
    if poll["status"] != "Published":
        raise HTTPException(status_code=400, detail="Poll is not active")
        
    if not poll["is_multiple_choice"] and len(payload.option_ids) > 1:
        raise HTTPException(status_code=400, detail="Multiple selection is not allowed for this poll")
        
    await db_execute("""
        DELETE FROM poll_votes 
        WHERE poll_id = %s AND account_id = %s
    """, (poll_id, account["account_id"]))
    
    for opt_id in payload.option_ids:
        await db_execute("""
            INSERT INTO poll_votes (poll_id, option_id, account_id)
            VALUES (%s, %s, %s)
        """, (poll_id, opt_id, account["account_id"]))
        
    return {"ok": True}


@api_router.post("/events/{eid}/like")
async def toggle_event_like(eid: str, account: dict = Depends(get_current_account)):
    row = await db_fetchone("SELECT 1 FROM events WHERE id=%s", (eid,))
    if not row:
        raise HTTPException(status_code=404, detail="Event not found")
        
    acc_id = account["account_id"]
    existing = await db_fetchone("SELECT 1 FROM event_likes WHERE event_id=%s AND account_id=%s", (eid, acc_id))
    
    if existing:
        await db_execute("DELETE FROM event_likes WHERE event_id=%s AND account_id=%s", (eid, acc_id))
        return {"ok": True, "action": "unliked"}
    else:
        await db_execute("INSERT INTO event_likes (event_id, account_id) VALUES (%s, %s)", (eid, acc_id))
        return {"ok": True, "action": "liked"}


@api_router.get("/events/{eid}/comments")
async def get_event_comments(eid: str, _: dict = Depends(get_current_account)):
    row = await db_fetchone("SELECT 1 FROM events WHERE id=%s", (eid,))
    if not row:
        raise HTTPException(status_code=404, detail="Event not found")
        
    rows = await db_fetchall("""
        SELECT c.id, c.body AS comment, c.created_at, c.account_id, COALESCE(ud.name, ac.email) AS author_name
        FROM event_comments c
        LEFT JOIN accounts ac ON ac.account_id = c.account_id
        LEFT JOIN user_details ud ON ud.user_id = ac.user_id
        WHERE c.event_id=%s
        ORDER BY c.created_at ASC
    """, (eid,))
    
    for r in rows:
        if isinstance(r["created_at"], (datetime, date)):
            r["created_at"] = r["created_at"].isoformat()
    return rows


@api_router.post("/events/{eid}/comments")
async def post_event_comment(eid: str, payload: CommentIn, account: dict = Depends(get_current_account)):
    row = await db_fetchone("SELECT 1 FROM events WHERE id=%s", (eid,))
    if not row:
        raise HTTPException(status_code=404, detail="Event not found")
        
    cid = await db_execute(
        "INSERT INTO event_comments (event_id, account_id, body) VALUES (%s, %s, %s)",
        (eid, account["account_id"], payload.comment)
    )
    
    c_row = await db_fetchone("""
        SELECT c.id, c.body AS comment, c.created_at, c.account_id, COALESCE(ud.name, ac.email) AS author_name
        FROM event_comments c
        LEFT JOIN accounts ac ON ac.account_id = c.account_id
        LEFT JOIN user_details ud ON ud.user_id = ac.user_id
        WHERE c.id=%s
    """, (cid,))
    
    if c_row and isinstance(c_row["created_at"], (datetime, date)):
        c_row["created_at"] = c_row["created_at"].isoformat()
        
    return {"ok": True, "data": c_row}


# ---------- Profile Endpoints ----------

@api_router.get("/profile/data")
async def get_profile_data(account: dict = Depends(get_current_account)):
    uid = account.get("user_id")
    eid = account.get("employee_id")
    
    user_details = {}
    family_members = []
    unit_homeowners = []
    vehicles = []
    pets = []
    education = []
    experience = []

    if eid:
        # It's an employee
        emp = await db_fetchone("SELECT * FROM employees WHERE employee_id=%s", (eid,))
        if emp:
            user_details = {
                "name": emp.get("name"),
                "first_name": emp.get("first_name"),
                "last_name": emp.get("last_name"),
                "email": emp.get("email"),
                "contact_number": emp.get("contact_number"),
                "address": emp.get("address"),
                "profile_pic_url": emp.get("profile_pic_url")
            }
        education = await db_fetchall("SELECT * FROM employee_education WHERE employee_id=%s ORDER BY created_at DESC", (eid,))
        experience = await db_fetchall("SELECT * FROM employee_experience WHERE employee_id=%s ORDER BY created_at DESC", (eid,))
    elif uid:
        user_details = await db_fetchone("SELECT * FROM user_details WHERE user_id=%s", (uid,))
        if not user_details:
            user_details = {}

        family_members = await db_fetchall("SELECT * FROM family_members WHERE user_id=%s", (uid,))
        
        unit_id = user_details.get("unit_id") if user_details else None
        if unit_id:
            vehicles = await db_fetchall(
                "SELECT v.* FROM vehicles v JOIN user_details ud ON v.user_id = ud.user_id WHERE ud.unit_id=%s", 
                (unit_id,)
            )
            pets = await db_fetchall(
                "SELECT p.* FROM pets p JOIN user_details ud ON p.user_id = ud.user_id WHERE ud.unit_id=%s", 
                (unit_id,)
            )
        else:
            vehicles = await db_fetchall("SELECT * FROM vehicles WHERE user_id=%s", (uid,))
            pets = await db_fetchall("SELECT * FROM pets WHERE user_id=%s", (uid,))
        
        if user_details and user_details.get("unit_id"):
            # Fetch other homeowners in the same unit
            unit_homeowners = await db_fetchall(
                "SELECT * FROM user_details WHERE unit_id=%s AND user_id != %s", 
                (user_details["unit_id"], uid)
            )

    if not user_details and (account.get("role_code") == "super_admin" or not (uid or eid)):
        user_details = {
            "name": "Super Administrator" if account.get("role_code") == "super_admin" else "Platform Administrator",
            "email": account.get("email"),
            "contact_number": "-",
            "address": "Global Platform Scope",
            "profile_pic_url": None,
            "association_name": "Nestora Platform"
        }

    # Fetch association info for the homeowner / user
    assoc_info = await db_fetchone("""
        SELECT 
            COALESCE(a.name, a2.name, a3.name, a4.name) as association_name,
            COALESCE(a.id, a2.id, a3.id, a4.id) as association_id
        FROM accounts ac
        LEFT JOIN user_details ud ON ac.user_id = ud.user_id
        LEFT JOIN units u ON ud.unit_id = u.id
        LEFT JOIN blocks b ON u.block_id = b.id
        LEFT JOIN associations a ON b.association_id = a.id
        LEFT JOIN associations a2 ON ud.association_id = a2.id
        LEFT JOIN admin_associations aa ON aa.admin_id = ac.account_id
        LEFT JOIN associations a3 ON aa.association_id = a3.id
        LEFT JOIN board_members bm ON bm.account_id = ac.account_id AND bm.status = 'active'
        LEFT JOIN associations a4 ON bm.association_id = a4.id
        WHERE ac.account_id = %s OR ud.user_id = %s
        LIMIT 1
    """, (account.get("account_id"), uid))

    assoc_name = assoc_info.get("association_name") if assoc_info else None
    assoc_id = assoc_info.get("association_id") if assoc_info else None

    if user_details:
        if not user_details.get("association_name"):
            user_details["association_name"] = assoc_name or "Nestora Platform"
        if assoc_id and not user_details.get("association_id"):
            user_details["association_id"] = assoc_id

    # ensure they are lists
    if not isinstance(family_members, list):
        family_members = []
    if not isinstance(vehicles, list):
        vehicles = []
    if not isinstance(pets, list):
        pets = []
    if not isinstance(unit_homeowners, list):
        unit_homeowners = []
    if not isinstance(education, list):
        education = []
    if not isinstance(experience, list):
        experience = []


    def format_dates(rows):
        for r in rows:
            for k, v in r.items():
                if hasattr(v, "isoformat"):
                    r[k] = v.isoformat()
        return rows

    def format_dict_dates(d):
        for k, v in d.items():
            if hasattr(v, "isoformat"):
                d[k] = v.isoformat()
        return d

    user_details = format_dict_dates(user_details)
    family_members = format_dates(family_members)
    vehicles = format_dates(vehicles)
    pets = format_dates(pets)
    education = format_dates(education)
    experience = format_dates(experience)
    unit_homeowners = format_dates(unit_homeowners)

    return {
        "ok": True,
        "user_details": user_details,
        "family_members": family_members,
        "vehicles": vehicles,
        "pets": pets,
        "education": education,
        "experience": experience,
        "unit_homeowners": unit_homeowners,
        "association_name": assoc_name,
        "association_id": assoc_id
    }

@api_router.put("/profile/user-details")
async def update_profile_details(payload: ProfileUpdateIn, account: dict = Depends(get_current_account)):
    uid = account.get("user_id")
    eid = account.get("employee_id")
    
    updates = []
    vals = []
    if payload.name is not None:
        updates.append("name=%s")
        vals.append(payload.name)
    if payload.address is not None:
        updates.append("address=%s")
        vals.append(payload.address)
    if payload.email is not None:
        updates.append("email=%s")
        vals.append(payload.email)
    if payload.contact_number is not None:
        updates.append("contact_number=%s")
        vals.append(payload.contact_number)
    if payload.alt_contact_number is not None and uid:
        updates.append("alt_contact_number=%s")
        vals.append(payload.alt_contact_number)
    if "profile_pic_url" in payload.model_fields_set:
        updates.append("profile_pic_url=%s")
        vals.append(payload.profile_pic_url if payload.profile_pic_url else None)
        
    if updates:
        if uid:
            vals.append(uid)
            q = f"UPDATE user_details SET {', '.join(updates)} WHERE user_id=%s"
            await db_execute(q, tuple(vals))
        elif eid:
            vals.append(eid)
            q = f"UPDATE employees SET {', '.join(updates)} WHERE employee_id=%s"
            await db_execute(q, tuple(vals))
    
    return {"ok": True}

@api_router.post("/profile/family-members")
async def add_family_member(payload: FamilyMemberIn, account: dict = Depends(get_current_account)):
    uid = account["user_id"]
    
    fid = await db_execute(
        "INSERT INTO family_members (user_id, name, email, contact_number, alt_contact_number) VALUES (%s,%s,%s,%s,%s)",
        (uid, payload.name, payload.email, payload.contact_number, payload.alt_contact_number)
    )
    return {"ok": True, "id": fid}

@api_router.post("/profile/education")
async def add_education(payload: EmployeeEducationIn, account: dict = Depends(get_current_account)):
    eid = account.get("employee_id")
    if not eid:
        raise HTTPException(status_code=403, detail="Only employees can add education")
    
    insert_id = await db_execute(
        "INSERT INTO employee_education (employee_id, education_level, degree, field_of_study, institution, board_university, start_date, end_date, currently_studying, grade, location, description, certificate_url) VALUES (%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s)",
        (eid, payload.education_level, payload.degree, payload.field_of_study, payload.institution, payload.board_university, payload.start_date, payload.end_date, payload.currently_studying, payload.grade, payload.location, payload.description, payload.certificate_url)
    )
    return {"ok": True, "id": insert_id}

@api_router.put("/profile/education/{eid_id}")
async def update_education(eid_id: str, payload: EmployeeEducationIn, account: dict = Depends(get_current_account)):
    eid = account.get("employee_id")
    if not eid:
        raise HTTPException(status_code=403, detail="Only employees can update education")
    
    row = await db_fetchone("SELECT id FROM employee_education WHERE id=%s AND employee_id=%s", (eid_id, eid))
    if not row:
        raise HTTPException(status_code=404, detail="Education record not found")
        
    await db_execute(
        "UPDATE employee_education SET education_level=%s, degree=%s, field_of_study=%s, institution=%s, board_university=%s, start_date=%s, end_date=%s, currently_studying=%s, grade=%s, location=%s, description=%s, certificate_url=%s WHERE id=%s",
        (payload.education_level, payload.degree, payload.field_of_study, payload.institution, payload.board_university, payload.start_date, payload.end_date, payload.currently_studying, payload.grade, payload.location, payload.description, payload.certificate_url, eid_id)
    )
    return {"ok": True}

@api_router.delete("/profile/education/{eid_id}")
async def delete_education(eid_id: str, account: dict = Depends(get_current_account)):
    eid = account.get("employee_id")
    if not eid:
        raise HTTPException(status_code=403, detail="Only employees can delete education")
        
    row = await db_fetchone("SELECT id FROM employee_education WHERE id=%s AND employee_id=%s", (eid_id, eid))
    if not row:
        raise HTTPException(status_code=404, detail="Education record not found")
        
    await db_execute("DELETE FROM employee_education WHERE id=%s", (eid_id,))
    return {"ok": True}

@api_router.post("/profile/experience")
async def add_experience(payload: EmployeeExperienceIn, account: dict = Depends(get_current_account)):
    eid = account.get("employee_id")
    if not eid:
        raise HTTPException(status_code=403, detail="Only employees can add experience")
    
    insert_id = await db_execute(
        "INSERT INTO employee_experience (employee_id, job_title, employment_type, company, industry, location, work_mode, start_date, end_date, currently_working, description, skills, website_url, certificate_url) VALUES (%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s)",
        (eid, payload.job_title, payload.employment_type, payload.company, payload.industry, payload.location, payload.work_mode, payload.start_date, payload.end_date, payload.currently_working, payload.description, payload.skills, payload.website_url, payload.certificate_url)
    )
    return {"ok": True, "id": insert_id}

@api_router.put("/profile/experience/{exp_id}")
async def update_experience(exp_id: str, payload: EmployeeExperienceIn, account: dict = Depends(get_current_account)):
    eid = account.get("employee_id")
    if not eid:
        raise HTTPException(status_code=403, detail="Only employees can update experience")
    
    row = await db_fetchone("SELECT id FROM employee_experience WHERE id=%s AND employee_id=%s", (exp_id, eid))
    if not row:
        raise HTTPException(status_code=404, detail="Experience record not found")
        
    await db_execute(
        "UPDATE employee_experience SET job_title=%s, employment_type=%s, company=%s, industry=%s, location=%s, work_mode=%s, start_date=%s, end_date=%s, currently_working=%s, description=%s, skills=%s, website_url=%s, certificate_url=%s WHERE id=%s",
        (payload.job_title, payload.employment_type, payload.company, payload.industry, payload.location, payload.work_mode, payload.start_date, payload.end_date, payload.currently_working, payload.description, payload.skills, payload.website_url, payload.certificate_url, exp_id)
    )
    return {"ok": True}

@api_router.delete("/profile/experience/{exp_id}")
async def delete_experience(exp_id: str, account: dict = Depends(get_current_account)):
    eid = account.get("employee_id")
    if not eid:
        raise HTTPException(status_code=403, detail="Only employees can delete experience")
        
    row = await db_fetchone("SELECT id FROM employee_experience WHERE id=%s AND employee_id=%s", (exp_id, eid))
    if not row:
        raise HTTPException(status_code=404, detail="Experience record not found")
        
    await db_execute("DELETE FROM employee_experience WHERE id=%s", (exp_id,))
    return {"ok": True}

@api_router.post("/profile/vehicles")
async def add_vehicle(payload: VehicleIn, account: dict = Depends(get_current_account)):
    uid = account["user_id"]
    
    vid = await db_execute(
        "INSERT INTO vehicles (user_id, type, registration_number, insurance_url, puc_url) VALUES (%s,%s,%s,%s,%s)",
        (uid, payload.type, payload.registration_number, payload.insurance_url, payload.puc_url)
    )
    return {"ok": True, "id": vid}

@api_router.put("/profile/vehicles/{vid}")
async def edit_vehicle(vid: str, payload: VehicleIn, account: dict = Depends(get_current_account)):
    uid = account["user_id"]
    row = await db_fetchone(
        "SELECT v.id FROM vehicles v LEFT JOIN user_details ud ON v.user_id = ud.user_id WHERE v.id=%s AND (v.user_id=%s OR (ud.unit_id IS NOT NULL AND ud.unit_id = (SELECT unit_id FROM user_details WHERE user_id=%s)))", 
        (vid, uid, uid)
    )
    if not row:
        raise HTTPException(status_code=404, detail="Vehicle not found or unauthorized.")
        
    await db_execute(
        "UPDATE vehicles SET type=%s, registration_number=%s, insurance_url=%s, puc_url=%s WHERE id=%s",
        (payload.type, payload.registration_number, payload.insurance_url, payload.puc_url, vid)
    )
    return {"ok": True}

@api_router.delete("/profile/vehicles/{vid}")
async def delete_vehicle(vid: str, account: dict = Depends(get_current_account)):
    uid = account["user_id"]
    row = await db_fetchone(
        "SELECT v.id FROM vehicles v LEFT JOIN user_details ud ON v.user_id = ud.user_id WHERE v.id=%s AND (v.user_id=%s OR (ud.unit_id IS NOT NULL AND ud.unit_id = (SELECT unit_id FROM user_details WHERE user_id=%s)))", 
        (vid, uid, uid)
    )
    if not row:
        raise HTTPException(status_code=404, detail="Vehicle not found or unauthorized.")
        
    await db_execute("DELETE FROM vehicles WHERE id=%s", (vid,))
    return {"ok": True}

@api_router.post("/profile/pets")
async def add_pet(payload: PetIn, account: dict = Depends(get_current_account)):
    uid = account["user_id"]
    
    pid = await db_execute(
        "INSERT INTO pets (user_id, type, name, breed, vaccinated, vaccination_date, next_vaccination_reminder, vaccination_certificate_url, reminder_date) VALUES (%s,%s,%s,%s,%s,%s,%s,%s,%s)",
        (uid, payload.type, payload.name, payload.breed, payload.vaccinated, payload.vaccination_date, payload.next_vaccination_reminder, payload.vaccination_certificate_url, payload.reminder_date)
    )
    return {"ok": True, "id": pid}

@api_router.put("/profile/pets/{pid}")
async def edit_pet(pid: str, payload: PetIn, account: dict = Depends(get_current_account)):
    uid = account["user_id"]
    row = await db_fetchone(
        "SELECT p.id FROM pets p LEFT JOIN user_details ud ON p.user_id = ud.user_id WHERE p.id=%s AND (p.user_id=%s OR (ud.unit_id IS NOT NULL AND ud.unit_id = (SELECT unit_id FROM user_details WHERE user_id=%s)))", 
        (pid, uid, uid)
    )
    if not row:
        raise HTTPException(status_code=404, detail="Pet not found or unauthorized.")
        
    await db_execute(
        "UPDATE pets SET type=%s, name=%s, breed=%s, vaccinated=%s, vaccination_date=%s, next_vaccination_reminder=%s, vaccination_certificate_url=%s, reminder_date=%s WHERE id=%s",
        (payload.type, payload.name, payload.breed, payload.vaccinated, payload.vaccination_date, payload.next_vaccination_reminder, payload.vaccination_certificate_url, payload.reminder_date, pid)
    )
    return {"ok": True}

@api_router.delete("/profile/pets/{pid}")
async def delete_pet(pid: str, account: dict = Depends(get_current_account)):
    uid = account["user_id"]
    row = await db_fetchone(
        "SELECT p.id FROM pets p LEFT JOIN user_details ud ON p.user_id = ud.user_id WHERE p.id=%s AND (p.user_id=%s OR (ud.unit_id IS NOT NULL AND ud.unit_id = (SELECT unit_id FROM user_details WHERE user_id=%s)))", 
        (pid, uid, uid)
    )
    if not row:
        raise HTTPException(status_code=404, detail="Pet not found or unauthorized.")
        
    await db_execute("DELETE FROM pets WHERE id=%s", (pid,))
    return {"ok": True}

@api_router.get("/admin/associations")
async def get_all_associations(account: dict = Depends(require_role(["Super admin", "Admin", "Accountant", "Board member"]))):
    rows = []
    if account["role"] == "Super admin":
        rows = await db_fetchall("SELECT a.*, (SELECT COUNT(*) FROM units u JOIN blocks b ON u.block_id = b.id WHERE b.association_id = a.id) as unit_count, sp.name as plan_name FROM associations a LEFT JOIN subscription_plans sp ON a.current_plan_id = sp.id ORDER BY a.created_at DESC")
    elif account["role"] in ["Admin", "Accountant"]:
        rows = await db_fetchall(
            """
            SELECT a.*, (SELECT COUNT(*) FROM units u JOIN blocks b ON u.block_id = b.id WHERE b.association_id = a.id) as unit_count, sp.name as plan_name FROM associations a
            JOIN admin_associations aa ON a.id = aa.association_id
            LEFT JOIN subscription_plans sp ON a.current_plan_id = sp.id
            WHERE aa.admin_id = %s
            ORDER BY a.created_at DESC
            """, (account["account_id"],)
        )
    else:
        # Board member
        assoc_id = await get_board_member_association(account["account_id"])
        if assoc_id:
            rows = await db_fetchall("SELECT a.*, (SELECT COUNT(*) FROM units u JOIN blocks b ON u.block_id = b.id WHERE b.association_id = a.id) as unit_count, sp.name as plan_name FROM associations a LEFT JOIN subscription_plans sp ON a.current_plan_id = sp.id WHERE a.id=%s", (assoc_id,))
            
    # Inject allowed features for each association
    if rows:
        plan_features = await db_fetchall("""
            SELECT spf.plan_id, f.name 
            FROM subscription_plan_features spf
            JOIN features f ON spf.feature_id = f.id
        """)
        feature_map = {}
        for pf in plan_features:
            feature_map.setdefault(pf["plan_id"], []).append(pf["name"])
            
        for r in rows:
            r["allowed_features"] = feature_map.get(r["current_plan_id"], []) if r.get("current_plan_id") else []
            
    return rows

@api_router.get("/admin/associations/{id}/stats")
async def get_association_stats(id: str, _: dict = Depends(require_admin)):
    assoc = await db_fetchone("SELECT name, contract_url FROM associations WHERE id = %s", (id,))
    if not assoc:
        raise HTTPException(status_code=404, detail="Association not found")
        
    blocks = await db_fetchone("SELECT count(*) as count FROM blocks WHERE association_id = %s", (id,))
    units = await db_fetchone("SELECT count(*) as count FROM units u JOIN blocks b ON u.block_id = b.id WHERE b.association_id = %s", (id,))
    
    floors = await db_fetchall("SELECT b.name as block_name, count(DISTINCT u.floor) as floors FROM units u JOIN blocks b ON u.block_id = b.id WHERE b.association_id = %s GROUP BY b.id", (id,))
    
    rented = await db_fetchone('''
        SELECT count(*) as count 
        FROM units u 
        JOIN blocks b ON u.block_id = b.id 
        WHERE b.association_id = %s AND u.is_rented = TRUE
    ''', (id,))
    
    return {
        "name": assoc["name"],
        "contract_url": assoc["contract_url"],
        "total_blocks": blocks["count"],
        "total_units": units["count"],
        "floors_per_block": floors,
        "rented_units": rented["count"]
    }

@api_router.get("/admin/associations/excel-template")
async def download_association_excel_template():
    import io
    from openpyxl import Workbook
    from fastapi.responses import StreamingResponse
    
    wb = Workbook()
    
    # Sheet 1: Association Details
    ws1 = wb.active
    ws1.title = "Association Details"
    ws1.append(["Association Name", "Association URL", "Address 1", "Address 2", "City", "State", "Pin Code", "Country"])
    
    # Sheet 2: Details of Unit
    ws2 = wb.create_sheet("Unit Details")
    ws2.append(["Block Name", "Floor", "Unit Number"])
    
    # Sheet 3: Homeowner Details
    ws3 = wb.create_sheet("Homeowner Details")
    ws3.append([
        "Block Name", "Unit Number", "First Name", "Last Name", "Email", "Phone Number", 
        "Rented", "Tenant First Name", "Tenant Last Name", "Tenant Email Id", "Tenant Contact Number"
    ])
    
    # Add Data Validation for Rented column (Yes/No)
    from openpyxl.worksheet.datavalidation import DataValidation
    dv = DataValidation(type="list", formula1='"Yes,No"', allow_blank=True)
    ws3.add_data_validation(dv)
    # Apply to G2:G1048576 (Rented column)
    dv.add("G2:G1048576")
    
    # Sheet 4: Board and Committee Members
    ws4 = wb.create_sheet("Board & Committee Members")
    ws4.append(["Block Name", "Unit Number", "Role", "Committee Name"])
    
    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    
    response = StreamingResponse(iter([output.getvalue()]), media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    response.headers["Content-Disposition"] = "attachment; filename=onboarding_template.xlsx"
    return response


@api_router.post("/admin/associations/onboard")
async def onboard_association(
    entity_id: str = Form(...),
    plan_id: str = Form(None),
    num_blocks: str = Form(None),
    floors_per_block: str = Form(None),
    units_per_floor: str = Form(None),
    contract_file: UploadFile = File(None),
    csv_file: UploadFile = File(...),
    account: dict = Depends(require_admin)
):
    import pandas as pd
    import secrets
    import hashlib
    
    try:
        import io
        contents = await csv_file.read()
        xl = pd.ExcelFile(io.BytesIO(contents))
    except Exception as e:
        raise HTTPException(status_code=400, detail="Invalid Excel file.")
    
    if len(xl.sheet_names) < 3:
        raise HTTPException(status_code=400, detail="Excel file must have at least 3 sheets (Association, Units, Homeowners).")
    
    df1 = xl.parse(0)
    df2 = xl.parse(1)
    df3 = xl.parse(2)
    df4 = xl.parse(3) if len(xl.sheet_names) > 3 else pd.DataFrame()
    
    if df1.empty:
        raise HTTPException(status_code=400, detail="Association details sheet is empty.")
    
    # Fetch roles mapping
    roles = await db_fetchall("SELECT id, name FROM roles")
    role_map = {r['name'].lower(): r['id'] for r in roles}
    homeowner_role_id = role_map.get('homeowner')
    tenant_role_id = role_map.get('tenant')
    board_role_id = role_map.get('board member')
    
    # 1. Parse Association
    assoc_data = df1.iloc[0]
    assoc_name = assoc_data.get("Association Name", "Unknown")
    assoc_url = assoc_data.get("Association URL", "")
    addr1 = assoc_data.get("Address 1", "")
    addr2 = assoc_data.get("Address 2", "")
    city = assoc_data.get("City", "")
    state = assoc_data.get("State", "")
    pincode = assoc_data.get("Pin Code", "")
    country = assoc_data.get("Country", "")
    if pd.isna(country):
        country = ""
        
    def _clean_str(v):
        if pd.isna(v): return ""
        s = str(v).strip()
        return "" if s.lower() in ["nan", "none"] else s

    merged_address_base = ", ".join(filter(None, [
        _clean_str(addr1), 
        _clean_str(addr2), 
        _clean_str(city), 
        _clean_str(state), 
        _clean_str(pincode)
    ]))
    
    # Save contract file logic here if needed, for now just placeholder
    contract_url = None
    if contract_file and contract_file.filename:
        contract_url = f"/uploads/{contract_file.filename}"
    
    # Create Association
    # Create Association Code
    assoc_code = assoc_name.upper().replace(" ", "")[:4]
    if len(assoc_code) < 4:
        assoc_code = assoc_code.ljust(4, "0")
        
    from datetime import datetime
    assoc_id = await db_execute(
        """INSERT INTO associations (name, association_code, entity_id, address_line_1, address_line_2, city, state, pincode, country, url, contract_url, current_plan_id, subscription_status, subscription_start) 
           VALUES (%s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s)""",
        (assoc_name, assoc_code, entity_id, addr1, addr2, city, state, pincode, country, assoc_url, contract_url, plan_id, 'Active' if plan_id else 'Trial', datetime.utcnow() if plan_id else None)
    )
    
    # 2. Parse Blocks & Units
    block_map = {} # block_name -> block_id
    unit_map = {} # (block_name, unit_number) -> unit_id
    
    for _, row in df2.iterrows():
        b_name = str(row.get("Block Name", "")).strip()
        u_num = str(row.get("Unit Number", "")).strip()
        floor = str(row.get("Floor", "")).strip()
        
        if not b_name or not u_num:
            continue
            
        if b_name not in block_map:
            b_id = await db_execute("INSERT INTO blocks (association_id, name) VALUES (%s, %s)", (assoc_id, b_name))
            block_map[b_name] = b_id
            
        b_id = block_map[b_name]
        u_id = await db_execute("INSERT INTO units (block_id, floor, unit_number) VALUES (%s, %s, %s)", (b_id, floor, u_num))
        unit_map[(b_name, u_num)] = u_id

    # 3. Parse Homeowners
    for _, row in df3.iterrows():
        b_name = str(row.get("Block Name", "")).strip()
        u_num = str(row.get("Unit Number", "")).strip()
        
        if not b_name or not u_num:
            continue
            
        if (b_name, u_num) not in unit_map:
            raise HTTPException(status_code=400, detail=f"Homeowner assigned to invalid Unit: Block {b_name}, Unit {u_num}. Ensure it exists in Sheet 2.")
            
        u_id = unit_map[(b_name, u_num)]
        fname = str(row.get("First Name", ""))
        lname = str(row.get("Last Name", ""))
        email = str(row.get("Email", ""))
        phone = str(row.get("Phone Number", ""))
        
        import string
        def gen_code():
            chars = string.ascii_letters + string.digits + "!@#$%^&*"
            while True:
                code = ''.join(secrets.choice(chars) for _ in range(8))
                if (any(c.islower() for c in code) and any(c.isupper() for c in code) 
                    and any(c.isdigit() for c in code) and any(c in "!@#$%^&*" for c in code)):
                    return code

        login_code = gen_code()
        code_id = await db_execute("INSERT INTO user_codes (login_code, status) VALUES (%s, 'active')", (login_code,))
        
        full_address = f"{b_name}-{u_num}, {merged_address_base}" if merged_address_base else f"{b_name}-{u_num}"
        
        user_id = await db_execute(
            """INSERT INTO user_details (name, first_name, last_name, email, contact_number, unit_id, address, code_id, role_id) 
               VALUES (%s, %s, %s, %s, %s, %s, %s, %s, %s)""",
            (f"{fname} {lname}", fname, lname, email, phone, u_id, full_address, code_id, homeowner_role_id)
        )
        
        # Check if Rented
        rented = str(row.get("Rented", "")).strip().title()
        if rented == "Yes":
            await db_execute("UPDATE units SET is_rented=TRUE WHERE id=%s", (u_id,))
            t_fname = str(row.get("Tenant First Name", ""))
            t_lname = str(row.get("Tenant Last Name", ""))
            t_email = str(row.get("Tenant Email Id", ""))
            t_phone = str(row.get("Tenant Contact Number", ""))
            
            if t_email:
                t_login_code = gen_code()
                t_code_id = await db_execute("INSERT INTO user_codes (login_code, status) VALUES (%s, 'active')", (t_login_code,))
                
                t_user_id = await db_execute(
                    """INSERT INTO user_details (name, first_name, last_name, email, contact_number, unit_id, address, code_id, role_id) 
                       VALUES (%s, %s, %s, %s, %s, %s, %s, %s, %s)""",
                    (f"{t_fname} {t_lname}", t_fname, t_lname, t_email, t_phone, u_id, full_address, t_code_id, tenant_role_id)
                )
        
    # 4. Parse Board Members (Optional)
    if not df4.empty:
        committee_map = {}
        for _, row in df4.iterrows():
            b_name = str(row.get("Block Name", "")).strip()
            u_num = str(row.get("Unit Number", "")).strip()
            role = str(row.get("Role", "")).strip()
            c_name = str(row.get("Committee Name", "")).strip()
            
            if not c_name:
                continue
                
            if c_name not in committee_map:
                c_id = await db_execute("INSERT INTO committees (association_id, name) VALUES (%s, %s)", (assoc_id, c_name))
                committee_map[c_name] = c_id
                
            if (b_name, u_num) in unit_map:
                u_id = unit_map[(b_name, u_num)]
                # Find user for this unit
                res = await db_fetchall("SELECT user_id FROM user_details WHERE unit_id=%s LIMIT 1", (u_id,))
                if res:
                    uid = res[0]['user_id']
                    await db_execute("INSERT INTO committee_members (committee_id, user_id, role) VALUES (%s, %s, %s)", (committee_map[c_name], uid, role))

    return {"ok": True, "message": "Association Onboarded"}


# ---------- Service Requests & Uploads ----------
from fastapi import File, UploadFile
import shutil
import uuid
from fastapi.staticfiles import StaticFiles

UPLOAD_DIR = ROOT_DIR / "uploads"
os.makedirs(UPLOAD_DIR, exist_ok=True)

# Mount uploads directory so they can be accessed at /uploads/filename
app.mount("/uploads", StaticFiles(directory=UPLOAD_DIR), name="uploads")

@api_router.post("/upload")
async def upload_file(file: UploadFile = File(...)):
    file_ext = file.filename.split(".")[-1] if "." in file.filename else "bin"
    file_name = f"{uuid.uuid4()}.{file_ext}"
    file_path = UPLOAD_DIR / file_name
    with open(file_path, "wb") as buffer:
        shutil.copyfileobj(file.file, buffer)
    return {"ok": True, "url": f"/uploads/{file_name}"}

# ==============================================================================
# CLOUD UPLOAD ENDPOINTS (UploadThing / Cloudinary Ready)
# ==============================================================================
from cloud_uploader import upload_media_asset, upload_document_asset

@api_router.post("/upload-media-asset")
async def api_upload_media_asset(file: UploadFile = File(...)):
    """Upload Images and Videos to Cloud Storage (UploadThing / Cloudinary)."""
    return await upload_media_asset(file)

@api_router.post("/upload-document-asset")
async def api_upload_document_asset(file: UploadFile = File(...)):
    """Upload Documents and Files to Cloud Storage (UploadThing)."""
    return await upload_document_asset(file)

class ServiceRequestIn(BaseModel):
    service_type: str
    sub_category: Optional[str] = None
    custom_title: Optional[str] = None
    description: Optional[str] = None
    image_url: Optional[str] = None
    user_id: Optional[str] = None
    incoming_call_no: Optional[str] = None

@api_router.post("/service-requests")
async def create_service_request(payload: ServiceRequestIn, account: dict = Depends(get_current_account)):
    if payload.user_id and account["role"] in ["Super admin", "Admin"]:
        user_id = payload.user_id
    else:
        user_id = account["account_id"]
    
    # Get association_id and unit_id for the user
    res = await db_fetchall('''
        SELECT a.id as assoc_id, u.id as unit_id, a.association_code, a.name as association_name, b.name as block_name, u.unit_number, ud.name, ac.email
        FROM accounts ac
        LEFT JOIN user_details ud ON ac.user_id = ud.user_id
        LEFT JOIN units u ON ud.unit_id = u.id
        LEFT JOIN blocks b ON u.block_id = b.id
        LEFT JOIN associations a ON b.association_id = a.id
        WHERE ac.account_id = %s
    ''', (user_id,))
    
    if not res or not res[0]['assoc_id']:
        raise HTTPException(400, "User not mapped to any association")
    assoc_id = res[0]['assoc_id']
    unit_id = res[0]['unit_id']
    block_name = res[0]['block_name'] or "NA"
    unit_number = res[0]['unit_number'] or "NA"
    association_name = res[0]['association_name'] or ""
    homeowner_email = res[0]['email']
    homeowner_name = res[0]['name']
    
    assoc_initials = "".join([word[0].upper() for word in association_name.split() if word]) or "XX"
    
    count_res = await db_fetchone('SELECT COUNT(id) as cnt FROM service_requests WHERE unit_id = %s', (unit_id,))
    sr_count = (count_res['cnt'] + 1) if count_res else 1
    
    sr_display_id = f"SR-{assoc_initials}-{block_name}-{unit_number}#{sr_count}"
    status = 'In Progress' if (payload.user_id and account["role"] in ["Super admin", "Admin"]) else 'New'

    req_id = await db_execute('''
        INSERT INTO service_requests (user_id, association_id, unit_id, sr_display_id, service_type, sub_category, custom_title, description, image_url, incoming_call_no, status)
        VALUES (%s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s)
    ''', (user_id, assoc_id, unit_id, sr_display_id, payload.service_type, payload.sub_category, payload.custom_title, payload.description, payload.image_url, payload.incoming_call_no, status))

    if payload.user_id and account["role"] in ["Super admin", "Admin"] and homeowner_email:
        mock_send_email(
            to_email=homeowner_email,
            subject="Service Request Created",
            body=f"Hi {homeowner_name},\n\nA new service request ({sr_display_id}) has been created on your behalf by the admin.\nService Type: {payload.service_type}\n\nThank you,\nNestora Community Team"
        )
        await db_execute('''
            INSERT INTO notifications (account_id, title, message)
            VALUES (%s, %s, %s)
        ''', (
            user_id,
            f"Service Request Created: {sr_display_id}",
            f"An admin has created a new {payload.service_type} request on your behalf."
        ))
    elif account["role"] == "Homeowner":
        admin_accounts = await db_fetchall('SELECT admin_id FROM admin_associations WHERE association_id = %s', (assoc_id,))
        for admin in admin_accounts:
            await db_execute('''
                INSERT INTO notifications (account_id, title, message)
                VALUES (%s, %s, %s)
            ''', (
                admin['admin_id'],
                f"New Service Request: {sr_display_id}",
                f"{homeowner_name} created a new {payload.service_type} request for Unit {unit_number}."
            ))
    
    return {"ok": True, "id": req_id}

@api_router.get("/notifications")
async def list_notifications(account: dict = Depends(get_current_account)):
    rows = await db_fetchall('''
        SELECT id, title, message, is_read, created_at 
        FROM notifications 
        WHERE account_id = %s
        ORDER BY created_at DESC LIMIT 50
    ''', (account["account_id"],))
    return {"ok": True, "data": rows}

@api_router.patch("/notifications/{notif_id}/read")
async def mark_notification_read(notif_id: str, account: dict = Depends(get_current_account)):
    await db_execute('UPDATE notifications SET is_read = 1 WHERE id = %s AND account_id = %s', (notif_id, account["account_id"]))
    return {"ok": True}

@api_router.get("/service-requests")
async def list_service_requests(association_id: str | None = None, account: dict = Depends(get_current_account)):
    user_id = account["account_id"]
    role_id = account.get("role_id")
    
    query = '''
        SELECT sr.*, u.unit_number, b.name as block_name, a.name as association_name, ud.name as requestor_name, ud.contact_number as requestor_phone, COALESCE(ac.email, ud.email) as requestor_email
        FROM service_requests sr
        LEFT JOIN units u ON sr.unit_id = u.id
        LEFT JOIN blocks b ON u.block_id = b.id
        LEFT JOIN associations a ON sr.association_id = a.id
        LEFT JOIN accounts ac ON sr.user_id = ac.account_id
        LEFT JOIN user_details ud ON (ac.user_id = ud.user_id OR sr.user_id = ud.user_id)
    '''
    
    # Check if user has admin-level view (can see all for association)
    # Super admin, Admin, Security, Board member, and Committee member can see association-wide requests
    if account["role"] in ["Super admin", "Admin", "Security", "Board member", "Committee member"]:
        if association_id:
            if account["role"] in ["Super admin", "Admin"]:
                query += " WHERE (sr.association_id = %s OR sr.association_id IS NULL) ORDER BY sr.created_at DESC"
            else:
                query += " WHERE sr.association_id = %s ORDER BY sr.created_at DESC"
            rows = await db_fetchall(query, (association_id,))
        else:
            # For Super admin / Admin who might have multiple associations mapped in admin_associations
            res = await db_fetchall("SELECT association_id FROM admin_associations WHERE admin_id = %s", (account["account_id"],))
            if res and account["role"] in ["Super admin", "Admin"]:
                assoc_ids = [r['association_id'] for r in res]
                format_strings = ','.join(['%s'] * len(assoc_ids))
                query += f" WHERE (sr.association_id IN ({format_strings}) OR sr.association_id IS NULL) ORDER BY sr.created_at DESC"
                rows = await db_fetchall(query, tuple(assoc_ids))
            else:
                # If they are a Board/Committee member, they belong to one association via user_details
                if account["role"] in ["Board member", "Committee member", "Security"]:
                    assoc_res = await db_fetchall('''
                        SELECT b.association_id 
                        FROM user_details ud
                        JOIN units u ON ud.unit_id = u.id
                        JOIN blocks b ON u.block_id = b.id
                        WHERE ud.user_id = %s
                    ''', (account.get("user_id"),))
                    if assoc_res:
                        query += " WHERE sr.association_id = %s ORDER BY sr.created_at DESC"
                        rows = await db_fetchall(query, (assoc_res[0]['association_id'],))
                    else:
                        rows = await db_fetchall(query + " ORDER BY sr.created_at DESC")
                else:
                    rows = await db_fetchall(query + " ORDER BY sr.created_at DESC")
    else:
        query += " WHERE sr.user_id = %s ORDER BY sr.created_at DESC"
        rows = await db_fetchall(query, (user_id,))
        
    print("DEBUG GET SRs for", account.get("email"), ":", len(rows), "rows found. Role:", account.get("role"))
    for r in rows:
        print("  - id:", r.get('id'), "status:", r.get('status'))
        
    return {"ok": True, "data": rows}

@api_router.get("/service-requests/{req_id}")
async def get_service_request_by_id(req_id: str, account: dict = Depends(get_current_account)):
    query = '''
        SELECT sr.*, u.unit_number, b.name as block_name, a.name as association_name, ud.name as requestor_name, ud.contact_number as requestor_phone, COALESCE(ac.email, ud.email) as requestor_email
        FROM service_requests sr
        LEFT JOIN units u ON sr.unit_id = u.id
        LEFT JOIN blocks b ON u.block_id = b.id
        LEFT JOIN associations a ON sr.association_id = a.id
        LEFT JOIN accounts ac ON sr.user_id = ac.account_id
        LEFT JOIN user_details ud ON (ac.user_id = ud.user_id OR sr.user_id = ud.user_id)
        WHERE sr.id = %s
    '''
    row = await db_fetchone(query, (req_id,))
    if not row:
        raise HTTPException(404, "Service request not found")
    return {"ok": True, "data": row}

class ServiceRequestStatusIn(BaseModel):
    status: str

@api_router.patch("/service-requests/{req_id}/status")
async def update_service_request_status(req_id: str, payload: ServiceRequestStatusIn, account: dict = Depends(get_current_account)):
    await db_execute("UPDATE service_requests SET status = %s WHERE id = %s", (payload.status, req_id))
        
    res = await db_fetchall("SELECT user_id, sr_display_id FROM service_requests WHERE id = %s", (req_id,))
    if res and account["role"] in ["Super admin", "Admin"] and res[0]['user_id']:
        account_exists = await db_fetchone("SELECT account_id FROM accounts WHERE account_id = %s", (res[0]['user_id'],))
        if account_exists:
            await db_execute('''
                INSERT INTO notifications (account_id, title, message)
                VALUES (%s, %s, %s)
            ''', (
                res[0]['user_id'],
                f"Service Request Updated: {res[0]['sr_display_id']}",
                f"Your service request status was updated to '{payload.status}' by the admin."
            ))
        
    return {"ok": True}
    
class ServiceRequestMappingIn(BaseModel):
    association_id: str
    unit_id: str
    user_id: str

@api_router.patch("/service-requests/{req_id}/map")
async def map_service_request(req_id: str, payload: ServiceRequestMappingIn, account: dict = Depends(get_current_account)):
    if account["role"] not in ["Super admin", "Admin"]:
        raise HTTPException(status_code=403, detail="Admin access required")
    
    # Check if request exists and is unassigned
    res = await db_fetchall("SELECT association_id FROM service_requests WHERE id = %s", (req_id,))
    if not res:
        raise HTTPException(status_code=404, detail="Service request not found")
    
    await db_execute(
        "UPDATE service_requests SET association_id = %s, unit_id = %s, user_id = %s WHERE id = %s",
        (payload.association_id, payload.unit_id, payload.user_id, req_id)
    )
    return {"ok": True, "message": "Service request mapped successfully."}

@api_router.get("/admin/associations/{id}/blocks")
async def get_association_blocks(id: str, _: dict = Depends(require_admin)):
    rows = await db_fetchall("SELECT id, name FROM blocks WHERE association_id = %s ORDER BY name ASC", (id,))
    return {"blocks": rows}

@api_router.get("/associations/{id}/blocks")
async def get_assoc_blocks_public(id: str, account: dict = Depends(get_current_account)):
    if id == "me":
        res = await db_fetchall('''
            SELECT b.association_id 
            FROM accounts ac
            JOIN user_details ud ON ac.user_id = ud.user_id
            JOIN units u ON ud.unit_id = u.id
            JOIN blocks b ON u.block_id = b.id
            WHERE ac.account_id = %s
        ''', (account["account_id"],))
        if res:
            id = res[0]['association_id']
        else:
            return {"blocks": []}
            
    # Any authenticated user can fetch blocks for an association
    rows = await db_fetchall("SELECT id, name FROM blocks WHERE association_id = %s ORDER BY name ASC", (id,))
    return {"blocks": rows}

@api_router.get("/admin/associations/{id}/all-units")
async def get_association_all_units(id: str, _: dict = Depends(require_admin)):
    rows = await db_fetchall("""
        SELECT u.id, u.unit_number, b.name as block_name 
        FROM units u 
        JOIN blocks b ON u.block_id = b.id 
        WHERE b.association_id = %s 
        ORDER BY b.name ASC, u.unit_number ASC
    """, (id,))
    return {"ok": True, "data": rows}

@api_router.get("/admin/blocks/{id}/units")
async def get_block_units(id: str, _: dict = Depends(require_admin)):
    rows = await db_fetchall("SELECT id, unit_number FROM units WHERE block_id = %s ORDER BY unit_number ASC", (id,))
    return {"ok": True, "data": rows}

@api_router.get("/admin/units/{id}/homeowners")
async def get_unit_homeowners(id: str, _: dict = Depends(require_admin)):
    rows = await db_fetchall('''
        SELECT COALESCE(ac.account_id, ud.user_id) as user_id, ud.name, ud.first_name, ud.last_name, ud.email 
        FROM user_details ud 
        LEFT JOIN accounts ac ON ud.user_id = ac.user_id 
        WHERE ud.unit_id = %s
    ''', (id,))
    return {"ok": True, "data": rows}

class ServiceRequestMessageIn(BaseModel):
    message: Optional[str] = None
    attachment_url: Optional[str] = None

@api_router.post("/service-requests/{req_id}/messages")
async def send_message(req_id: str, payload: ServiceRequestMessageIn, account: dict = Depends(get_current_account)):
    sender_id = account["account_id"]
    msg_id = await db_execute('''
        INSERT INTO service_request_messages (service_request_id, sender_id, message, attachment_url)
        VALUES (%s, %s, %s, %s)
    ''', (req_id, sender_id, payload.message, payload.attachment_url))
    
    res = await db_fetchall("SELECT user_id, association_id, sr_display_id FROM service_requests WHERE id = %s", (req_id,))
    if res:
        homeowner_id = res[0]['user_id']
        assoc_id = res[0]['association_id']
        sr_display_id = res[0]['sr_display_id']
        
        if sender_id == homeowner_id:
            admin_accounts = await db_fetchall('SELECT admin_id FROM admin_associations WHERE association_id = %s', (assoc_id,))
            for admin in admin_accounts:
                await db_execute('''
                    INSERT INTO notifications (account_id, title, message)
                    VALUES (%s, %s, %s)
                ''', (
                    admin['admin_id'],
                    f"New Message: {sr_display_id}",
                    f"A homeowner sent a new message on their service request."
                ))
        else:
            await db_execute('''
                INSERT INTO notifications (account_id, title, message)
                VALUES (%s, %s, %s)
            ''', (
                homeowner_id,
                f"New Message: {sr_display_id}",
                f"An admin sent a new message on your service request."
            ))
            
    return {"ok": True, "id": msg_id}

@api_router.get("/service-requests/{req_id}/messages")
async def get_messages(req_id: str, account: dict = Depends(get_current_account)):
    rows = await db_fetchall('''
        SELECT m.*, a.email, a.role_id, ud.name as sender_name
        FROM service_request_messages m
        JOIN accounts a ON m.sender_id = a.account_id
        LEFT JOIN user_details ud ON a.user_id = ud.user_id
        WHERE m.service_request_id = %s
        ORDER BY m.created_at ASC
    ''', (req_id,))
    return {"ok": True, "data": rows}

@api_router.get("/associations/{assoc_id}/board-members")
async def get_association_board_members(assoc_id: str, account: dict = Depends(get_current_account)):
    actual_assoc_id = assoc_id
    if assoc_id == "me":
        res = await db_fetchall('''
            SELECT b.association_id 
            FROM accounts ac
            JOIN user_details ud ON ac.user_id = ud.user_id
            JOIN units u ON ud.unit_id = u.id
            JOIN blocks b ON u.block_id = b.id
            WHERE ac.account_id = %s
        ''', (account["account_id"],))
        if not res:
            return {"ok": True, "data": []}
        actual_assoc_id = res[0]['association_id']

    rows = await db_fetchall('''
        SELECT ac.account_id, ud.name, ac.email, ud.contact_number, ud.profile_pic_url, ud.board_member_since
        FROM accounts ac
        JOIN roles r ON ac.role_id = r.id
        JOIN user_details ud ON ac.user_id = ud.user_id
        JOIN units u ON ud.unit_id = u.id
        JOIN blocks b ON u.block_id = b.id
        WHERE b.association_id = %s AND r.name = 'Board member'
        ORDER BY ud.name ASC
    ''', (actual_assoc_id,))
    
    for r in rows:
        r["board_member_since"] = r["board_member_since"].isoformat() if r.get("board_member_since") else None

    return {"ok": True, "data": rows}

@api_router.get("/associations/{assoc_id}/committee-members")
async def get_association_committee_members(assoc_id: str, account: dict = Depends(get_current_account)):
    actual_assoc_id = assoc_id
    if assoc_id == "me":
        res = await db_fetchall('''
            SELECT b.association_id 
            FROM accounts ac
            JOIN user_details ud ON ac.user_id = ud.user_id
            JOIN units u ON ud.unit_id = u.id
            JOIN blocks b ON u.block_id = b.id
            WHERE ac.account_id = %s
        ''', (account["account_id"],))
        if not res:
            return {"ok": True, "data": []}
        actual_assoc_id = res[0]['association_id']

    query = """
        SELECT 
            cm.id as committee_member_id, cm.committee_id, cm.user_id, cm.start_date as role_start_date, cm.end_date as role_end_date, cm.created_at,
            c.name as committee_name, c.association_id, assoc.name as association_name,
            ud.name, a.email, ud.contact_number as phone, ud.profile_pic_url
        FROM committee_members cm
        JOIN committees c ON cm.committee_id = c.id
        LEFT JOIN associations assoc ON c.association_id = assoc.id
        JOIN accounts a ON cm.user_id = a.user_id
        JOIN user_details ud ON cm.user_id = ud.user_id
        WHERE c.association_id = %s
        ORDER BY c.name ASC, ud.name ASC
    """
    rows = await db_fetchall(query, (actual_assoc_id,))
    
    for r in rows:
        if r.get('role_start_date'): r['role_start_date'] = str(r['role_start_date'])
        if r.get('role_end_date'): r['role_end_date'] = str(r['role_end_date'])
        if r.get('created_at'): r['created_at'] = str(r['created_at'])

    return {"ok": True, "data": rows}

class BoardTaskIn(BaseModel):
    association_id: str | None = None
    title: str
    description: str
    supervised_by: str
    image_url: str | None = None

@api_router.post("/board-tasks")
async def create_board_task(payload: BoardTaskIn, account: dict = Depends(get_current_account)):
    if len(payload.title) > 100:
        raise HTTPException(400, "Title cannot exceed 100 characters")
    if len(payload.description) > 5000:
        raise HTTPException(400, "Description cannot exceed 5000 characters")
        
    assoc_id = payload.association_id
    if not assoc_id or assoc_id == "me":
        if account["role"] == "Board member":
            res = await db_fetchall('''
                SELECT b.association_id 
                FROM accounts ac
                JOIN user_details ud ON ac.user_id = ud.user_id
                JOIN units u ON ud.unit_id = u.id
                JOIN blocks b ON u.block_id = b.id
                WHERE ac.account_id = %s
            ''', (account["account_id"],))
            if not res:
                raise HTTPException(400, "Board member has no association assigned")
            assoc_id = res[0]['association_id']
        else:
            raise HTTPException(400, "association_id is required")
            
    task_id = await db_execute('''
        INSERT INTO board_tasks (association_id, created_by, title, description, supervised_by, image_url)
        VALUES (%s, %s, %s, %s, %s, %s)
    ''', (assoc_id, account["account_id"], payload.title, payload.description, payload.supervised_by, payload.image_url))

    # Send notifications
    board_members = await db_fetchall('''
        SELECT ac.account_id 
        FROM accounts ac
        JOIN roles r ON ac.role_id = r.id
        JOIN user_details ud ON ac.user_id = ud.user_id
        JOIN units u ON ud.unit_id = u.id
        JOIN blocks b ON u.block_id = b.id
        WHERE b.association_id = %s AND r.name = 'Board member'
    ''', (assoc_id,))
    
    admins = await db_fetchall('SELECT admin_id FROM admin_associations WHERE association_id = %s', (assoc_id,))
    
    notify_accounts = set([bm['account_id'] for bm in board_members])
    if account["role"] == "Board member":
        notify_accounts.update([ad['admin_id'] for ad in admins])
        
    notify_accounts.discard(account["account_id"]) # Don't notify self
    
    for acc_id in notify_accounts:
        await db_execute('''
            INSERT INTO notifications (account_id, title, message)
            VALUES (%s, %s, %s)
        ''', (acc_id, "New Board Task Created", f"A new board task '{payload.title}' was created."))

    return {"ok": True}

@api_router.get("/board-tasks")
async def list_board_tasks(association_id: str | None = None, account: dict = Depends(get_current_account)):
    query = '''
        SELECT bt.*, a.name as association_name, 
               COALESCE(ud.name, emp.name, ac.email) as supervised_by_name, 
               COALESCE(ud2.name, emp2.name, ac2.email) as created_by_name
        FROM board_tasks bt
        JOIN associations a ON bt.association_id = a.id
        LEFT JOIN accounts ac ON bt.supervised_by = ac.account_id
        LEFT JOIN user_details ud ON ac.user_id = ud.user_id
        LEFT JOIN employees emp ON ac.employee_id = emp.employee_id
        LEFT JOIN accounts ac2 ON bt.created_by = ac2.account_id
        LEFT JOIN user_details ud2 ON ac2.user_id = ud2.user_id
        LEFT JOIN employees emp2 ON ac2.employee_id = emp2.employee_id
    '''
    
    if account["role"] in ["Super admin", "Admin"]:
        if association_id and association_id != "ALL":
            query += " WHERE bt.association_id = %s ORDER BY bt.created_at DESC"
            rows = await db_fetchall(query, (association_id,))
        else:
            res = await db_fetchall("SELECT association_id FROM admin_associations WHERE admin_id = %s", (account["account_id"],))
            if res:
                assoc_ids = [r['association_id'] for r in res]
                format_strings = ','.join(['%s'] * len(assoc_ids))
                query += f" WHERE bt.association_id IN ({format_strings}) ORDER BY bt.created_at DESC"
                rows = await db_fetchall(query, tuple(assoc_ids))
            else:
                rows = await db_fetchall(query + " ORDER BY bt.created_at DESC")
    elif account["role"] == "Board member":
        res = await db_fetchall('''
            SELECT b.association_id 
            FROM accounts ac
            JOIN user_details ud ON ac.user_id = ud.user_id
            JOIN units u ON ud.unit_id = u.id
            JOIN blocks b ON u.block_id = b.id
            WHERE ac.account_id = %s
        ''', (account["account_id"],))
        if res:
            query += " WHERE bt.association_id = %s ORDER BY bt.created_at DESC"
            rows = await db_fetchall(query, (res[0]['association_id'],))
        else:
            rows = []
    else:
        rows = []
        
    return {"ok": True, "data": rows}

@api_router.get("/board-tasks/{task_id}")
async def get_board_task_by_id(task_id: str, account: dict = Depends(get_current_account)):
    query = '''
        SELECT bt.*, a.name as association_name, 
               COALESCE(ud.name, emp.name, ac.email) as supervised_by_name, 
               COALESCE(ud2.name, emp2.name, ac2.email) as created_by_name
        FROM board_tasks bt
        JOIN associations a ON bt.association_id = a.id
        LEFT JOIN accounts ac ON bt.supervised_by = ac.account_id
        LEFT JOIN user_details ud ON ac.user_id = ud.user_id
        LEFT JOIN employees emp ON ac.employee_id = emp.employee_id
        LEFT JOIN accounts ac2 ON bt.created_by = ac2.account_id
        LEFT JOIN user_details ud2 ON ac2.user_id = ud2.user_id
        LEFT JOIN employees emp2 ON ac2.employee_id = emp2.employee_id
        WHERE bt.id = %s
    '''
    row = await db_fetchone(query, (task_id,))
    if not row:
        raise HTTPException(status_code=404, detail="Board task not found")
    return {"ok": True, "data": row}

class BoardTaskStatusIn(BaseModel):
    status: str

@api_router.patch("/board-tasks/{task_id}/status")
async def update_board_task_status(task_id: str, payload: BoardTaskStatusIn, account: dict = Depends(get_current_account)):
    # Update status
    await db_execute("UPDATE board_tasks SET status = %s WHERE id = %s", (payload.status, task_id))
    
    # Send notification
    res = await db_fetchall("SELECT created_by, supervised_by, title FROM board_tasks WHERE id = %s", (task_id,))
    if res:
        notify_accounts = set()
        notify_accounts.add(res[0]['created_by'])
        notify_accounts.add(res[0]['supervised_by'])
        notify_accounts.discard(account["account_id"]) # Don't notify self
        
        for acc_id in notify_accounts:
            await db_execute('''
                INSERT INTO notifications (account_id, title, message)
                VALUES (%s, %s, %s)
            ''', (acc_id, "Board Task Updated", f"The status of '{res[0]['title']}' was changed to {payload.status}."))
            
    return {"ok": True, "message": f"Task status updated to {payload.status}"}

class BoardTaskMessageIn(BaseModel):
    message: str
    attachment_url: Optional[str] = None

@api_router.post("/board-tasks/{task_id}/messages")
async def send_board_task_message(task_id: str, payload: BoardTaskMessageIn, account: dict = Depends(get_current_account)):
    sender_id = account["account_id"]
    await db_execute('''
        INSERT INTO board_task_messages (board_task_id, sender_id, message, attachment_url)
        VALUES (%s, %s, %s, %s)
    ''', (task_id, sender_id, payload.message, payload.attachment_url))
    
    # Notify board members & admins of the association
    res = await db_fetchall("SELECT association_id, title FROM board_tasks WHERE id = %s", (task_id,))
    if res:
        assoc_id = res[0]['association_id']
        title = res[0]['title']
        
        board_members = await db_fetchall('''
            SELECT ac.account_id 
            FROM accounts ac
            JOIN roles r ON ac.role_id = r.id
            JOIN user_details ud ON ac.user_id = ud.user_id
            JOIN units u ON ud.unit_id = u.id
            JOIN blocks b ON u.block_id = b.id
            WHERE b.association_id = %s AND r.name = 'Board member'
        ''', (assoc_id,))
        
        admins = await db_fetchall('SELECT admin_id FROM admin_associations WHERE association_id = %s', (assoc_id,))
        
        notify_accounts = set([bm['account_id'] for bm in board_members])
        notify_accounts.update([ad['admin_id'] for ad in admins])
        notify_accounts.discard(sender_id) # Don't notify self
        
        for acc_id in notify_accounts:
            await db_execute('''
                INSERT INTO notifications (account_id, title, message)
                VALUES (%s, %s, %s)
            ''', (acc_id, "New message on Board Task", f"New message on '{title}'."))
            
    return {"ok": True}

@api_router.get("/board-tasks/{task_id}/messages")
async def get_board_task_messages(task_id: str, account: dict = Depends(get_current_account)):
    rows = await db_fetchall('''
        SELECT m.*, a.email, a.role_id, r.name as role_name,
               COALESCE(ud.name, emp.name, a.email) as sender_name
        FROM board_task_messages m
        JOIN accounts a ON m.sender_id = a.account_id
        LEFT JOIN roles r ON a.role_id = r.id
        LEFT JOIN user_details ud ON a.user_id = ud.user_id
        LEFT JOIN employees emp ON a.employee_id = emp.employee_id
        WHERE m.board_task_id = %s
        ORDER BY m.created_at ASC
    ''', (task_id,))
    return {"ok": True, "messages": rows}

class MeetingIn(BaseModel):
    association_id: str | None = None
    title: str
    meeting_type: str
    priority: str
    audience: str
    agenda: str
    description: str
    meeting_date: str
    meeting_time: str
    duration: str
    venue: str
    meeting_link: str | None = None
    organizer: str
    attachment_url: str | None = None
    target_block_id: str | None = None

@api_router.post("/meetings")
async def create_meeting(payload: MeetingIn, account: dict = Depends(get_current_account)):
    if len(payload.title) > 100:
        raise HTTPException(400, "Title cannot exceed 100 characters")
    if len(payload.agenda) > 100:
        raise HTTPException(400, "Agenda cannot exceed 100 characters")
    if len(payload.description) > 5000:
        raise HTTPException(400, "Description cannot exceed 5000 characters")
        
    assoc_id = payload.association_id
    if not assoc_id or assoc_id == "me":
        if account["role"] == "Board member":
            res = await db_fetchall('''
                SELECT b.association_id 
                FROM accounts ac
                JOIN user_details ud ON ac.user_id = ud.user_id
                JOIN units u ON ud.unit_id = u.id
                JOIN blocks b ON u.block_id = b.id
                WHERE ac.account_id = %s
            ''', (account["account_id"],))
            if not res:
                raise HTTPException(400, "Board member has no association assigned")
            assoc_id = res[0]['association_id']
        else:
            raise HTTPException(400, "association_id is required")

    link = payload.meeting_link
    if not link:
        import uuid
        mock_id = str(uuid.uuid4())[:10]
        link = f"https://meet.google.com/{mock_id}"
            
    await db_execute('''
        INSERT INTO meetings (association_id, created_by, title, meeting_type, priority, audience, agenda, description, meeting_date, meeting_time, duration, venue, meeting_link, organizer, attachment_url, target_block_id)
        VALUES (%s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s)
    ''', (assoc_id, account["account_id"], payload.title, payload.meeting_type, payload.priority, payload.audience, payload.agenda, payload.description, payload.meeting_date, payload.meeting_time, payload.duration, payload.venue, link, payload.organizer, payload.attachment_url, payload.target_block_id))

    # Send notifications
    role_to_fetch = payload.audience
    if payload.audience == "Board Members":
        role_to_fetch = "Board member"
    elif payload.audience == "Homeowner":
        role_to_fetch = "Homeowner"
        
    query_str = '''
        SELECT ac.account_id 
        FROM accounts ac
        JOIN roles r ON ac.role_id = r.id
        JOIN user_details ud ON ac.user_id = ud.user_id
        JOIN units u ON ud.unit_id = u.id
        JOIN blocks b ON u.block_id = b.id
        WHERE b.association_id = %s AND r.name = %s
    '''
    params = [assoc_id, role_to_fetch]
    
    if payload.target_block_id:
        query_str += " AND b.id = %s"
        params.append(payload.target_block_id)
        
    notify_accounts_raw = await db_fetchall(query_str, tuple(params))
    
    notify_accounts = set([a['account_id'] for a in notify_accounts_raw])
    notify_accounts.discard(account["account_id"]) # Don't notify self
    
    for acc_id in notify_accounts:
        await db_execute('''
            INSERT INTO notifications (account_id, title, message)
            VALUES (%s, %s, %s)
        ''', (acc_id, f"New Meeting: {payload.title}", f"A new {payload.meeting_type} has been scheduled for {payload.meeting_date} at {payload.meeting_time}."))

    return {"ok": True}

@api_router.get("/meetings")
async def list_meetings(association_id: str | None = None, account: dict = Depends(get_current_account)):
    query = '''
        SELECT m.*, a.name as association_name, 
               COALESCE(ud.name, emp.name, ac.email) as organizer_name, 
               COALESCE(ud2.name, emp2.name, ac2.email) as created_by_name,
               (SELECT name FROM blocks WHERE id = m.target_block_id) as target_block_name
        FROM meetings m
        JOIN associations a ON m.association_id = a.id
        LEFT JOIN accounts ac ON m.organizer = ac.account_id
        LEFT JOIN user_details ud ON ac.user_id = ud.user_id
        LEFT JOIN employees emp ON ac.employee_id = emp.employee_id
        LEFT JOIN accounts ac2 ON m.created_by = ac2.account_id
        LEFT JOIN user_details ud2 ON ac2.user_id = ud2.user_id
        LEFT JOIN employees emp2 ON ac2.employee_id = emp2.employee_id
    '''
    
    if account["role"] in ["Super admin", "Admin", "Accountant"]:
        if association_id and association_id != "ALL":
            query += " WHERE m.association_id = %s ORDER BY m.created_at DESC"
            rows = await db_fetchall(query, (association_id,))
        else:
            res = await db_fetchall("SELECT association_id FROM admin_associations WHERE admin_id = %s", (account["account_id"],))
            if res:
                assoc_ids = [r['association_id'] for r in res]
                format_strings = ','.join(['%s'] * len(assoc_ids))
                query += f" WHERE m.association_id IN ({format_strings}) ORDER BY m.created_at DESC"
                rows = await db_fetchall(query, tuple(assoc_ids))
            else:
                rows = await db_fetchall(query + " ORDER BY m.created_at DESC")
    elif account["role"] == "Board member":
        res = await db_fetchall('''
            SELECT b.association_id 
            FROM accounts ac
            JOIN user_details ud ON ac.user_id = ud.user_id
            JOIN units u ON ud.unit_id = u.id
            JOIN blocks b ON u.block_id = b.id
            WHERE ac.account_id = %s
        ''', (account["account_id"],))
        if res:
            query += " WHERE m.association_id = %s ORDER BY m.created_at DESC"
            rows = await db_fetchall(query, (res[0]['association_id'],))
        else:
            rows = []
    elif account["role"] == "Homeowner":
        res = await db_fetchall('''
            SELECT b.association_id, b.id as block_id 
            FROM accounts ac
            JOIN user_details ud ON ac.user_id = ud.user_id
            JOIN units u ON ud.unit_id = u.id
            JOIN blocks b ON u.block_id = b.id
            WHERE ac.account_id = %s
        ''', (account["account_id"],))
        if res:
            query += " WHERE m.association_id = %s AND m.audience = 'Homeowner' AND (m.target_block_id IS NULL OR m.target_block_id = '' OR m.target_block_id = %s) ORDER BY m.created_at DESC"
            rows = await db_fetchall(query, (res[0]['association_id'], res[0]['block_id']))
        else:
            rows = []
    else:
        rows = []
        
    if rows:
        meeting_ids = [row['id'] for row in rows]
        format_strings = ','.join(['%s'] * len(meeting_ids))
        
        my_att_rows = await db_fetchall(f"SELECT meeting_id, status FROM meeting_attendance WHERE account_id = %s AND meeting_id IN ({format_strings})", (account["account_id"], *meeting_ids))
        my_att_map = {r['meeting_id']: r['status'] for r in my_att_rows}
        
        stat_rows = await db_fetchall(f"SELECT meeting_id, status, COUNT(*) as count FROM meeting_attendance WHERE meeting_id IN ({format_strings}) GROUP BY meeting_id, status", tuple(meeting_ids))
        stat_map = {}
        for r in stat_rows:
            mid = r['meeting_id']
            if mid not in stat_map:
                stat_map[mid] = {"yes_count": 0, "no_count": 0, "maybe_count": 0}
            status_lower = r['status'].lower()
            if status_lower == "yes": stat_map[mid]["yes_count"] = r['count']
            elif status_lower == "no": stat_map[mid]["no_count"] = r['count']
            elif status_lower == "maybe": stat_map[mid]["maybe_count"] = r['count']
            
        assoc_ids = list(set([row['association_id'] for row in rows]))
        assoc_format_strings = ','.join(['%s'] * len(assoc_ids))
        
        ho_rows = await db_fetchall(f'''
            SELECT b.association_id, COUNT(DISTINCT ac.account_id) as count
            FROM accounts ac
            JOIN roles r ON ac.role_id = r.id
            JOIN user_details ud ON ac.user_id = ud.user_id
            JOIN units u ON ud.unit_id = u.id
            JOIN blocks b ON u.block_id = b.id
            WHERE r.name = 'Homeowner' AND b.association_id IN ({assoc_format_strings})
            GROUP BY b.association_id
        ''', tuple(assoc_ids))
        ho_map = {r['association_id']: r['count'] for r in ho_rows}
        
        board_rows = await db_fetchall(f'''
            SELECT aa.association_id, COUNT(DISTINCT ac.account_id) as count
            FROM admin_associations aa
            JOIN accounts ac ON aa.admin_id = ac.account_id
            JOIN roles r ON ac.role_id = r.id
            WHERE r.name = 'Board member' AND aa.association_id IN ({assoc_format_strings})
            GROUP BY aa.association_id
        ''', tuple(assoc_ids))
        board_map = {r['association_id']: r['count'] for r in board_rows}

        for row in rows:
            row['my_attendance_status'] = my_att_map.get(row['id'], None)
            
            stats = stat_map.get(row['id'], {"yes_count": 0, "no_count": 0, "maybe_count": 0})
            total_eligible = ho_map.get(row['association_id'], 0) + board_map.get(row['association_id'], 0)
            stats["no_response_count"] = max(0, total_eligible - (stats["yes_count"] + stats["no_count"] + stats["maybe_count"]))
            
            row['attendance_stats'] = stats
            
            mt = row.get("meeting_time")
            if isinstance(mt, timedelta):
                total_seconds = int(mt.total_seconds())
                hours, remainder = divmod(total_seconds, 3600)
                minutes, seconds = divmod(remainder, 60)
                row["meeting_time"] = f"{hours:02d}:{minutes:02d}:{seconds:02d}"
                
            if isinstance(row.get('meeting_date'), (date, datetime)):
                row['meeting_date'] = row['meeting_date'].isoformat()
            if isinstance(row.get('created_at'), (date, datetime)):
                row['created_at'] = row['created_at'].isoformat()
            if isinstance(row.get('updated_at'), (date, datetime)):
                row['updated_at'] = row['updated_at'].isoformat()

    return {"ok": True, "data": rows}

class MeetingStatusUpdate(BaseModel):
    status: str
    
class PollOptionIn(BaseModel):
    text: str

class PollIn(BaseModel):
    question: str
    description: Optional[str] = None
    visibility: Optional[str] = 'All'
    status: Optional[str] = 'Published'
    is_multiple_choice: Optional[bool] = False
    end_date: Optional[str] = None
    association_id: Optional[str] = None
    options: List[PollOptionIn]
    
class PollVoteIn(BaseModel):
    option_ids: List[str]

class MeetingAttendanceIn(BaseModel):
    status: str

@api_router.post("/meetings/{meeting_id}/attendance")
async def update_meeting_attendance(meeting_id: str, payload: MeetingAttendanceIn, account: dict = Depends(get_current_account)):
    if payload.status not in ["Yes", "No", "Maybe"]:
        raise HTTPException(400, "Invalid status")
        
    meeting = await db_fetchone('SELECT meeting_date, meeting_time FROM meetings WHERE id = %s', (meeting_id,))
    if not meeting:
        raise HTTPException(404, "Meeting not found")
        
    mt_dt = datetime.combine(meeting['meeting_date'], (datetime.min + meeting['meeting_time']).time())
    if (mt_dt - datetime.now()).total_seconds() < 24 * 3600:
        raise HTTPException(400, "RSVP is locked within 24 hours of the meeting")
    
    await db_execute('''
        INSERT INTO meeting_attendance (meeting_id, account_id, status)
        VALUES (%s, %s, %s)
        ON DUPLICATE KEY UPDATE status = VALUES(status)
    ''', (meeting_id, account["account_id"], payload.status))
    
    return {"ok": True}

class MeetingDetailsUpdateIn(BaseModel):
    title: str
    meeting_type: str
    priority: str
    audience: str
    agenda: str
    description: str
    meeting_date: str
    meeting_time: str
    duration: str
    venue: str
    organizer: str

@api_router.patch("/meetings/{meeting_id}/details")
async def update_meeting_details(meeting_id: str, payload: MeetingDetailsUpdateIn, account: dict = Depends(get_current_account)):
    # Check if admin
    if account["role"] not in ["Super admin", "Admin"]:
        raise HTTPException(403, "Only admins can edit meetings")
        
    res = await db_fetchall("SELECT created_at FROM meetings WHERE id = %s", (meeting_id,))
    if not res:
        raise HTTPException(404, "Meeting not found")
        
    created_at = res[0]['created_at']
    import datetime
    if (datetime.datetime.now() - created_at).total_seconds() > 2 * 3600:
        raise HTTPException(400, "Edit time limit (2 hours) has expired.")
        
    await db_execute('''
        UPDATE meetings SET 
            title = %s, meeting_type = %s, priority = %s, audience = %s, 
            agenda = %s, description = %s, meeting_date = %s, meeting_time = %s, 
            duration = %s, venue = %s, organizer = %s
        WHERE id = %s
    ''', (payload.title, payload.meeting_type, payload.priority, payload.audience, payload.agenda, payload.description, payload.meeting_date, payload.meeting_time, payload.duration, payload.venue, payload.organizer, meeting_id))
    
    return {"ok": True, "message": "Meeting updated successfully"}

class MeetingMinutesIn(BaseModel):
    meeting_minutes: str
    discussed_topic: str

@api_router.patch("/meetings/{meeting_id}/minutes")
async def add_meeting_minutes(meeting_id: str, payload: MeetingMinutesIn, account: dict = Depends(get_current_account)):
    if account["role"] not in ["Super admin", "Admin"]:
        raise HTTPException(403, "Only admins can add meeting minutes")
        
    await db_execute('''
        UPDATE meetings SET 
            meeting_minutes = %s, discussed_topic = %s, status = 'Completed'
        WHERE id = %s
    ''', (payload.meeting_minutes, payload.discussed_topic, meeting_id))
    
    return {"ok": True, "message": "Meeting minutes added successfully"}

@api_router.get("/timeline")
async def get_timeline(limit: int = 10, offset: int = 0, account: dict = Depends(get_current_account)):
    is_homeowner = account["role"] == "Homeowner"
    assoc_id = None
    
    if is_homeowner:
        res = await db_fetchall('''
            SELECT b.association_id 
            FROM accounts ac
            JOIN user_details ud ON ac.user_id = ud.user_id
            JOIN units u ON ud.unit_id = u.id
            JOIN blocks b ON u.block_id = b.id
            WHERE ac.account_id = %s
        ''', (account["account_id"],))
        if res:
            assoc_id = res[0]['association_id']
    else:
        res = await db_fetchall("SELECT association_id FROM admin_associations WHERE admin_id = %s LIMIT 1", (account["account_id"],))
        if res:
            assoc_id = res[0]['association_id']
            
    if not assoc_id and is_homeowner:
        return {"ok": True, "data": []}
        
    meeting_where = "WHERE m.association_id = %s"
    meeting_params = [assoc_id] if assoc_id else []
    if not assoc_id:
        meeting_where = "WHERE 1=1" 
        
    if is_homeowner:
        meeting_where += " AND m.audience = 'Homeowner'"

    announcement_where = "WHERE a.association_id = %s" if assoc_id else "WHERE 1=1"
    announcement_params = [assoc_id] if assoc_id else []
    
    if is_homeowner:
        announcement_where += " AND a.audience IN ('Homeowner', 'Homeowners')"

    event_where = "WHERE e.status = 'Published'"
    if assoc_id:
        event_where += " AND e.association_id = %s"
        
    event_params = [assoc_id] if assoc_id else []
    
    if is_homeowner:
        event_where += " AND e.audience IN ('All', 'Homeowner', 'Homeowners')"

    poll_where = "WHERE p.status = 'Published'"
    if assoc_id:
        poll_where += " AND p.association_id = %s"
    
    poll_params = [assoc_id] if assoc_id else []
    
    if is_homeowner:
        poll_where += " AND p.visibility IN ('All', 'Homeowner', 'Homeowners')"

    union_query = f"""
        SELECT 'meeting' AS item_type, m.id, m.created_at 
        FROM meetings m {meeting_where}
        
        UNION ALL
        
        SELECT 'announcement' AS item_type, a.id, a.created_at
        FROM announcements a {announcement_where}
        
        UNION ALL
        
        SELECT 'event' AS item_type, e.id, e.created_at
        FROM events e {event_where}
        
        UNION ALL
        
        SELECT 'poll' AS item_type, p.id, p.created_at
        FROM polls p {poll_where}
        
        ORDER BY created_at DESC
        LIMIT %s OFFSET %s
    """
    
    params = tuple(meeting_params + announcement_params + event_params + poll_params + [limit, offset])
    timeline_ids = await db_fetchall(union_query, params)
    
    if not timeline_ids:
        return {"ok": True, "data": []}
        
    m_ids = [r['id'] for r in timeline_ids if r['item_type'] == 'meeting']
    a_ids = [r['id'] for r in timeline_ids if r['item_type'] == 'announcement']
    e_ids = [r['id'] for r in timeline_ids if r['item_type'] == 'event']
    p_ids = [r['id'] for r in timeline_ids if r['item_type'] == 'poll']
    
    timeline_items = []
    
    if m_ids:
        m_format = ','.join(['%s']*len(m_ids))
        m_query = f'''
            SELECT m.*, a.name as association_name, 
                   COALESCE(ud.name, emp.name, ac.email) as organizer_name, 
                   COALESCE(ud2.name, emp2.name, ac2.email) as created_by_name
            FROM meetings m
            JOIN associations a ON m.association_id = a.id
            LEFT JOIN accounts ac ON m.organizer = ac.account_id
            LEFT JOIN user_details ud ON ac.user_id = ud.user_id
            LEFT JOIN employees emp ON ac.employee_id = emp.employee_id
            LEFT JOIN accounts ac2 ON m.created_by = ac2.account_id
            LEFT JOIN user_details ud2 ON ac2.user_id = ud2.user_id
            LEFT JOIN employees emp2 ON ac2.employee_id = emp2.employee_id
            WHERE m.id IN ({m_format})
        '''
        m_rows = await db_fetchall(m_query, tuple(m_ids))
        
        assoc_ids = list(set([r['association_id'] for r in m_rows]))
        if assoc_ids:
            a_fmt = ','.join(['%s'] * len(assoc_ids))
            
            ho_rows = await db_fetchall(f'''
                SELECT b.association_id, COUNT(DISTINCT ac.account_id) as count
                FROM accounts ac
                JOIN roles r ON ac.role_id = r.id
                JOIN user_details ud ON ac.user_id = ud.user_id
                JOIN units u ON ud.unit_id = u.id
                JOIN blocks b ON u.block_id = b.id
                WHERE r.name = 'Homeowner' AND b.association_id IN ({a_fmt})
                GROUP BY b.association_id
            ''', tuple(assoc_ids))
            ho_map = {r['association_id']: r['count'] for r in ho_rows}
            
            board_rows = await db_fetchall(f'''
                SELECT aa.association_id, COUNT(DISTINCT ac.account_id) as count
                FROM admin_associations aa
                JOIN accounts ac ON aa.admin_id = ac.account_id
                JOIN roles r ON ac.role_id = r.id
                WHERE r.name = 'Board member' AND aa.association_id IN ({a_fmt})
                GROUP BY aa.association_id
            ''', tuple(assoc_ids))
            board_map = {r['association_id']: r['count'] for r in board_rows}
            
            my_att_rows = await db_fetchall(f"SELECT meeting_id, status FROM meeting_attendance WHERE account_id = %s AND meeting_id IN ({m_format})", (account["account_id"], *m_ids))
            my_att_map = {r['meeting_id']: r['status'] for r in my_att_rows}
            
            stat_rows = await db_fetchall(f"SELECT meeting_id, status, COUNT(*) as count FROM meeting_attendance WHERE meeting_id IN ({m_format}) GROUP BY meeting_id, status", tuple(m_ids))
            stat_map = {}
            for r in stat_rows:
                mid = r['meeting_id']
                if mid not in stat_map:
                    stat_map[mid] = {"yes_count": 0, "no_count": 0, "maybe_count": 0}
                status_lower = r['status'].lower()
                if status_lower == "yes": stat_map[mid]["yes_count"] = r['count']
                elif status_lower == "no": stat_map[mid]["no_count"] = r['count']
                elif status_lower == "maybe": stat_map[mid]["maybe_count"] = r['count']
                
            for row in m_rows:
                row['item_type'] = 'meeting'
                row['my_attendance_status'] = my_att_map.get(row['id'], None)
                stats = stat_map.get(row['id'], {"yes_count": 0, "no_count": 0, "maybe_count": 0})
                total_eligible = ho_map.get(row['association_id'], 0) + board_map.get(row['association_id'], 0)
                stats["no_response_count"] = max(0, total_eligible - (stats["yes_count"] + stats["no_count"] + stats["maybe_count"]))
                row['attendance_stats'] = stats
                
                mt = row.get("meeting_time")
                if isinstance(mt, timedelta):
                    total_seconds = int(mt.total_seconds())
                    hours, remainder = divmod(total_seconds, 3600)
                    minutes, seconds = divmod(remainder, 60)
                    row["meeting_time"] = f"{hours:02d}:{minutes:02d}:{seconds:02d}"
                timeline_items.append(row)

    if a_ids:
        a_format = ','.join(['%s']*len(a_ids))
        a_rows = await db_fetchall(f"""
            SELECT a.id, a.title, a.body, a.category, a.pinned, a.audience, a.attachment_url, a.created_at, a.updated_at,
                   COALESCE(ud.name, ac.email) AS author_name,
                   (SELECT COUNT(*) FROM announcement_likes WHERE announcement_id = a.id) as like_count,
                   (SELECT COUNT(*) FROM announcement_comments WHERE announcement_id = a.id) as comment_count,
                   EXISTS(SELECT 1 FROM announcement_likes WHERE announcement_id = a.id AND account_id = %s) as user_has_liked
            FROM announcements a
            LEFT JOIN accounts ac ON ac.account_id=a.created_by
            LEFT JOIN user_details ud ON ud.user_id=ac.user_id
            WHERE a.id IN ({a_format})
        """, (account["account_id"], *a_ids))
        for row in a_rows:
            row['user_has_liked'] = bool(row['user_has_liked'])
            sr = _serialize_announcement(row)
            sr['item_type'] = 'announcement'
            timeline_items.append(sr)
            
    if e_ids:
        e_format = ','.join(['%s']*len(e_ids))
        e_rows = await db_fetchall(f"""
            SELECT e.*, COALESCE(ud.name, ac.email) AS author_name,
                   (SELECT COUNT(*) FROM event_likes WHERE event_id = e.id) as like_count,
                   (SELECT COUNT(*) FROM event_comments WHERE event_id = e.id) as comment_count,
                   EXISTS(SELECT 1 FROM event_likes WHERE event_id = e.id AND account_id = %s) as user_has_liked
            FROM events e
            LEFT JOIN accounts ac ON ac.account_id=e.created_by
            LEFT JOIN user_details ud ON ud.user_id=ac.user_id
            WHERE e.id IN ({e_format})
        """, (account["account_id"], *e_ids))
        e_rows = await _attach_rsvp_data(e_rows, account["account_id"])
        for row in e_rows:
            row['user_has_liked'] = bool(row['user_has_liked'])
            sr = _serialize_event(row)
            sr['item_type'] = 'event'
            timeline_items.append(sr)

    if p_ids:
        p_format = ','.join(['%s']*len(p_ids))
        p_query = f'''
            SELECT p.*, a.name as association_name, COALESCE(ud.name, e.name, ac.email) as created_by_name
            FROM polls p
            LEFT JOIN associations a ON p.association_id = a.id
            LEFT JOIN accounts ac ON p.created_by = ac.account_id
            LEFT JOIN user_details ud ON ac.user_id = ud.user_id
            LEFT JOIN employees e ON ac.employee_id = e.employee_id
            WHERE p.id IN ({p_format})
        '''
        p_rows = await db_fetchall(p_query, tuple(p_ids))
        
        for row in p_rows:
            row['item_type'] = 'poll'
            # fetch options
            options = await db_fetchall("SELECT * FROM poll_options WHERE poll_id = %s ORDER BY created_at ASC", (row["id"],))
            row["options"] = options
            
            # fetch votes
            votes = await db_fetchall("SELECT option_id, account_id FROM poll_votes WHERE poll_id = %s", (row["id"],))
            row["total_votes"] = len(set(v["account_id"] for v in votes))
            
            row["my_votes"] = [v["option_id"] for v in votes if v["account_id"] == account["account_id"]]
            
            for opt in options:
                opt_votes = len([v for v in votes if v["option_id"] == opt["id"]])
                opt["vote_count"] = opt_votes
                opt["vote_percentage"] = round((opt_votes / row["total_votes"] * 100) if row["total_votes"] > 0 else 0)
                if opt.get("created_at"):
                    opt["created_at"] = opt["created_at"].isoformat()
            
            # fetch likes
            likes = await db_fetchall("SELECT account_id FROM poll_likes WHERE poll_id = %s", (row["id"],))
            row["like_count"] = len(likes)
            row["user_has_liked"] = any(l["account_id"] == account["account_id"] for l in likes)
            
            # fetch comments
            comments = await db_fetchall("SELECT id FROM poll_comments WHERE poll_id = %s", (row["id"],))
            row["comment_count"] = len(comments)
            
            row["is_multiple_choice"] = bool(row.get("is_multiple_choice"))
            timeline_items.append(row)

    for item in timeline_items:
        for k in ['created_at', 'updated_at', 'meeting_date', 'starts_at', 'ends_at']:
            if k in item and isinstance(item[k], (datetime, date)):
                item[k] = item[k].isoformat()

    timeline_items.sort(key=lambda x: str(x.get('created_at', '')), reverse=True)
    return {"ok": True, "data": timeline_items}

# ---------- Mount ----------


app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],
    allow_credentials=False,
    allow_methods=["*"],
    allow_headers=["*"],
)


async def get_board_member_association(account_id):
    res = await db_fetchone("SELECT association_id FROM board_members WHERE account_id = %s AND status = 'active' ORDER BY created_at DESC LIMIT 1", (account_id,))
    return res["association_id"] if res else None


@api_router.get("/admin/associations/{assoc_id}/homeowners")
async def get_assoc_homeowners(assoc_id: str, account: dict = Depends(require_role(["Super admin", "Admin", "Board member"]))):
    rows = await db_fetchall('''
        SELECT a.account_id, a.user_id, ud.name, a.email, ud.profile_pic_url
        FROM accounts a
        JOIN roles r ON a.role_id = r.id
        JOIN user_details ud ON a.user_id = ud.user_id
        JOIN units u ON ud.unit_id = u.id
        JOIN blocks b ON u.block_id = b.id
        WHERE b.association_id = %s AND r.name = 'Homeowner'
        ORDER BY ud.name ASC
    ''', (assoc_id,))
    return {"ok": True, "data": rows}

class BoardMemberIn(BaseModel):
    association_id: str
    account_id: str
    term_start_date: str
    term_end_date: str

@api_router.post("/admin/board-members")
async def create_board_member(payload: BoardMemberIn, account: dict = Depends(require_admin)):
    import datetime
    start = datetime.datetime.strptime(payload.term_start_date, "%Y-%m-%d").date()
    end = datetime.datetime.strptime(payload.term_end_date, "%Y-%m-%d").date()
    if end < start:
        raise HTTPException(status_code=400, detail="Term end date cannot be before start date.")
        
    role = await db_fetchone("SELECT id FROM roles WHERE name='Board member'")
    
    # insert board member record
    new_id = await db_execute(
        "INSERT INTO board_members (association_id, account_id, term_start_date, term_end_date, created_by) VALUES (%s, %s, %s, %s, %s)",
        (payload.association_id, payload.account_id, start, end, account["account_id"])
    )
    
    # update account role
    await db_execute("UPDATE accounts SET role_id=%s WHERE account_id=%s", (role["id"], payload.account_id))
    
    return {"ok": True, "data": {"id": new_id}}

@api_router.get("/admin/board-members")
async def get_admin_board_members(assoc_id: Optional[str] = None, account: dict = Depends(require_admin)):
    sql = '''
        SELECT bm.id, bm.term_start_date, bm.term_end_date, bm.status, 
               a.account_id, COALESCE(ud.name, a.email) as name, ud.profile_pic_url, ud.contact_number, a.email,
               assoc.name as association_name
        FROM board_members bm
        JOIN accounts a ON bm.account_id = a.account_id
        LEFT JOIN user_details ud ON a.user_id = ud.user_id
        JOIN associations assoc ON bm.association_id = assoc.id
    '''
    params = []
    if assoc_id:
        sql += " WHERE bm.association_id = %s"
        params.append(assoc_id)
        
    sql += " ORDER BY bm.created_at DESC"
    
    rows = await db_fetchall(sql, tuple(params))
    for r in rows:
        r["term_start_date"] = r["term_start_date"].isoformat() if r.get("term_start_date") else None
        r["term_end_date"] = r["term_end_date"].isoformat() if r.get("term_end_date") else None
        
    return {"ok": True, "data": rows}

@api_router.put("/admin/board-members/{bm_id}/end-term")
async def end_board_member_term(bm_id: str, account: dict = Depends(require_admin)):
    bm = await db_fetchone("SELECT account_id FROM board_members WHERE id=%s", (bm_id,))
    if not bm:
        raise HTTPException(status_code=404, detail="Board member not found")
        
    role = await db_fetchone("SELECT id FROM roles WHERE name='Homeowner'")
    
    import datetime
    now = datetime.datetime.now().date()
    
    # Set status to past and term_end_date to today
    await db_execute("UPDATE board_members SET status='past', term_end_date=%s WHERE id=%s", (now, bm_id))
    await db_execute("UPDATE accounts SET role_id=%s WHERE account_id=%s", (role["id"], bm["account_id"]))
    
    return {"ok": True}


class CommitteeMemberIn(BaseModel):
    user_id: str
    start_date: Optional[str] = None
    end_date: Optional[str] = None

class CommitteeMemberAssignIn(BaseModel):
    user_id: str
    committee_id: str
    start_date: Optional[str] = None
    end_date: Optional[str] = None

class CommitteeMemberUpdateIn(BaseModel):
    start_date: Optional[str] = None
    end_date: Optional[str] = None

class CommitteeIn(BaseModel):
    association_id: str
    name: str
    description: Optional[str] = None
    start_date: Optional[str] = None
    end_date: Optional[str] = None
    members: List[CommitteeMemberIn] = []

@api_router.get("/admin/homeowners")
async def get_homeowners(assoc_id: Optional[str] = None, account: dict = Depends(require_role(["Super admin", "Admin", "Board member"]))):
    if account["role"] == "Board member":
        assoc_id = await get_board_member_association(account["account_id"])

    query = '''
        SELECT ud.*, u.unit_number, b.name as block_name, a.name as association_name
        FROM user_details ud
        JOIN units u ON ud.unit_id = u.id
        JOIN blocks b ON u.block_id = b.id
        JOIN associations a ON b.association_id = a.id
        JOIN accounts acc ON ud.user_id = acc.user_id
        JOIN roles r ON acc.role_id = r.id
        WHERE r.name = 'Homeowner'
    '''
    params = ()
    if assoc_id and assoc_id != 'ALL':
        query += " AND a.id = %s"
        params = (assoc_id,)
        
    query += " ORDER BY a.name, b.name, u.unit_number"
    rows = await db_fetchall(query, params)
    for r in rows:
        if r.get('created_at'): r['created_at'] = str(r['created_at'])
        if r.get('updated_at'): r['updated_at'] = str(r['updated_at'])
        if r.get('dob'): r['dob'] = str(r['dob'])
        if r.get('move_in_date'): r['move_in_date'] = str(r['move_in_date'])
    return {"ok": True, "data": rows}

@api_router.get("/admin/committees")
async def get_committees(assoc_id: Optional[str] = None, account: dict = Depends(get_current_account)):
    if account["role"] not in ["Super admin", "Admin", "Board member"]:
        pass # Allow homeowners/tenants if they have it in their UI, we will rely on UI to hide it if no permission, but ideally we should check account.get('role_permissions')
    if account["role"] == "Board member":
        assoc_id = await get_board_member_association(account["account_id"])

    sql = '''
        SELECT c.*, a.name as association_name,
               (SELECT COUNT(*) FROM committee_members cm WHERE cm.committee_id = c.id) as member_count
        FROM committees c
        JOIN associations a ON c.association_id = a.id
    '''
    params = []
    if assoc_id and assoc_id != 'ALL':
        sql += " WHERE c.association_id = %s"
        params.append(assoc_id)
    sql += " ORDER BY c.created_at DESC"
    
    rows = await db_fetchall(sql, tuple(params))
    
    for row in rows:
        members_sql = '''
            SELECT cm.id, cm.user_id, cm.start_date, cm.end_date, ud.name, ud.profile_pic_url, a.email
            FROM committee_members cm
            JOIN user_details ud ON cm.user_id = ud.user_id
            JOIN accounts a ON ud.user_id = a.user_id
            WHERE cm.committee_id = %s
        '''
        members = await db_fetchall(members_sql, (row['id'],))
        unique_members = {m['user_id']: m for m in members}.values()
        row['members'] = list(unique_members)
        if row.get('start_date'): row['start_date'] = str(row['start_date'])
        if row.get('end_date'): row['end_date'] = str(row['end_date'])
        if row.get('created_at'): row['created_at'] = str(row['created_at'])
        for m in row['members']:
            if m.get('start_date'): m['start_date'] = str(m['start_date'])
            if m.get('end_date'): m['end_date'] = str(m['end_date'])
            
    return {"ok": True, "data": rows}

@api_router.post("/admin/committees")
async def create_committee(payload: dict, account: dict = Depends(get_current_account)):
    if account["role"] not in ["Super admin", "Admin", "Board member"]:
        if not account.get("role_permissions", {}).get("Committees", {}).get("can_create"):
            raise HTTPException(403, "Permission denied")
    
    assoc_id = payload.get("association_id")
    if account["role"] == "Board member":
        assoc_id = await get_board_member_association(account["account_id"])

    start = payload.get("start_date")
    end = payload.get("end_date")
    
    committee_id = await db_execute(
        "INSERT INTO committees (association_id, name, description, start_date, end_date) VALUES (%s, %s, %s, %s, %s)",
        (assoc_id, payload.get("name"), payload.get("description"), start, end)
    )
    
    cm_role = await db_fetchone("SELECT id FROM roles WHERE name='Committee member'")
    
    for member in payload.get("members", []):
        m_start = member.get("start_date")
        m_end = member.get("end_date")
        await db_execute(
            "INSERT INTO committee_members (committee_id, user_id, start_date, end_date) VALUES (%s, %s, %s, %s)",
            (committee_id, member.get("user_id"), m_start, m_end)
        )
        if cm_role:
            await db_execute("UPDATE accounts SET role_id=%s WHERE user_id=%s", (cm_role["id"], member.get("user_id")))
            
    return {"ok": True, "data": {"id": committee_id}}

@api_router.put("/admin/committees/{committee_id}")
async def update_committee(committee_id: str, payload: dict, account: dict = Depends(get_current_account)):
    if account["role"] not in ["Super admin", "Admin", "Board member"]:
        if not account.get("role_permissions", {}).get("Committees", {}).get("can_update"):
            raise HTTPException(403, "Permission denied")
    
    assoc_id = payload.get("association_id")
    if account["role"] == "Board member":
        assoc_id = await get_board_member_association(account["account_id"])
    
    start = payload.get("start_date")
    end = payload.get("end_date")
    
    await db_execute(
        "UPDATE committees SET association_id=%s, name=%s, description=%s, start_date=%s, end_date=%s WHERE id=%s",
        (assoc_id, payload.get("name"), payload.get("description"), start, end, committee_id)
    )
    members_list = payload.get("members")
    if members_list is not None:
        ho_role = await db_fetchone("SELECT id FROM roles WHERE name='Homeowner'")
        cm_role = await db_fetchone("SELECT id FROM roles WHERE name='Committee member'")
        
        current_members_rows = await db_fetchall("SELECT user_id FROM committee_members WHERE committee_id=%s", (committee_id,))
        old_user_ids = {r["user_id"] for r in current_members_rows}
        new_user_ids = {m.get("user_id") for m in members_list if m.get("user_id")}
        
        for old_user in old_user_ids:
            if old_user not in new_user_ids:
                await db_execute("DELETE FROM committee_members WHERE committee_id=%s AND user_id=%s", (committee_id, old_user))
                if ho_role:
                    await db_execute("UPDATE accounts SET role_id=%s WHERE user_id=%s", (ho_role["id"], old_user))
                    
        for member in members_list:
            u_id = member.get("user_id")
            if not u_id:
                continue
            m_start = member.get("start_date")
            m_end = member.get("end_date")
            if u_id in old_user_ids:
                await db_execute(
                    "UPDATE committee_members SET start_date=%s, end_date=%s WHERE committee_id=%s AND user_id=%s",
                    (m_start, m_end, committee_id, u_id)
                )
            else:
                await db_execute(
                    "INSERT INTO committee_members (committee_id, user_id, start_date, end_date) VALUES (%s, %s, %s, %s)",
                    (committee_id, u_id, m_start, m_end)
                )
                if cm_role:
                    await db_execute("UPDATE accounts SET role_id=%s WHERE user_id=%s", (cm_role["id"], u_id))
                    
    return {"ok": True}

@api_router.delete("/admin/committees/{committee_id}")
async def delete_committee(committee_id: str, account: dict = Depends(get_current_account)):
    if account["role"] not in ["Super admin", "Admin", "Board member"]:
        if not account.get("role_permissions", {}).get("Committees", {}).get("can_delete"):
            raise HTTPException(403, "Permission denied")
    ho_role = await db_fetchone("SELECT id FROM roles WHERE name='Homeowner'")
    old_members_rows = await db_fetchall("SELECT user_id FROM committee_members WHERE committee_id=%s", (committee_id,))
    
    if ho_role:
        for r in old_members_rows:
            await db_execute("UPDATE accounts SET role_id=%s WHERE user_id=%s", (ho_role["id"], r['user_id']))
            
    await db_execute("DELETE FROM committees WHERE id=%s", (committee_id,))
    return {"ok": True}

@api_router.get("/admin/committee-members")
async def get_committee_members(assoc_id: Optional[str] = None, account: dict = Depends(require_role(["Super admin", "Admin", "Board member"]))):
    if account["role"] == "Board member":
        assoc_id = await get_board_member_association(account["account_id"])

    query = """
        SELECT 
            cm.id as committee_member_id, cm.committee_id, cm.user_id, cm.start_date as role_start_date, cm.end_date as role_end_date, cm.created_at,
            c.name as committee_name, c.association_id, assoc.name as association_name,
            ud.name, a.email, ud.contact_number as phone, ud.profile_pic_url
        FROM committee_members cm
        JOIN committees c ON cm.committee_id = c.id
        LEFT JOIN associations assoc ON c.association_id = assoc.id
        JOIN accounts a ON cm.user_id = a.user_id
        JOIN user_details ud ON cm.user_id = ud.user_id
    """
    params = ()
    if assoc_id and assoc_id != "ALL":
        query += " WHERE c.association_id = %s"
        params = (assoc_id,)
        
    query += " ORDER BY c.name ASC, ud.name ASC"
    
    rows = await db_fetchall(query, params)
    for r in rows:
        if r.get('role_start_date'): r['role_start_date'] = str(r['role_start_date'])
        if r.get('role_end_date'): r['role_end_date'] = str(r['role_end_date'])
        if r.get('created_at'): r['created_at'] = str(r['created_at'])
    return {"ok": True, "data": rows}

@api_router.post("/admin/committee-members")
async def assign_committee_member(payload: CommitteeMemberAssignIn, account: dict = Depends(require_role(["Super admin", "Admin", "Board member"]))):
    start = payload.start_date if payload.start_date else None
    end = payload.end_date if payload.end_date else None
    
    existing = await db_fetchone(
        "SELECT * FROM committee_members WHERE committee_id=%s AND user_id=%s",
        (payload.committee_id, payload.user_id)
    )
    if existing:
        raise HTTPException(400, "User is already a member of this committee.")
        
    await db_execute(
        "INSERT INTO committee_members (committee_id, user_id, start_date, end_date) VALUES (%s, %s, %s, %s)",
        (payload.committee_id, payload.user_id, start, end)
    )
    
    cm_role = await db_fetchone("SELECT id FROM roles WHERE name='Committee member'")
    if cm_role:
        await db_execute("UPDATE accounts SET role_id=%s WHERE user_id=%s", (cm_role["id"], payload.user_id))
        
    return {"ok": True}

@api_router.put("/admin/committee-members/{committee_id}/{user_id}")
async def update_committee_member(committee_id: str, user_id: str, payload: CommitteeMemberUpdateIn, account: dict = Depends(require_role(["Super admin", "Admin", "Board member"]))):
    start = payload.start_date if payload.start_date else None
    end = payload.end_date if payload.end_date else None
    
    await db_execute(
        "UPDATE committee_members SET start_date=%s, end_date=%s WHERE committee_id=%s AND user_id=%s",
        (start, end, committee_id, user_id)
    )
    return {"ok": True}

@api_router.delete("/admin/committee-members/{committee_id}/{user_id}")
async def delete_committee_member(committee_id: str, user_id: str, account: dict = Depends(require_role(["Super admin", "Admin", "Board member"]))):
    await db_execute("DELETE FROM committee_members WHERE committee_id=%s AND user_id=%s", (committee_id, user_id))
    # Check if user has other active committee memberships
    other_cm = await db_fetchone("SELECT id FROM committee_members WHERE user_id=%s", (user_id,))
    if not other_cm:
        homeowner_role = await db_fetchone("SELECT id FROM roles WHERE name='Homeowner'")
        if homeowner_role:
            await db_execute(
                "UPDATE accounts SET role_id=%s WHERE user_id=%s AND role_id=(SELECT id FROM roles WHERE name='Committee member')",
                (homeowner_role["id"], user_id)
            )
    return {"ok": True}

@api_router.delete("/admin/committee-members/{member_id}")
async def delete_committee_member_by_id(member_id: str, account: dict = Depends(require_role(["Super admin", "Admin", "Board member"]))):
    cm = await db_fetchone("SELECT user_id, committee_id FROM committee_members WHERE id=%s", (member_id,))
    if cm:
        await db_execute("DELETE FROM committee_members WHERE id=%s", (member_id,))
        other_cm = await db_fetchone("SELECT id FROM committee_members WHERE user_id=%s", (cm["user_id"],))
        if not other_cm:
            homeowner_role = await db_fetchone("SELECT id FROM roles WHERE name='Homeowner'")
            if homeowner_role:
                await db_execute(
                    "UPDATE accounts SET role_id=%s WHERE user_id=%s AND role_id=(SELECT id FROM roles WHERE name='Committee member')",
                    (homeowner_role["id"], cm["user_id"])
                )
    return {"ok": True}


@api_router.get("/admin/subscription-plans")
async def get_subscription_plans(account: dict = Depends(require_role(["Super admin"]))):
    rows = await db_fetchall("SELECT * FROM subscription_plans ORDER BY created_at DESC")
    plan_features = await db_fetchall("SELECT plan_id, feature_id FROM subscription_plan_features")
    feature_map = {}
    for pf in plan_features:
        feature_map.setdefault(pf['plan_id'], []).append(pf['feature_id'])
    
    for r in rows:
        if r.get('created_at'): r['created_at'] = str(r['created_at'])
        r['monthly_price'] = float(r['monthly_price']) if r['monthly_price'] is not None else None
        r['yearly_price'] = float(r['yearly_price']) if r['yearly_price'] is not None else None
        r['features'] = feature_map.get(r['id'], [])
        
    return {"ok": True, "data": rows}

@api_router.post("/admin/subscription-plans")
async def create_subscription_plan(payload: SubscriptionPlanCreate, account: dict = Depends(require_role(["Super admin"]))):
    existing = await db_fetchone("SELECT id FROM subscription_plans WHERE name=%s AND country=%s", (payload.name, payload.country))
    if existing:
        raise HTTPException(status_code=400, detail="A plan with this name already exists in this country")
        
    new_id = await db_execute(
        "INSERT INTO subscription_plans (name, code, country, description, monthly_price, yearly_price, trial_days, is_active) VALUES (%s, %s, %s, %s, %s, %s, %s, %s)",
        (payload.name, payload.code, payload.country, payload.description, payload.monthly_price, payload.yearly_price, payload.trial_days, payload.is_active)
    )
    return {"ok": True, "id": new_id}

@api_router.put("/admin/subscription-plans/{plan_id}")
async def update_subscription_plan(plan_id: str, payload: SubscriptionPlanUpdate, account: dict = Depends(require_role(["Super admin"]))):
    existing = await db_fetchone("SELECT id FROM subscription_plans WHERE name=%s AND country=%s AND id != %s", (payload.name, payload.country, plan_id))
    if existing:
        raise HTTPException(status_code=400, detail="A plan with this name already exists in this country")
        
    await db_execute(
        "UPDATE subscription_plans SET name=%s, code=%s, country=%s, description=%s, monthly_price=%s, yearly_price=%s, trial_days=%s, is_active=%s WHERE id=%s",
        (payload.name, payload.code, payload.country, payload.description, payload.monthly_price, payload.yearly_price, payload.trial_days, payload.is_active, plan_id)
    )
    return {"ok": True}

@api_router.delete("/admin/subscription-plans/{plan_id}")
async def delete_subscription_plan(plan_id: str, account: dict = Depends(require_role(["Super admin"]))):
    await db_execute("DELETE FROM subscription_plans WHERE id=%s", (plan_id,))
    return {"ok": True}

@api_router.post("/admin/subscription-plans/{plan_id}/features")
async def set_subscription_plan_features(plan_id: str, payload: SubscriptionPlanFeatures, account: dict = Depends(require_role(["Super admin"]))):
    await db_execute("DELETE FROM subscription_plan_features WHERE plan_id=%s", (plan_id,))
    for fid in payload.feature_ids:
        await db_execute(
            "INSERT INTO subscription_plan_features (plan_id, feature_id) VALUES (%s, %s)",
            (plan_id, fid)
        )
    return {"ok": True}

@api_router.put("/admin/associations/{assoc_id}/subscription")
async def update_association_subscription(assoc_id: str, payload: AssociationSubscriptionUpdate, account: dict = Depends(require_role(["Super admin"]))):
    s_start = payload.subscription_start if payload.subscription_start else None
    s_end = payload.subscription_end if payload.subscription_end else None
    r_date = payload.renewal_date if payload.renewal_date else None
    
    await db_execute(
        "UPDATE associations SET current_plan_id=%s, subscription_start=%s, subscription_end=%s, payment_status=%s, renewal_date=%s, subscription_status=%s WHERE id=%s",
        (payload.plan_id, s_start, s_end, payload.payment_status, r_date, payload.subscription_status, assoc_id)
    )
    return {"ok": True}





class VisitorRequestIn(BaseModel):
    mobile: str
    name: str
    visitor_type: str
    number_of_visitors: int
    vehicle_number: Optional[str] = None
    photo_url: Optional[str] = None
    id_type: Optional[str] = None
    id_number: Optional[str] = None
    unit_id: str
    purpose: str
    expected_duration: Optional[str] = None
    notes: Optional[str] = None


class PreApprovedVisitorIn(BaseModel):
    visitor_name: str
    mobile: str
    visitor_type: str
    visit_date: str
    start_time: str
    end_time: str
    number_of_visitors: int = 1
    vehicle_number: Optional[str] = None
    purpose: Optional[str] = None
    pass_type: str = "Single Entry"

class CheckInIn(BaseModel):
    guard_id: Optional[str] = None
    gate: Optional[str] = None
    visitor_photo_url: Optional[str] = None
    remarks: Optional[str] = None


class DeliveryIn(BaseModel):
    unit_id: str
    delivery_type: str
    company_name: Optional[str] = None
    delivery_person_name: Optional[str] = None
    mobile: Optional[str] = None
    status: str
    gate: Optional[str] = None
    package_photo_url: Optional[str] = None


class StaffCheckIn(BaseModel):
    gate: Optional[str] = "Main Gate"
    remarks: Optional[str] = None

class VisitorCheckIn(BaseModel):
    pass_code: Optional[str] = None

# ==========================================
# VISITOR MANAGEMENT (SECURITY & RESIDENT)
# ==========================================


@api_router.get("/security/visitors/search")
async def search_visitor_by_mobile(mobile: str, account: dict = Depends(get_current_account)):
    # Quick entry search
    visitor = await db_fetchone("SELECT * FROM visitors WHERE mobile = %s", (mobile,))
    if not visitor:
        raise HTTPException(status_code=404, detail="Visitor not found")
    return visitor

@api_router.get("/security/residents/search")
async def search_residents(q: str, account: dict = Depends(get_current_account)):
    query = f"%{q}%"
    rows = await db_fetchall(
        """
        SELECT ud.user_id, ud.name, ud.contact_number, un.id as unit_id, un.unit_number, b.name as block_name
        FROM user_details ud
        JOIN units un ON ud.unit_id = un.id
        JOIN blocks b ON un.block_id = b.id
        WHERE ud.name LIKE %s OR ud.contact_number LIKE %s OR un.unit_number LIKE %s
        LIMIT 20
        """,
        (query, query, query)
    )
    return {"residents": rows}

@api_router.post("/security/visitor/request")
async def request_visitor_entry(payload: VisitorRequestIn, account: dict = Depends(get_current_account)):
    import uuid
    
    # 1. Ensure visitor exists
    visitor = await db_fetchone("SELECT id FROM visitors WHERE mobile = %s", (payload.mobile,))
    if not visitor:
        visitor_id = str(uuid.uuid4())
        await db_execute(
            "INSERT INTO visitors (id, name, mobile, photo_url, id_type, id_number) VALUES (%s, %s, %s, %s, %s, %s)",
            (visitor_id, payload.name, payload.mobile, payload.photo_url, payload.id_type, payload.id_number)
        )
    else:
        visitor_id = visitor["id"]
        # Update details if provided
        await db_execute(
            "UPDATE visitors SET name=%s, photo_url=COALESCE(%s, photo_url) WHERE id=%s",
            (payload.name, payload.photo_url, visitor_id)
        )

    # 2. Create Visit Request
    visit_id = str(uuid.uuid4())
    await db_execute(
        """
        INSERT INTO visitor_visits 
        (id, visitor_id, unit_id, purpose, visitor_type, number_of_visitors, vehicle_number, notes, expected_duration, status) 
        VALUES (%s, %s, %s, %s, %s, %s, %s, %s, %s, 'Pending')
        """,
        (visit_id, visitor_id, payload.unit_id, payload.purpose, payload.visitor_type, 
         payload.number_of_visitors, payload.vehicle_number, payload.notes, payload.expected_duration)
    )

    # 3. Notify the Resident(s) of the unit
    residents = await db_fetchall("SELECT a.id FROM accounts a JOIN user_details ud ON a.user_id = ud.user_id WHERE ud.unit_id = %s", (payload.unit_id,))
    notif_title = "New Visitor Request"
    notif_msg = f"Visitor '{payload.name}' wants to visit your flat for {payload.purpose}."
    for r in residents:
        await db_execute(
            "INSERT INTO notifications (account_id, title, message, visit_id) VALUES (%s, %s, %s, %s)",
            (r["id"], notif_title, notif_msg, visit_id)
        )
    
    return {"ok": True, "visit_id": visit_id, "message": "Request sent to resident."}

@api_router.get("/security/visitor/status/{visit_id}")
async def get_visitor_status(visit_id: str, account: dict = Depends(get_current_account)):
    visit = await db_fetchone("SELECT status, pass_code FROM visitor_visits WHERE id = %s", (visit_id,))
    if not visit:
        raise HTTPException(404, "Visit not found")
    return visit

@api_router.get("/resident/visitor/pending")
async def get_pending_visitor_requests(account: dict = Depends(get_current_account)):
    # Get requests for resident's unit
    user_detail = await db_fetchone("SELECT unit_id FROM user_details WHERE user_id = %s", (account["user_id"],))
    if not user_detail or not user_detail["unit_id"]:
        return {"requests": []}
        
    rows = await db_fetchall(
        """
        SELECT v.id as visit_id, vis.name, vis.mobile, vis.photo_url, v.purpose, v.visitor_type, v.number_of_visitors, v.created_at
        FROM visitor_visits v
        JOIN visitors vis ON v.visitor_id = vis.id
        WHERE v.unit_id = %s AND v.status = 'Pending'
        ORDER BY v.created_at DESC
        """,
        (user_detail["unit_id"],)
    )
    return {"requests": rows}

@api_router.post("/resident/visitor/{visit_id}/approve")
async def approve_visitor(visit_id: str, account: dict = Depends(get_current_account)):
    visit = await db_fetchone("SELECT * FROM visitor_visits WHERE id = %s AND status = 'Pending'", (visit_id,))
    if not visit:
        raise HTTPException(404, "Visit not found or already processed")
    
    # Check if resident has access
    user_detail = await db_fetchone("SELECT unit_id FROM user_details WHERE user_id = %s", (account["user_id"],))
    if visit["unit_id"] != user_detail["unit_id"]:
        raise HTTPException(403, "Not authorized to approve this visit")
        
    pass_code = f"PASS-{generate_code(4)}"
    await db_execute(
        "UPDATE visitor_visits SET status = 'Approved', pass_code = %s WHERE id = %s",
        (pass_code, visit_id)
    )
    # Mark notification as read
    await db_execute("UPDATE notifications SET is_read = 1 WHERE visit_id = %s AND account_id = %s", (visit_id, account["account_id"]))
    
    return {"ok": True, "pass_code": pass_code}

@api_router.post("/resident/visitor/{visit_id}/reject")
async def reject_visitor(visit_id: str, account: dict = Depends(get_current_account)):
    visit = await db_fetchone("SELECT * FROM visitor_visits WHERE id = %s AND status = 'Pending'", (visit_id,))
    if not visit:
        raise HTTPException(404, "Visit not found or already processed")
    
    user_detail = await db_fetchone("SELECT unit_id FROM user_details WHERE user_id = %s", (account["user_id"],))
    if visit["unit_id"] != user_detail["unit_id"]:
        raise HTTPException(403, "Not authorized to reject this visit")
        
    await db_execute("UPDATE visitor_visits SET status = 'Denied' WHERE id = %s", (visit_id,))
    await db_execute("UPDATE notifications SET is_read = 1 WHERE visit_id = %s AND account_id = %s", (visit_id, account["account_id"]))
    
    return {"ok": True}


# ==========================================
# PRE-APPROVED VISITORS
# ==========================================

@api_router.post("/resident/preapproved-visitors")
async def create_preapproved_visitor(payload: PreApprovedVisitorIn, account: dict = Depends(get_current_account)):
    import uuid
    import random
    import string
    
    # Get resident's unit
    user_detail = await db_fetchone("SELECT user_id, unit_id FROM user_details WHERE user_id = %s", (account["user_id"],))
    if not user_detail or not user_detail["unit_id"]:
        raise HTTPException(403, "Resident details not found or not assigned to a unit.")
        
    pass_code = "PA-" + "".join(random.choices(string.ascii_uppercase + string.digits, k=6))
    otp = "".join(random.choices(string.digits, k=6))
    
    pass_id = await db_execute(
        """
        INSERT INTO pre_approved_visitors 
        (resident_id, unit_id, visitor_name, mobile, visitor_type, pass_code, otp, visit_date, start_time, end_time, number_of_visitors, vehicle_number, purpose, pass_type, status, created_by)
        VALUES (%s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, 'Active', %s)
        """,
        (user_detail["user_id"], user_detail["unit_id"], payload.visitor_name, payload.mobile, payload.visitor_type, pass_code, otp, payload.visit_date, payload.start_time, payload.end_time, payload.number_of_visitors, payload.vehicle_number, payload.purpose, payload.pass_type, account["account_id"])
    )
    
    return {"ok": True, "pass_code": pass_code, "otp": otp, "id": pass_id}

@api_router.get("/resident/preapproved-visitors")
async def get_resident_preapproved(account: dict = Depends(get_current_account)):
    user_detail = await db_fetchone("SELECT user_id, unit_id FROM user_details WHERE user_id = %s", (account["user_id"],))
    if not user_detail or not user_detail["unit_id"]:
        return {"visitors": []}
        
    rows = await db_fetchall(
        """
        SELECT p.*, un.unit_number, b.name as block_name,
        (SELECT check_in FROM visitor_logs WHERE pre_approved_id = p.id ORDER BY check_in DESC LIMIT 1) as last_check_in
        FROM pre_approved_visitors p
        LEFT JOIN units un ON p.unit_id = un.id
        LEFT JOIN blocks b ON un.block_id = b.id
        WHERE p.unit_id = %s
        ORDER BY p.visit_date DESC, p.start_time DESC
        """,
        (user_detail["unit_id"],)
    )
    for r in rows:
        if r.get("start_time") is not None: r["start_time"] = str(r["start_time"])
        if r.get("end_time") is not None: r["end_time"] = str(r["end_time"])
    return {"visitors": rows}

@api_router.get("/public/visitor-passes/{pass_code}")
async def get_public_visitor_pass(pass_code: str):
    """Return the minimum pass details needed by a shared visitor-pass link."""
    row = await db_fetchone(
        """
        SELECT id, visitor_name, mobile, visitor_type, pass_code, otp,
               visit_date, start_time, end_time, number_of_visitors,
               vehicle_number, purpose, pass_type, status
        FROM pre_approved_visitors
        WHERE pass_code = %s
        LIMIT 1
        """,
        (pass_code,)
    )
    if not row:
        raise HTTPException(status_code=404, detail="Visitor pass not found")

    if row.get("visit_date") is not None:
        row["visit_date"] = str(row["visit_date"])
    if row.get("start_time") is not None:
        row["start_time"] = str(row["start_time"])
    if row.get("end_time") is not None:
        row["end_time"] = str(row["end_time"])
    return row

@api_router.delete("/resident/preapproved-visitors/{pass_id}")
async def cancel_preapproved_visitor(pass_id: str, account: dict = Depends(get_current_account)):
    user_detail = await db_fetchone("SELECT unit_id FROM user_details WHERE user_id = %s", (account["user_id"],))
    if not user_detail or not user_detail["unit_id"]:
        raise HTTPException(status_code=403, detail="Unauthorized")

    await db_execute("UPDATE pre_approved_visitors SET status = 'Cancelled' WHERE id = %s AND unit_id = %s", (pass_id, user_detail["unit_id"]))
    return {"ok": True}

@api_router.get("/security/preapproved-visitors/today")
async def get_today_preapproved(account: dict = Depends(get_current_account)):
    # Assuming today's date context. Using CURDATE() in SQL.
    rows = await db_fetchall(
        """
        SELECT p.*,
        (SELECT id FROM visitor_logs WHERE pre_approved_id = p.id AND check_out IS NULL LIMIT 1) as active_log_id,
        ud.name as resident_name, un.unit_number, b.name as block_name
        FROM pre_approved_visitors p
        JOIN user_details ud ON p.resident_id = ud.user_id
        JOIN units un ON p.unit_id = un.id
        JOIN blocks b ON un.block_id = b.id
        WHERE p.visit_date = CURDATE() AND p.status IN ('Active', 'Used')
        ORDER BY p.start_time ASC
        """
    )
    for r in rows:
        if r.get("start_time") is not None: r["start_time"] = str(r["start_time"])
        if r.get("end_time") is not None: r["end_time"] = str(r["end_time"])
    return {"visitors": rows}

@api_router.get("/security/preapproved-visitors/search")
async def search_preapproved(q: str, account: dict = Depends(get_current_account)):
    query = f"%{q}%"
    rows = await db_fetchall(
        """
        SELECT p.*,
        (SELECT id FROM visitor_logs WHERE pre_approved_id = p.id AND check_out IS NULL LIMIT 1) as active_log_id,
        ud.name as resident_name, un.unit_number, b.name as block_name
        FROM pre_approved_visitors p
        JOIN user_details ud ON p.resident_id = ud.user_id
        JOIN units un ON p.unit_id = un.id
        JOIN blocks b ON un.block_id = b.id
        WHERE p.mobile LIKE %s OR p.pass_code LIKE %s OR p.otp = %s
        ORDER BY p.visit_date DESC
        LIMIT 10
        """,
        (query, query, q)
    )
    for r in rows:
        if r.get("start_time") is not None: r["start_time"] = str(r["start_time"])
        if r.get("end_time") is not None: r["end_time"] = str(r["end_time"])
    return {"visitors": rows}

@api_router.post("/security/preapproved-visitors/{pass_id}/check-in")
async def check_in_preapproved(pass_id: str, payload: CheckInIn, account: dict = Depends(get_current_account)):
    import uuid
    pass_obj = await db_fetchone("SELECT * FROM pre_approved_visitors WHERE id = %s", (pass_id,))
    if not pass_obj:
        raise HTTPException(404, "Pass not found")
        
    if pass_obj["status"] not in ("Active", "Used"):
        raise HTTPException(400, "Pass is not valid for check-in.")
        
    log_id = await db_execute(
        """
        INSERT INTO visitor_logs (pre_approved_id, gate, guard_id, visitor_photo_url, remarks)
        VALUES (%s, %s, %s, %s, %s)
        """,
        (pass_id, payload.gate, payload.guard_id, payload.visitor_photo_url, payload.remarks)
    )
    
    if pass_obj["pass_type"] == "Single Entry":
        await db_execute("UPDATE pre_approved_visitors SET status = 'Used' WHERE id = %s", (pass_id,))
        
    return {"ok": True, "log_id": log_id}

@api_router.post("/security/preapproved-visitors/{pass_id}/check-out")
async def check_out_preapproved(pass_id: str, account: dict = Depends(get_current_account)):
    # Find active log
    log = await db_fetchone("SELECT id FROM visitor_logs WHERE pre_approved_id = %s AND check_out IS NULL ORDER BY check_in DESC LIMIT 1", (pass_id,))
    if not log:
        raise HTTPException(404, "No active check-in found for this pass")
        
    await db_execute("UPDATE visitor_logs SET check_out = CURRENT_TIMESTAMP WHERE id = %s", (log["id"],))
    
    pass_obj = await db_fetchone("SELECT pass_type FROM pre_approved_visitors WHERE id = %s", (pass_id,))
    if pass_obj and pass_obj["pass_type"] == "Single Entry":
        await db_execute("UPDATE pre_approved_visitors SET status = 'Expired' WHERE id = %s", (pass_id,))
        
    return {"ok": True}


@api_router.post("/security/visitors/log/{log_id}/check-out")
async def check_out_visitor_log_id(log_id: str, account: dict = Depends(get_current_account)):
    log = await db_fetchone("SELECT * FROM visitor_logs WHERE id = %s AND check_out IS NULL", (log_id,))
    if not log:
        raise HTTPException(404, "No active check-in found for this log")
        
    await db_execute("UPDATE visitor_logs SET check_out = CURRENT_TIMESTAMP WHERE id = %s", (log_id,))
    
    if log.get("pre_approved_id"):
        pass_obj = await db_fetchone("SELECT pass_type FROM pre_approved_visitors WHERE id = %s", (log["pre_approved_id"],))
        if pass_obj and pass_obj["pass_type"] == "Single Entry":
            await db_execute("UPDATE pre_approved_visitors SET status = 'Expired' WHERE id = %s", (log["pre_approved_id"],))
            
    return {"ok": True}


@api_router.get("/security/visitors/checkin-list")
async def get_checkin_list(account: dict = Depends(get_current_account)):
    rows = await db_fetchall(
        """
        SELECT l.id as log_id, l.id, p.id as pass_id, p.visitor_name, p.mobile, p.visitor_type, p.pass_type, p.pass_code, l.check_in, un.unit_number, b.name as block_name
        FROM visitor_logs l
        JOIN pre_approved_visitors p ON l.pre_approved_id = p.id
        JOIN units un ON p.unit_id = un.id
        JOIN blocks b ON un.block_id = b.id
        WHERE l.check_out IS NULL
        ORDER BY l.check_in DESC
        """
    )
    return {"visitors": rows}

@api_router.get("/security/visitors/checkout-list")
async def get_checkout_list(account: dict = Depends(get_current_account)):
    rows = await db_fetchall(
        """
        SELECT l.id, p.visitor_name, p.mobile, p.visitor_type, p.pass_type, p.pass_code, l.check_in, l.check_out, un.unit_number, b.name as block_name
        FROM visitor_logs l
        JOIN pre_approved_visitors p ON l.pre_approved_id = p.id
        JOIN units un ON p.unit_id = un.id
        JOIN blocks b ON un.block_id = b.id
        WHERE l.check_out IS NOT NULL AND l.check_out >= NOW() - INTERVAL 24 HOUR
        ORDER BY l.check_out DESC
        """
    )
    return {"visitors": rows}


@api_router.get("/security/visitors/history")
async def get_visitor_history(account: dict = Depends(get_current_account)):
    # ONLY past/completed data - NO active visitors!
    
    # 1. Pre-Approved Visitors who have checked out
    pre_approved = await db_fetchall(
        """
        SELECT l.id as visit_id, p.visitor_name, p.mobile, p.purpose, p.visitor_type, 
               l.check_in, l.check_out, 'Exited' as status, un.unit_number, b.name as block_name, 'Pre-Approved' as entry_type
        FROM visitor_logs l
        JOIN pre_approved_visitors p ON l.pre_approved_id = p.id
        JOIN units un ON p.unit_id = un.id
        JOIN blocks b ON un.block_id = b.id
        WHERE l.check_out IS NOT NULL
        ORDER BY l.check_out DESC
        LIMIT 100
        """
    )
    
    # 2. Regular Visitors who have checked out / completed
    regular = await db_fetchall(
        """
        SELECT v.id as visit_id, vis.name as visitor_name, vis.mobile, v.purpose, v.visitor_type, 
               COALESCE(v.check_in_time, v.created_at) as check_in, v.check_out_time as check_out, v.status, un.unit_number, b.name as block_name, 'Regular' as entry_type
        FROM visitor_visits v
        JOIN visitors vis ON v.visitor_id = vis.id
        JOIN units un ON v.unit_id = un.id
        JOIN blocks b ON un.block_id = b.id
        WHERE v.status NOT IN ('Active', 'Entered', 'Pending', 'Checked-In')
        ORDER BY v.created_at DESC
        LIMIT 100
        """
    )
    
    # 3. Expired or Cancelled pre-approved passes that were never checked in
    inactive_passes = await db_fetchall(
        """
        SELECT p.id as visit_id, p.visitor_name, p.mobile, p.purpose, p.visitor_type, 
               p.visit_date as check_in, NULL as check_out, p.status, un.unit_number, b.name as block_name, 'Pre-Approved' as entry_type
        FROM pre_approved_visitors p
        JOIN units un ON p.unit_id = un.id
        JOIN blocks b ON un.block_id = b.id
        WHERE p.status IN ('Expired', 'Cancelled')
        AND NOT EXISTS (SELECT 1 FROM visitor_logs WHERE pre_approved_id = p.id AND check_out IS NOT NULL)
        ORDER BY p.visit_date DESC
        LIMIT 50
        """
    )
    
    combined = pre_approved + regular + inactive_passes
    combined.sort(key=lambda x: str(x.get('check_out') or x.get('check_in') or ''), reverse=True)
    
    return {"visitors": combined[:100]}


# ==========================================
# DELIVERY MANAGEMENT
# ==========================================

@api_router.post("/security/deliveries")
async def create_delivery(payload: DeliveryIn, account: dict = Depends(get_current_account)):
    import uuid
    
    # Try to find a resident for the unit to notify
    # For now we'll just associate with the unit_id
    
    delivery_id = await db_execute(
        """
        INSERT INTO deliveries 
        (unit_id, delivery_type, company_name, delivery_person_name, mobile, status, gate, guard_id, package_photo_url)
        VALUES (%s, %s, %s, %s, %s, %s, %s, %s, %s)
        """,
        (payload.unit_id, payload.delivery_type, payload.company_name, 
         payload.delivery_person_name, payload.mobile, payload.status, payload.gate, 
         account["account_id"], payload.package_photo_url)
    )
    
    # Send notification
    # Get all accounts for this unit
    residents = await db_fetchall("SELECT account_id FROM user_details WHERE unit_id = %s", (payload.unit_id,))
    
    message = f"{payload.delivery_type} delivery from {payload.company_name or 'Unknown'} has arrived."
    if payload.status == "Collected at Gate":
        message += " It was collected by security at the gate."
    else:
        message += " They have been allowed inside the premises."
        
    for res in residents:
        if res.get("account_id"):
            await db_execute(
                """
                INSERT INTO notifications (account_id, title, message, type)
                VALUES (%s, %s, %s, %s)
                """,
                (res["account_id"], "Delivery Update", message, "delivery")
            )
    
    return {"ok": True, "id": delivery_id}

@api_router.get("/security/deliveries/active")
async def get_active_deliveries(account: dict = Depends(get_current_account)):
    rows = await db_fetchall(
        """
        SELECT d.*, un.unit_number, b.name as block_name
        FROM deliveries d
        JOIN units un ON d.unit_id = un.id
        JOIN blocks b ON un.block_id = b.id
        WHERE d.status IN ('Inside Premises', 'Collected at Gate')
        ORDER BY d.check_in DESC
        """
    )
    return {"deliveries": rows}

@api_router.post("/security/deliveries/{delivery_id}/complete")
async def complete_delivery(delivery_id: str, account: dict = Depends(get_current_account)):
    await db_execute("UPDATE deliveries SET status = 'Completed', check_out = CURRENT_TIMESTAMP WHERE id = %s", (delivery_id,))
    return {"ok": True}

@api_router.get("/security/deliveries/history")
async def get_delivery_history(account: dict = Depends(get_current_account)):
    rows = await db_fetchall(
        """
        SELECT d.*, un.unit_number, b.name as block_name
        FROM deliveries d
        JOIN units un ON d.unit_id = un.id
        JOIN blocks b ON un.block_id = b.id
        WHERE d.status IN ('Completed', 'Cancelled')
        ORDER BY COALESCE(d.check_out, d.check_in) DESC
        LIMIT 50
        """
    )
    return {"deliveries": rows}

@api_router.get("/resident/deliveries")
async def get_resident_deliveries(account: dict = Depends(get_current_account)):
    user_detail = await db_fetchone("SELECT unit_id FROM user_details WHERE user_id = %s", (account["user_id"],))
    if not user_detail or not user_detail["unit_id"]:
        return {"deliveries": []}
        
    rows = await db_fetchall(
        """
        SELECT d.*
        FROM deliveries d
        WHERE d.unit_id = %s
        ORDER BY d.check_in DESC
        """,
        (user_detail["unit_id"],)
    )
    return {"deliveries": rows}


@api_router.get("/security/vehicles")
async def get_all_vehicles(account: dict = Depends(get_current_account)):
    rows = await db_fetchall(
        """
        SELECT v.id, v.type as vehicle_type, v.registration_number, 
               ud.name as homeowner_name, un.unit_number, b.name as block_name
        FROM vehicles v
        JOIN accounts u ON v.user_id = u.user_id
        JOIN user_details ud ON u.user_id = ud.user_id
        JOIN units un ON ud.unit_id = un.id
        JOIN blocks b ON un.block_id = b.id
        """
    )
    return {"vehicles": rows}


# ==========================================
# STAFF ENTRY MANAGEMENT
# ==========================================

@api_router.post("/security/staff/seed")
async def seed_staff(account: dict = Depends(get_current_account)):
    import uuid
    # Check if any staff exist
    existing = await db_fetchone("SELECT id FROM service_staff LIMIT 1")
    if existing:
        return {"msg": "Staff already seeded"}

    # Get some unit ids
    units = await db_fetchall("SELECT id FROM units LIMIT 5")
    if len(units) < 2:
        return {"msg": "Not enough units"}
        
    u1, u2 = units[0]["id"], units[1]["id"]
    
    # Insert Staff 1 (Cook)
    s1_id = str(uuid.uuid4())
    await db_execute(
        """
        INSERT INTO service_staff (id, staff_code, name, mobile, category, status, police_verified, id_verified)
        VALUES (%s, %s, %s, %s, %s, %s, %s, %s)
        """,
        (s1_id, "STF-001", "Ramesh Patel", "9876543210", "Cook", "Active", True, True)
    )
    await db_execute("INSERT INTO service_staff_assignments (id, staff_id, unit_id) VALUES (%s, %s, %s)", (str(uuid.uuid4()), s1_id, u1))
    await db_execute("INSERT INTO service_staff_assignments (id, staff_id, unit_id) VALUES (%s, %s, %s)", (str(uuid.uuid4()), s1_id, u2))
    
    # Insert Staff 2 (Maid)
    s2_id = str(uuid.uuid4())
    await db_execute(
        """
        INSERT INTO service_staff (id, staff_code, name, mobile, category, status, police_verified)
        VALUES (%s, %s, %s, %s, %s, %s, %s)
        """,
        (s2_id, "STF-002", "Sunita Sharma", "8765432109", "Housekeeper / Maid", "Active", True)
    )
    await db_execute("INSERT INTO service_staff_assignments (id, staff_id, unit_id) VALUES (%s, %s, %s)", (str(uuid.uuid4()), s2_id, u1))
    
    # Insert Staff 3 (Blocked)
    s3_id = str(uuid.uuid4())
    await db_execute(
        """
        INSERT INTO service_staff (id, staff_code, name, mobile, category, status, police_verified)
        VALUES (%s, %s, %s, %s, %s, %s, %s)
        """,
        (s3_id, "STF-003", "Rahul Driver", "7654321098", "Driver", "Blacklisted", False)
    )
    await db_execute("INSERT INTO service_staff_assignments (id, staff_id, unit_id) VALUES (%s, %s, %s)", (str(uuid.uuid4()), s3_id, u2))
    
    return {"msg": "Successfully seeded service staff"}

@api_router.get("/security/staff/search")
async def search_staff(q: str, account: dict = Depends(get_current_account)):
    # Search by staff_code, mobile, or name
    q = f"%{q}%"
    staff = await db_fetchall(
        """
        SELECT * FROM service_staff 
        WHERE staff_code LIKE %s OR mobile LIKE %s OR name LIKE %s
        """,
        (q, q, q)
    )
    
    # Fetch assignments for each staff
    for s in staff:
        assignments = await db_fetchall(
            """
            SELECT a.unit_id, un.unit_number, b.name as block_name
            FROM service_staff_assignments a
            JOIN units un ON a.unit_id = un.id
            JOIN blocks b ON un.block_id = b.id
            WHERE a.staff_id = %s AND a.status = 'Active'
            """,
            (s["id"],)
        )
        s["assignments"] = assignments
        
        # Check active log
        active_log = await db_fetchone(
            "SELECT id, check_in FROM service_staff_logs WHERE staff_id = %s AND check_out IS NULL",
            (s["id"],)
        )
        s["active_log"] = active_log
        
    return {"staff": staff}

@api_router.post("/security/staff/{staff_id}/check-in")
async def staff_checkin(staff_id: str, payload: StaffCheckIn, account: dict = Depends(get_current_account)):
    import uuid
    # Validate staff
    staff = await db_fetchone("SELECT * FROM service_staff WHERE id = %s", (staff_id,))
    if not staff:
        raise HTTPException(404, "Staff not found")
        
    if staff["status"] != "Active":
        raise HTTPException(400, f"Cannot check in: Staff is {staff['status']}")
        
    # Create log
    log_id = await db_execute(
        """
        INSERT INTO service_staff_logs (staff_id, gate, guard_id, remarks)
        VALUES (%s, %s, %s, %s)
        """,
        (staff_id, payload.gate, account["account_id"], payload.remarks)
    )
    
    # Notify residents
    assignments = await db_fetchall("SELECT unit_id FROM service_staff_assignments WHERE staff_id = %s", (staff_id,))
    for a in assignments:
        residents = await db_fetchall("SELECT account_id FROM user_details WHERE unit_id = %s", (a["unit_id"],))
        for r in residents:
            if r.get("account_id"):
                await db_execute(
                    """
                    INSERT INTO notifications (account_id, title, message, type)
                    VALUES (%s, %s, %s, %s)
                    """,
                    (r["account_id"], "Staff Entry", f"{staff['name']} ({staff['category']}) has entered the society.", "security")
                )
    
    return {"ok": True, "log_id": log_id}

@api_router.post("/security/staff/{staff_id}/check-out")
async def staff_checkout(staff_id: str, payload: StaffCheckIn, account: dict = Depends(get_current_account)):
    import uuid
    # Find active log
    log = await db_fetchone("SELECT id FROM service_staff_logs WHERE staff_id = %s AND check_out IS NULL", (staff_id,))
    if not log:
        raise HTTPException(400, "Staff is not checked in")
        
    await db_execute(
        "UPDATE service_staff_logs SET check_out = CURRENT_TIMESTAMP WHERE id = %s",
        (log["id"],)
    )
    
    # Notify residents
    staff = await db_fetchone("SELECT name, category FROM service_staff WHERE id = %s", (staff_id,))
    assignments = await db_fetchall("SELECT unit_id FROM service_staff_assignments WHERE staff_id = %s", (staff_id,))
    for a in assignments:
        residents = await db_fetchall("SELECT account_id FROM user_details WHERE unit_id = %s", (a["unit_id"],))
        for r in residents:
            if r.get("account_id"):
                await db_execute(
                    """
                    INSERT INTO notifications (account_id, title, message, type)
                    VALUES (%s, %s, %s, %s)
                    """,
                    (r["account_id"], "Staff Exit", f"{staff['name']} ({staff['category']}) has left the society.", "security")
                )
                
    return {"ok": True}

@api_router.get("/security/staff/active")
async def get_active_staff(account: dict = Depends(get_current_account)):
    rows = await db_fetchall(
        """
        SELECT l.id as log_id, l.check_in, s.id as staff_id, s.name, s.staff_code, s.category, s.mobile
        FROM service_staff_logs l
        JOIN service_staff s ON l.staff_id = s.id
        WHERE l.check_out IS NULL
        ORDER BY l.check_in DESC
        """
    )
    return {"staff": rows}

@api_router.get("/security/staff/attendance")
async def get_staff_attendance(account: dict = Depends(get_current_account)):
    rows = await db_fetchall(
        """
        SELECT l.id as log_id, l.check_in, l.check_out, s.id as staff_id, s.name, s.staff_code, s.category, s.mobile
        FROM service_staff_logs l
        JOIN service_staff s ON l.staff_id = s.id
        WHERE DATE(l.check_in) = CURDATE() OR DATE(l.check_out) = CURDATE()
        ORDER BY l.check_in DESC
        LIMIT 100
        """
    )
    return {"attendance": rows}

@api_router.get("/security/staff/blocked")
async def get_blocked_staff(account: dict = Depends(get_current_account)):
    rows = await db_fetchall(
        """
        SELECT id, staff_code, name, mobile, category, status
        FROM service_staff
        WHERE status IN ('Suspended', 'Blacklisted')
        ORDER BY name ASC
        """
    )
    return {"staff": rows}


@api_router.post("/security/staff/auto-checkout")
async def auto_checkout_staff():
    # Called by a CRON job at midnight
    await db_execute(
        "UPDATE service_staff_logs SET check_out = CURRENT_TIMESTAMP, remarks = 'Auto-checkout at midnight' WHERE check_out IS NULL"
    )
    return {"ok": True, "msg": "Successfully auto-checked out all remaining staff"}


class IncidentCreate(BaseModel):
    title: str
    category: str
    incident_type: str
    severity: str
    description: Optional[str] = None
    block: Optional[str] = None
    unit: Optional[str] = None

class IncidentInvestigationUpdate(BaseModel):
    investigation_summary: Optional[str] = None
    root_cause: Optional[str] = None
    corrective_action: Optional[str] = None
    preventive_action: Optional[str] = None
    resolution_notes: Optional[str] = None
    status: Optional[str] = None
    assigned_to: Optional[str] = None

class IncidentStatusUpdate(BaseModel):
    status: str

class IncidentCommentCreate(BaseModel):
    comment: str


# ==========================================
# INCIDENT MANAGEMENT
# ==========================================

@api_router.post("/incidents")
async def create_incident(payload: IncidentCreate, account: dict = Depends(get_current_account)):
    new_id = await db_execute(
        """
        INSERT INTO incidents (title, category, incident_type, severity, description, block, unit, status, reported_by)
        VALUES (%s, %s, %s, %s, %s, %s, %s, 'Open', %s)
        """,
        (payload.title, payload.category, payload.incident_type, payload.severity, payload.description, payload.block, payload.unit, account.get("account_id"))
    )
    await db_execute(
        "INSERT INTO incident_history (incident_id, status, changed_by) VALUES (%s, %s, %s)",
        (new_id, 'Open', account.get("account_id"))
    )
    return {"id": new_id, "message": "Incident created successfully"}

@api_router.get("/incidents")
async def get_incidents(status: Optional[str] = None, category: Optional[str] = None, account: dict = Depends(get_current_account)):
    query = "SELECT i.*, COALESCE(e.name, u.name, a.email) as reporter_name FROM incidents i LEFT JOIN accounts a ON i.reported_by = a.account_id LEFT JOIN employees e ON a.employee_id = e.employee_id LEFT JOIN user_details u ON a.user_id = u.user_id WHERE 1=1"
    params = []
    
    if status:
        query += " AND i.status = %s"
        params.append(status)
        
    if category:
        query += " AND i.category = %s"
        params.append(category)
        
    query += " ORDER BY i.created_at DESC"
    
    incidents = await db_fetchall(query, tuple(params))
    return incidents

@api_router.get("/incidents/{incident_id}")
async def get_incident(incident_id: str, account: dict = Depends(get_current_account)):
    incident = await db_fetchone("SELECT i.*, COALESCE(e.name, u.name, a.email) as reporter_name FROM incidents i LEFT JOIN accounts a ON i.reported_by = a.account_id LEFT JOIN employees e ON a.employee_id = e.employee_id LEFT JOIN user_details u ON a.user_id = u.user_id WHERE i.id = %s", (incident_id,))
    if not incident:
        raise HTTPException(status_code=404, detail="Incident not found")
        
    media = await db_fetchall("SELECT * FROM incident_media WHERE incident_id = %s ORDER BY created_at DESC", (incident_id,))
    comments = await db_fetchall("SELECT c.*, COALESCE(e.name, u.name, a.email) as user_name FROM incident_comments c LEFT JOIN accounts a ON c.user_id = a.account_id LEFT JOIN employees e ON a.employee_id = e.employee_id LEFT JOIN user_details u ON a.user_id = u.user_id WHERE c.incident_id = %s ORDER BY c.created_at ASC", (incident_id,))
    history = await db_fetchall("SELECT h.*, COALESCE(e.name, u.name, a.email) as user_name FROM incident_history h LEFT JOIN accounts a ON h.changed_by = a.account_id LEFT JOIN employees e ON a.employee_id = e.employee_id LEFT JOIN user_details u ON a.user_id = u.user_id WHERE h.incident_id = %s ORDER BY h.timestamp ASC", (incident_id,))
    
    incident['media'] = media
    incident['comments'] = comments
    incident['history'] = history
    return incident

@api_router.put("/incidents/{incident_id}/status")
async def update_incident_status(incident_id: str, payload: IncidentStatusUpdate, account: dict = Depends(get_current_account)):
    incident = await db_fetchone("SELECT id, status FROM incidents WHERE id = %s", (incident_id,))
    if not incident:
        raise HTTPException(status_code=404, detail="Incident not found")
        
    if incident["status"] != payload.status:
        closed_at_sql = ", closed_at = CURRENT_TIMESTAMP" if payload.status == "Closed" else ""
        await db_execute(f"UPDATE incidents SET status = %s{closed_at_sql} WHERE id = %s", (payload.status, incident_id))
        await db_execute(
            "INSERT INTO incident_history (incident_id, status, changed_by) VALUES (%s, %s, %s)",
            (incident_id, payload.status, account.get("account_id"))
        )
    return {"message": "Status updated successfully"}

@api_router.put("/incidents/{incident_id}/investigation")
async def update_incident_investigation(incident_id: str, payload: IncidentInvestigationUpdate, account: dict = Depends(get_current_account)):
    if account.get("role") not in ["Super admin", "Admin", "Security"]:
        raise HTTPException(status_code=403, detail="Unauthorized")
        
    await db_execute(
        """
        UPDATE incidents SET 
            investigation_summary = COALESCE(%s, investigation_summary),
            root_cause = COALESCE(%s, root_cause),
            corrective_action = COALESCE(%s, corrective_action),
            preventive_action = COALESCE(%s, preventive_action),
            resolution_notes = COALESCE(%s, resolution_notes),
            assigned_to = COALESCE(%s, assigned_to)
        WHERE id = %s
        """,
        (payload.investigation_summary, payload.root_cause, payload.corrective_action, payload.preventive_action, payload.resolution_notes, payload.assigned_to, incident_id)
    )
    
    if payload.status:
        incident = await db_fetchone("SELECT status FROM incidents WHERE id = %s", (incident_id,))
        if incident["status"] != payload.status:
            closed_at_sql = ", closed_at = CURRENT_TIMESTAMP" if payload.status == "Closed" else ""
            await db_execute(f"UPDATE incidents SET status = %s{closed_at_sql} WHERE id = %s", (payload.status, incident_id))
            await db_execute(
                "INSERT INTO incident_history (incident_id, status, changed_by) VALUES (%s, %s, %s)",
                (incident_id, payload.status, account.get("account_id"))
            )
            
    return {"message": "Investigation details updated"}

@api_router.post("/incidents/{incident_id}/comments")
async def add_incident_comment(incident_id: str, payload: IncidentCommentCreate, account: dict = Depends(get_current_account)):
    new_id = await db_execute(
        "INSERT INTO incident_comments (incident_id, comment, user_id) VALUES (%s, %s, %s)",
        (incident_id, payload.comment, account.get("account_id"))
    )
    return {"id": new_id, "message": "Comment added"}

@api_router.get("/incidents/reports/summary")
async def get_incident_summary(account: dict = Depends(get_current_account)):
    if account.get("role") not in ["Super admin", "Admin", "Security"]:
        raise HTTPException(status_code=403, detail="Unauthorized")
        
    total = await db_fetchone("SELECT COUNT(*) as c FROM incidents")
    open_count = await db_fetchone("SELECT COUNT(*) as c FROM incidents WHERE status = 'Open'")
    in_progress = await db_fetchone("SELECT COUNT(*) as c FROM incidents WHERE status IN ('Under Review', 'Assigned', 'In Progress')")
    resolved = await db_fetchone("SELECT COUNT(*) as c FROM incidents WHERE status IN ('Resolved', 'Closed')")
    
    by_category = await db_fetchall("SELECT category, COUNT(*) as count FROM incidents GROUP BY category")
    by_severity = await db_fetchall("SELECT severity, COUNT(*) as count FROM incidents GROUP BY severity")
    
    return {
        "overview": {
            "total": total["c"],
            "open": open_count["c"],
            "in_progress": in_progress["c"],
            "resolved": resolved["c"]
        },
        "by_category": by_category,
        "by_severity": by_severity
    }


# --- Unit Documents ---

class UnitDocumentIn(BaseModel):
    association_id: str
    unit_id: Optional[str] = None
    name: str
    type: str
    description: Optional[str] = None
    file_name: Optional[str] = None
    file_url: Optional[str] = None
    file_type: Optional[str] = None
    file_size_kb: Optional[int] = None

@api_router.get("/unit-documents")
async def get_unit_documents(assoc_id: Optional[str] = None, account: dict = Depends(get_current_account)):
    # Admins can see all unit documents for an association
    if account["role"] in ["Super admin", "Admin"]:
        query = """
            SELECT ud.*, a.name as association_name, COALESCE(ud_acc.name, emp.name, 'Unknown User') as user_name, u.unit_number
            FROM unit_documents ud
            JOIN associations a ON ud.association_id = a.id
            JOIN accounts acc ON ud.user_id = acc.account_id
            LEFT JOIN user_details ud_acc ON acc.user_id = ud_acc.user_id
            LEFT JOIN employees emp ON acc.employee_id = emp.employee_id
            LEFT JOIN units u ON ud.unit_id = u.id
        """
        params = []
        if assoc_id:
            query += " WHERE ud.association_id = %s"
            params.append(assoc_id)
        query += " ORDER BY ud.created_at DESC"
        rows = await db_fetchall(query, tuple(params))
    else:
        # Get user's unit_id
        ud_row = None
        if account.get("user_id"):
            ud_row = await db_fetchone("SELECT unit_id FROM user_details WHERE user_id=%s", (account["user_id"],))
        user_unit_id = ud_row["unit_id"] if ud_row else None

        # Homeowners/Tenants only see their own
        query = """
            SELECT ud.*, a.name as association_name, COALESCE(ud_acc.name, emp.name, 'Unknown User') as user_name, u.unit_number
            FROM unit_documents ud
            JOIN associations a ON ud.association_id = a.id
            JOIN accounts acc ON ud.user_id = acc.account_id
            LEFT JOIN user_details ud_acc ON acc.user_id = ud_acc.user_id
            LEFT JOIN employees emp ON acc.employee_id = emp.employee_id
            LEFT JOIN units u ON ud.unit_id = u.id
            WHERE (ud.unit_id = %s OR ud.user_id = %s)
        """
        params = [user_unit_id, account["account_id"]]
        if assoc_id:
            query += " AND ud.association_id = %s"
            params.append(assoc_id)
        query += " ORDER BY ud.created_at DESC"
        rows = await db_fetchall(query, tuple(params))
        
    for r in rows:
        if r.get('created_at'): r['created_at'] = str(r['created_at'])
    return {"ok": True, "data": rows}

@api_router.post("/unit-documents")
async def create_unit_document(payload: UnitDocumentIn, account: dict = Depends(get_current_account)):
    unit_id = payload.unit_id
    if not unit_id:
        if account.get("user_id"):
            ud_row = await db_fetchone("SELECT unit_id FROM user_details WHERE user_id=%s", (account["user_id"],))
            if ud_row and ud_row["unit_id"]:
                unit_id = ud_row["unit_id"]

    await db_execute(
        """
        INSERT INTO unit_documents (
            association_id, unit_id, user_id, name, type, description,
            file_url, file_name, file_type, file_size_kb
        ) VALUES (%s, %s, %s, %s, %s, %s, %s, %s, %s, %s)
        """,
        (
            payload.association_id, unit_id, account["account_id"], payload.name, payload.type, payload.description,
            payload.file_url, payload.file_name, payload.file_type, payload.file_size_kb
        )
    )
    return {"ok": True}

@api_router.delete("/unit-documents/{doc_id}")
async def delete_unit_document(doc_id: str, account: dict = Depends(get_current_account)):
    # Check if they own it or if they are admin
    doc_rows = await db_fetchall("SELECT user_id, unit_id FROM unit_documents WHERE id=%s", (doc_id,))
    if not doc_rows:
        raise HTTPException(status_code=404, detail="Document not found")
        
    if account["role"] not in ["Super admin", "Admin"]:
        ud_row = None
        if account.get("user_id"):
            ud_row = await db_fetchone("SELECT unit_id FROM user_details WHERE user_id=%s", (account["user_id"],))
        user_unit_id = ud_row["unit_id"] if ud_row else None
        if doc_rows[0]["user_id"] != account["account_id"] and (not user_unit_id or doc_rows[0]["unit_id"] != user_unit_id):
            raise HTTPException(status_code=403, detail="Forbidden")
        
    await db_execute("DELETE FROM unit_documents WHERE id=%s", (doc_id,))
    return {"ok": True}

class UnitDocumentUpdateIn(BaseModel):
    name: str
    type: str
    description: Optional[str] = None

@api_router.put("/unit-documents/{doc_id}")
async def update_unit_document(doc_id: str, payload: UnitDocumentUpdateIn, account: dict = Depends(get_current_account)):
    doc_rows = await db_fetchall("SELECT user_id, unit_id FROM unit_documents WHERE id=%s", (doc_id,))
    if not doc_rows:
        raise HTTPException(status_code=404, detail="Document not found")
        
    if account["role"] not in ["Super admin", "Admin"]:
        ud_row = None
        if account.get("user_id"):
            ud_row = await db_fetchone("SELECT unit_id FROM user_details WHERE user_id=%s", (account["user_id"],))
        user_unit_id = ud_row["unit_id"] if ud_row else None
        if doc_rows[0]["user_id"] != account["account_id"] and (not user_unit_id or doc_rows[0]["unit_id"] != user_unit_id):
            raise HTTPException(status_code=403, detail="Forbidden")
        
    await db_execute(
        """
        UPDATE unit_documents SET
            name=%s, type=%s, description=%s
        WHERE id=%s
        """,
        (payload.name, payload.type, payload.description, doc_id)
    )
    return {"ok": True}



# --- Marketplace ---

class MarketplaceItemIn(BaseModel):
    title: str
    description: Optional[str] = None
    category_id: str
    price: float
    condition_state: Optional[str] = None
    brand: Optional[str] = None
    item_age: Optional[str] = None
    location: str
    contact_number: str
    listing_type: str
    is_negotiable: Optional[bool] = False
    status: Optional[str] = 'Active'
    images: List[str] = []

class MarketplaceItemUpdate(BaseModel):
    title: Optional[str] = None
    description: Optional[str] = None
    category_id: Optional[str] = None
    price: Optional[float] = None
    condition_state: Optional[str] = None
    brand: Optional[str] = None
    item_age: Optional[str] = None
    location: Optional[str] = None
    contact_number: Optional[str] = None
    listing_type: Optional[str] = None
    is_negotiable: Optional[bool] = None
    status: Optional[str] = None
    images: Optional[List[str]] = None

class MarketplaceChatMessageIn(BaseModel):
    message: str
    receiver_id: Optional[str] = None

class MarketplaceReportIn(BaseModel):
    reason: str

@api_router.get("/marketplace/categories")
async def get_marketplace_categories():
    rows = await db_fetchall("SELECT * FROM marketplace_categories ORDER BY name")
    return {"ok": True, "data": rows}

@api_router.get("/marketplace/items")
async def get_marketplace_items(
    account: dict = Depends(get_current_account),
    category_id: Optional[str] = None,
    condition: Optional[str] = None,
    min_price: Optional[float] = None,
    max_price: Optional[float] = None,
    is_negotiable: Optional[bool] = None,
    status: Optional[str] = 'Active',
    sort: Optional[str] = 'Newest'
):
    query = '''
        SELECT i.*, c.name as category_name, 
               COALESCE(ud.name, emp.name, acc.email) as seller_name,
               (SELECT COUNT(*) FROM marketplace_favorites f WHERE f.item_id = i.id AND f.user_id = %s) as is_saved,
               (SELECT COUNT(*) FROM marketplace_favorites f WHERE f.item_id = i.id) as saves_count
        FROM marketplace_items i
        LEFT JOIN marketplace_categories c ON i.category_id = c.id
        LEFT JOIN accounts acc ON i.user_id = acc.account_id
        LEFT JOIN user_details ud ON acc.user_id = ud.user_id
        LEFT JOIN employees emp ON acc.employee_id = emp.employee_id
        WHERE 1=1
    '''
    params = [account["account_id"]]
    
    if status:
        query += " AND i.status = %s"
        params.append(status)
    if category_id:
        query += " AND i.category_id = %s"
        params.append(category_id)
    if condition:
        query += " AND i.condition_state = %s"
        params.append(condition)
    if min_price is not None:
        query += " AND i.price >= %s"
        params.append(min_price)
    if max_price is not None:
        query += " AND i.price <= %s"
        params.append(max_price)
    if is_negotiable is not None:
        query += " AND i.is_negotiable = %s"
        params.append(is_negotiable)
        
    if sort == 'Oldest':
        query += " ORDER BY i.created_at ASC"
    else:
        query += " ORDER BY i.created_at DESC"
        
    items = await db_fetchall(query, tuple(params))
    
    # Fetch images for all items
    if items:
        item_ids = [str(item['id']) for item in items]
        placeholders = ','.join(['%s'] * len(item_ids))
        img_query = f"SELECT item_id, image_url FROM marketplace_images WHERE item_id IN ({placeholders})"
        images = await db_fetchall(img_query, tuple(item_ids))
        
        img_map = {}
        for img in images:
            img_map.setdefault(img['item_id'], []).append(img['image_url'])
            
        for item in items:
            item['images'] = img_map.get(item['id'], [])
            if item.get('created_at'): item['created_at'] = str(item['created_at'])
            if item.get('updated_at'): item['updated_at'] = str(item['updated_at'])
            
    return {"ok": True, "data": items}

@api_router.get("/marketplace/items/my")
async def get_my_marketplace_items(account: dict = Depends(get_current_account)):
    query = '''
        SELECT i.*, c.name as category_name,
               (SELECT COUNT(*) FROM marketplace_favorites f WHERE f.item_id = i.id) as saves_count
        FROM marketplace_items i
        LEFT JOIN marketplace_categories c ON i.category_id = c.id
        WHERE i.user_id = %s
        ORDER BY i.created_at DESC
    '''
    items = await db_fetchall(query, (account["account_id"],))
    
    if items:
        item_ids = [str(item['id']) for item in items]
        placeholders = ','.join(['%s'] * len(item_ids))
        img_query = f"SELECT item_id, image_url FROM marketplace_images WHERE item_id IN ({placeholders})"
        images = await db_fetchall(img_query, tuple(item_ids))
        
        img_map = {}
        for img in images:
            img_map.setdefault(img['item_id'], []).append(img['image_url'])
            
        for item in items:
            item['images'] = img_map.get(item['id'], [])
            if item.get('created_at'): item['created_at'] = str(item['created_at'])
            if item.get('updated_at'): item['updated_at'] = str(item['updated_at'])
            
    return {"ok": True, "data": items}

@api_router.post("/marketplace/items")
async def create_marketplace_item(payload: MarketplaceItemIn, account: dict = Depends(get_current_account)):
    # Get association_id for the user
    assoc_id = None
    if account.get("user_id"):
        ud_row = await db_fetchone("SELECT ud.unit_id, u.block_id, b.association_id FROM user_details ud LEFT JOIN units u ON ud.unit_id=u.id LEFT JOIN blocks b ON u.block_id=b.id WHERE ud.user_id=%s", (account["user_id"],))
        if ud_row:
            assoc_id = ud_row.get("association_id")
    elif account.get("employee_id"):
        emp_row = await db_fetchone("SELECT association_id FROM employees WHERE employee_id=%s", (account["employee_id"],))
        if emp_row:
            assoc_id = emp_row.get("association_id")
            
    if not assoc_id:
        assoc_id = "00000000-0000-0000-0000-000000000000" # Fallback if no association found (Admins might not have one)
        
    item_id = await db_execute('''
        INSERT INTO marketplace_items (
            association_id, user_id, category_id, title, description, price, condition_state, 
            brand, item_age, location, contact_number, is_negotiable, status, listing_type
        ) VALUES (%s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s)
    ''', (
        assoc_id, account["account_id"], payload.category_id, payload.title, payload.description,
        payload.price, payload.condition_state, payload.brand, payload.item_age, payload.location,
        payload.contact_number, payload.is_negotiable, payload.status, payload.listing_type
    ))
    
    for img_url in payload.images:
        await db_execute("INSERT INTO marketplace_images (item_id, image_url) VALUES (%s, %s)", (item_id, img_url))
        
    return {"ok": True, "id": item_id}

@api_router.put("/marketplace/items/{item_id}")
async def update_marketplace_item(item_id: str, payload: MarketplaceItemUpdate, account: dict = Depends(get_current_account)):
    item = await db_fetchone("SELECT user_id FROM marketplace_items WHERE id=%s", (item_id,))
    if not item:
        raise HTTPException(status_code=404, detail="Item not found")
        
    if item['user_id'] != account["account_id"] and account["role"] not in ["Super admin", "Admin"]:
        raise HTTPException(status_code=403, detail="Forbidden")
        
    updates = []
    params = []
    for k, v in payload.dict(exclude_unset=True).items():
        if k != "images":
            updates.append(f"{k}=%s")
            params.append(v)
            
    if updates:
        params.append(item_id)
        query = f"UPDATE marketplace_items SET {', '.join(updates)} WHERE id=%s"
        await db_execute(query, tuple(params))
        
    if payload.images is not None:
        await db_execute("DELETE FROM marketplace_images WHERE item_id=%s", (item_id,))
        for img_url in payload.images:
            await db_execute("INSERT INTO marketplace_images (item_id, image_url) VALUES (%s, %s)", (item_id, img_url))
            
    return {"ok": True}

@api_router.delete("/marketplace/items/{item_id}")
async def delete_marketplace_item(item_id: str, account: dict = Depends(get_current_account)):
    item = await db_fetchone("SELECT user_id FROM marketplace_items WHERE id=%s", (item_id,))
    if not item:
        raise HTTPException(status_code=404, detail="Item not found")
        
    if item['user_id'] != account["account_id"] and account["role"] not in ["Super admin", "Admin"]:
        raise HTTPException(status_code=403, detail="Forbidden")
        
    await db_execute("DELETE FROM marketplace_items WHERE id=%s", (item_id,))
    return {"ok": True}

@api_router.get("/marketplace/favorites")
async def get_marketplace_favorites(account: dict = Depends(get_current_account)):
    query = '''
        SELECT i.*, c.name as category_name, 1 as is_saved,
               (SELECT COUNT(*) FROM marketplace_favorites f WHERE f.item_id = i.id) as saves_count
        FROM marketplace_items i
        JOIN marketplace_favorites mf ON i.id = mf.item_id
        LEFT JOIN marketplace_categories c ON i.category_id = c.id
        WHERE mf.user_id = %s
        ORDER BY mf.created_at DESC
    '''
    items = await db_fetchall(query, (account["account_id"],))
    
    if items:
        item_ids = [str(item['id']) for item in items]
        placeholders = ','.join(['%s'] * len(item_ids))
        img_query = f"SELECT item_id, image_url FROM marketplace_images WHERE item_id IN ({placeholders})"
        images = await db_fetchall(img_query, tuple(item_ids))
        
        img_map = {}
        for img in images:
            img_map.setdefault(img['item_id'], []).append(img['image_url'])
            
        for item in items:
            item['images'] = img_map.get(item['id'], [])
            if item.get('created_at'): item['created_at'] = str(item['created_at'])
            if item.get('updated_at'): item['updated_at'] = str(item['updated_at'])
            
    return {"ok": True, "data": items}

@api_router.post("/marketplace/favorites/{item_id}")
async def toggle_marketplace_favorite(item_id: str, account: dict = Depends(get_current_account)):
    existing = await db_fetchone("SELECT id FROM marketplace_favorites WHERE user_id=%s AND item_id=%s", (account["account_id"], item_id))
    if existing:
        await db_execute("DELETE FROM marketplace_favorites WHERE id=%s", (existing["id"],))
        return {"ok": True, "saved": False}
    else:
        await db_execute("INSERT INTO marketplace_favorites (user_id, item_id) VALUES (%s, %s)", (account["account_id"], item_id))
        return {"ok": True, "saved": True}

@api_router.post("/marketplace/reports/{item_id}")
async def report_marketplace_item(item_id: str, payload: MarketplaceReportIn, account: dict = Depends(get_current_account)):
    await db_execute("INSERT INTO marketplace_reports (item_id, reporter_id, reason) VALUES (%s, %s, %s)", (item_id, account["account_id"], payload.reason))
    return {"ok": True}

@api_router.get("/marketplace/chat/{item_id}/threads")
async def get_marketplace_chat_threads(item_id: str, account: dict = Depends(get_current_account)):
    item = await db_fetchone("SELECT user_id FROM marketplace_items WHERE id=%s", (item_id,))
    if not item:
        raise HTTPException(status_code=404, detail="Item not found")
    if item['user_id'] != account["account_id"]:
        raise HTTPException(status_code=403, detail="Only the owner can view threads")
        
    query = '''
        SELECT 
            buyer.account_id as buyer_id,
            COALESCE(ud.name, emp.name, buyer.email) as buyer_name,
            buyer.email as buyer_email,
            MAX(c.created_at) as last_message_at,
            (SELECT message FROM marketplace_chat WHERE item_id = %s AND (sender_id = buyer.account_id OR receiver_id = buyer.account_id) ORDER BY created_at DESC LIMIT 1) as last_message
        FROM marketplace_chat c
        JOIN accounts buyer ON (buyer.account_id = c.sender_id OR buyer.account_id = c.receiver_id) AND buyer.account_id != %s
        LEFT JOIN user_details ud ON buyer.user_id = ud.user_id
        LEFT JOIN employees emp ON buyer.employee_id = emp.employee_id
        WHERE c.item_id = %s
        GROUP BY buyer.account_id, buyer_name, buyer_email
        ORDER BY last_message_at DESC
    '''
    threads = await db_fetchall(query, (item_id, account["account_id"], item_id))
    for t in threads:
        if t.get('last_message_at'): t['last_message_at'] = str(t['last_message_at'])
        
    return {"ok": True, "data": threads}

@api_router.get("/marketplace/chat/{item_id}")
async def get_marketplace_chat(item_id: str, buyer_id: Optional[str] = None, account: dict = Depends(get_current_account)):
    item = await db_fetchone("SELECT user_id, title FROM marketplace_items WHERE id=%s", (item_id,))
    if not item:
        raise HTTPException(status_code=404, detail="Item not found")
        
    is_owner = item['user_id'] == account["account_id"]
    
    if is_owner:
        if not buyer_id:
            return {"ok": True, "data": []}
        target_buyer = buyer_id
    else:
        target_buyer = account["account_id"]
        
    query = '''
        SELECT c.*, acc.email, COALESCE(ud.name, emp.name, acc.email) as sender_name,
               (c.sender_id = %s) as is_mine
        FROM marketplace_chat c
        JOIN accounts acc ON c.sender_id = acc.account_id
        LEFT JOIN user_details ud ON acc.user_id = ud.user_id
        LEFT JOIN employees emp ON acc.employee_id = emp.employee_id
        WHERE c.item_id = %s 
          AND ((c.sender_id = %s AND c.receiver_id = %s) OR (c.sender_id = %s AND c.receiver_id = %s))
        ORDER BY c.created_at ASC
    '''
    messages = await db_fetchall(query, (
        account["account_id"], 
        item_id, 
        target_buyer, item['user_id'], 
        item['user_id'], target_buyer
    ))
    
    for msg in messages:
        if msg.get('created_at'): msg['created_at'] = str(msg['created_at'])
        
    return {"ok": True, "data": messages}

@api_router.post("/marketplace/chat/{item_id}")
async def send_marketplace_chat(item_id: str, payload: MarketplaceChatMessageIn, account: dict = Depends(get_current_account)):
    item = await db_fetchone("SELECT user_id FROM marketplace_items WHERE id=%s", (item_id,))
    if not item:
        raise HTTPException(status_code=404, detail="Item not found")
        
    is_owner = account["account_id"] == item['user_id']
    
    if is_owner:
        if not payload.receiver_id:
            raise HTTPException(status_code=400, detail="Must specify receiver_id when replying as owner")
        receiver_id = payload.receiver_id
    else:
        receiver_id = item['user_id']
            
    await db_execute("INSERT INTO marketplace_chat (item_id, sender_id, receiver_id, message) VALUES (%s, %s, %s, %s)", (item_id, account["account_id"], receiver_id, payload.message))
    return {"ok": True}

@api_router.post("/marketplace/views/{item_id}")
async def record_marketplace_view(item_id: str, account: dict = Depends(get_current_account)):
    try:
        await db_execute("INSERT INTO marketplace_views (item_id, user_id) VALUES (%s, %s)", (item_id, account["account_id"]))
        await db_execute("UPDATE marketplace_items SET views_count = views_count + 1 WHERE id=%s", (item_id,))
    except Exception:
        pass # Ignore duplicates
    return {"ok": True}




class AssessmentRuleBase(BaseModel):
    frequency: str
    default_amount: float
    due_day_of_month: int

class FineRuleBase(BaseModel):
    fine_type: str
    amount: float
    grace_period_days: int

class AssociationSettingsUpdate(BaseModel):
    end_date: Optional[str] = None
    assessment_rules: Optional[AssessmentRuleBase] = None
    fine_rules: Optional[List[FineRuleBase]] = None

@api_router.get("/admin/associations/{id}/settings")
async def get_association_settings(id: str, account: dict = Depends(require_role(["Super admin", "Admin"]))):
    assoc = await db_fetchone("SELECT onboarding_date, end_date FROM associations WHERE id = %s", (id,))
    if not assoc:
        raise HTTPException(status_code=404, detail="Association not found")
        
    assessment = await db_fetchone("SELECT frequency, default_amount, due_day_of_month FROM assessment_rules WHERE association_id = %s LIMIT 1", (id,))
    fines = await db_fetchall("SELECT id, fine_type, amount, grace_period_days FROM fine_rules WHERE association_id = %s", (id,))
    
    # Format dates
    if assoc.get("onboarding_date"):
        assoc["onboarding_date"] = assoc["onboarding_date"].strftime("%Y-%m-%d") if hasattr(assoc["onboarding_date"], 'strftime') else str(assoc["onboarding_date"])
    if assoc.get("end_date"):
        assoc["end_date"] = assoc["end_date"].strftime("%Y-%m-%d") if hasattr(assoc["end_date"], 'strftime') else str(assoc["end_date"])
        
    return {
        "onboarding_date": assoc.get("onboarding_date"),
        "end_date": assoc.get("end_date"),
        "assessment_rules": assessment,
        "fine_rules": fines
    }

@api_router.put("/admin/associations/{id}/settings")
async def update_association_settings(id: str, payload: AssociationSettingsUpdate, account: dict = Depends(require_role(["Super admin", "Admin"]))):
    import uuid
    # Update end_date in associations
    if payload.end_date is not None:
        await db_execute("UPDATE associations SET end_date = %s WHERE id = %s", (payload.end_date if payload.end_date != "" else None, id))
        
    # Update assessment_rules
    if payload.assessment_rules:
        existing = await db_fetchone("SELECT id FROM assessment_rules WHERE association_id = %s", (id,))
        if existing:
            await db_execute("UPDATE assessment_rules SET frequency=%s, default_amount=%s, due_day_of_month=%s WHERE association_id=%s", 
                             (payload.assessment_rules.frequency, payload.assessment_rules.default_amount, payload.assessment_rules.due_day_of_month, id))
        else:
            await db_execute("INSERT INTO assessment_rules (id, association_id, frequency, default_amount, due_day_of_month) VALUES (%s, %s, %s, %s, %s)",
                             (str(uuid.uuid4()), id, payload.assessment_rules.frequency, payload.assessment_rules.default_amount, payload.assessment_rules.due_day_of_month))
                             
    # Update fine_rules
    if payload.fine_rules is not None:
        # For simplicity, we delete existing and recreate
        await db_execute("DELETE FROM fine_rules WHERE association_id = %s", (id,))
        for fr in payload.fine_rules:
            await db_execute("INSERT INTO fine_rules (id, association_id, fine_type, amount, grace_period_days) VALUES (%s, %s, %s, %s, %s)",
                             (str(uuid.uuid4()), id, fr.fine_type, fr.amount, fr.grace_period_days))
                             
    return {"ok": True, "message": "Settings updated"}




# ==========================================
# Wallet APIs
# ==========================================

from pydantic import BaseModel
from typing import Optional, List

class PinSetupIn(BaseModel):
    pin: str

class AddMoneyIn(BaseModel):
    amount: float
    method: str

class PayDuesIn(BaseModel):
    amount: float
    pin: str

@api_router.get("/wallet")
async def get_wallet(account: dict = Depends(get_current_account)):
    wallet = await db_fetchone("SELECT id, balance, reward_points, security_pin FROM wallets WHERE account_id = %s", (account["account_id"],))
    if not wallet:
        raise HTTPException(status_code=404, detail="Wallet not found.")
    
    return {
        "ok": True,
        "data": {
            "id": wallet["id"],
            "balance": float(wallet["balance"]),
            "reward_points": wallet["reward_points"],
            "has_pin": bool(wallet["security_pin"])
        }
    }

@api_router.post("/wallet/pin")
async def setup_wallet_pin(payload: PinSetupIn, account: dict = Depends(get_current_account)):
    if len(payload.pin) != 4 or not payload.pin.isdigit():
        raise HTTPException(status_code=400, detail="PIN must be exactly 4 digits.")
    
    wallet = await db_fetchone("SELECT id FROM wallets WHERE account_id = %s", (account["account_id"],))
    if not wallet:
        raise HTTPException(status_code=404, detail="Wallet not found.")
        
    await db_execute("UPDATE wallets SET security_pin = %s WHERE id = %s", (payload.pin, wallet["id"]))
    return {"ok": True, "message": "PIN set successfully."}

@api_router.post("/wallet/add-money")
async def add_wallet_money(payload: AddMoneyIn, account: dict = Depends(get_current_account)):
    if payload.amount <= 0:
        raise HTTPException(status_code=400, detail="Amount must be positive.")
        
    wallet = await db_fetchone("SELECT id, balance FROM wallets WHERE account_id = %s", (account["account_id"],))
    if not wallet:
        raise HTTPException(status_code=404, detail="Wallet not found.")
        
    # Mock adding money
    new_balance = float(wallet["balance"]) + payload.amount
    await db_execute("UPDATE wallets SET balance = %s WHERE id = %s", (new_balance, wallet["id"]))
    
    # Record transaction
    await db_execute(
        "INSERT INTO wallet_transactions (wallet_id, type, amount, status, description) VALUES (%s, %s, %s, %s, %s)",
        (wallet["id"], "Credit", payload.amount, "Completed", f"Added via {payload.method}")
    )
    
    return {"ok": True, "message": "Money added successfully.", "balance": new_balance}

@api_router.post("/wallet/pay-dues")
async def pay_wallet_dues(payload: PayDuesIn, account: dict = Depends(get_current_account)):
    if payload.amount <= 0:
        raise HTTPException(status_code=400, detail="Amount must be positive.")
        
    wallet = await db_fetchone("SELECT id, balance, security_pin FROM wallets WHERE account_id = %s", (account["account_id"],))
    if not wallet:
        raise HTTPException(status_code=404, detail="Wallet not found.")
        
    if not wallet["security_pin"] or wallet["security_pin"] != payload.pin:
        raise HTTPException(status_code=403, detail="Invalid PIN.")
        
    if float(wallet["balance"]) < payload.amount:
        raise HTTPException(status_code=400, detail="Insufficient wallet balance.")
        
    # Prevent double payment for the same month
    user_unit = await db_fetchone("SELECT unit_id FROM user_details WHERE user_id = %s", (account["user_id"],))
    unit_id = user_unit["unit_id"] if user_unit else None

    now = datetime.now()
    current_month_start = datetime(now.year, now.month, 1).strftime('%Y-%m-%d 00:00:00')
    existing_payment = None

    if unit_id:
        existing_payment = await db_fetchone("""
            SELECT wt.id 
            FROM wallet_transactions wt
            JOIN wallets w ON wt.wallet_id = w.id
            JOIN accounts a ON w.account_id = a.account_id
            JOIN user_details ud ON a.user_id = ud.user_id
            WHERE ud.unit_id = %s
              AND wt.type IN ('Debit', 'UPI') 
              AND wt.description = 'Maintenance Dues Paid' 
              AND wt.created_at >= %s
            LIMIT 1
        """, (unit_id, current_month_start))
    else:
        existing_payment = await db_fetchone(
            "SELECT id FROM wallet_transactions WHERE wallet_id = %s AND type IN ('Debit', 'UPI') AND description = 'Maintenance Dues Paid' AND created_at >= %s",
            (wallet["id"], current_month_start)
        )
    if existing_payment:
        raise HTTPException(status_code=400, detail="Dues already paid for this month.")
        
    # Process payment
    new_balance = float(wallet["balance"]) - payload.amount
    await db_execute("UPDATE wallets SET balance = %s WHERE id = %s", (new_balance, wallet["id"]))
    
    # Record transaction
    await db_execute(
        "INSERT INTO wallet_transactions (wallet_id, type, amount, status, description) VALUES (%s, %s, %s, %s, %s)",
        (wallet["id"], "Debit", payload.amount, "Completed", "Maintenance Dues Paid")
    )
    
    # TODO: We might need to update a 'payments' table or similar to mark dues as paid, 
    # but the current schema does not have a distinct dues tracker for homeowners except assessment_rules.
    
    return {"ok": True, "message": "Dues paid successfully.", "balance": new_balance}

class PayDuesUPIRequest(BaseModel):
    amount: float

@api_router.post("/wallet/pay-dues-upi")
async def pay_dues_upi(payload: PayDuesUPIRequest, account: dict = Depends(get_current_account)):
    wallet = await db_fetchone("SELECT id, balance FROM wallets WHERE account_id = %s", (account["account_id"],))
    if not wallet:
        raise HTTPException(status_code=404, detail="Wallet not found. Please verify details first.")
        
    # Prevent double payment for the same month
    user_unit = await db_fetchone("SELECT unit_id FROM user_details WHERE user_id = %s", (account["user_id"],))
    unit_id = user_unit["unit_id"] if user_unit else None

    now = datetime.now()
    current_month_start = datetime(now.year, now.month, 1).strftime('%Y-%m-%d 00:00:00')
    existing_payment = None

    if unit_id:
        existing_payment = await db_fetchone("""
            SELECT wt.id 
            FROM wallet_transactions wt
            JOIN wallets w ON wt.wallet_id = w.id
            JOIN accounts a ON w.account_id = a.account_id
            JOIN user_details ud ON a.user_id = ud.user_id
            WHERE ud.unit_id = %s
              AND wt.type IN ('Debit', 'UPI') 
              AND wt.description = 'Maintenance Dues Paid' 
              AND wt.created_at >= %s
            LIMIT 1
        """, (unit_id, current_month_start))
    else:
        existing_payment = await db_fetchone(
            "SELECT id FROM wallet_transactions WHERE wallet_id = %s AND type IN ('Debit', 'UPI') AND description = 'Maintenance Dues Paid' AND created_at >= %s",
            (wallet["id"], current_month_start)
        )
    if existing_payment:
        raise HTTPException(status_code=400, detail="Dues already paid for this month.")
        
    # Record transaction (does not affect wallet balance)
    await db_execute(
        "INSERT INTO wallet_transactions (wallet_id, type, amount, status, description) VALUES (%s, %s, %s, %s, %s)",
        (wallet["id"], "UPI", payload.amount, "Completed", "Maintenance Dues Paid")
    )
    
    return {"ok": True, "message": "Dues paid successfully via UPI."}

@api_router.get("/wallet/transactions")
async def get_wallet_transactions(account: dict = Depends(get_current_account)):
    wallet = await db_fetchone("SELECT id FROM wallets WHERE account_id = %s", (account["account_id"],))
    if not wallet:
        return {"ok": True, "data": []}
        
    txs = await db_fetchall(
        "SELECT id, type, amount, status, description, created_at FROM wallet_transactions WHERE wallet_id = %s ORDER BY created_at DESC", 
        (wallet["id"],)
    )
    
    for tx in txs:
        tx["amount"] = float(tx["amount"])
        tx["created_at"] = tx["created_at"].isoformat()
        
    return {"ok": True, "data": txs}


class SendMoneyIn(BaseModel):
    recipient_email: str
    amount: float
    pin: str
    purpose: Optional[str] = None
    notes: Optional[str] = None

@api_router.post("/wallet/send")
async def send_wallet_money(payload: SendMoneyIn, account: dict = Depends(get_current_account)):
    if payload.amount <= 0:
        raise HTTPException(status_code=400, detail="Amount must be positive.")
        
    sender_wallet = await db_fetchone("SELECT id, balance, security_pin FROM wallets WHERE account_id = %s", (account["account_id"],))
    if not sender_wallet:
        raise HTTPException(status_code=404, detail="Your wallet was not found.")
        
    if not sender_wallet["security_pin"] or sender_wallet["security_pin"] != payload.pin:
        raise HTTPException(status_code=403, detail="Invalid PIN.")
        
    if float(sender_wallet["balance"]) < payload.amount:
        raise HTTPException(status_code=400, detail="Insufficient wallet balance.")
        
    recipient = await db_fetchone("SELECT account_id FROM accounts WHERE email = %s", (payload.recipient_email.lower(),))
    if not recipient:
        raise HTTPException(status_code=404, detail="Recipient not found.")
        
    if recipient["account_id"] == account["account_id"]:
        raise HTTPException(status_code=400, detail="Cannot send money to yourself.")
        
    recipient_wallet = await db_fetchone("SELECT id, balance FROM wallets WHERE account_id = %s", (recipient["account_id"],))
    if not recipient_wallet:
        raise HTTPException(status_code=404, detail="Recipient does not have a wallet set up.")
        
    # Process transfer
    new_sender_balance = float(sender_wallet["balance"]) - payload.amount
    new_recipient_balance = float(recipient_wallet["balance"]) + payload.amount
    
    await db_execute("UPDATE wallets SET balance = %s WHERE id = %s", (new_sender_balance, sender_wallet["id"]))
    await db_execute("UPDATE wallets SET balance = %s WHERE id = %s", (new_recipient_balance, recipient_wallet["id"]))
    
    # Record transactions
    desc_suffix = f" for {payload.purpose}" if payload.purpose else ""
    await db_execute(
        "INSERT INTO wallet_transactions (wallet_id, type, amount, status, description) VALUES (%s, %s, %s, %s, %s)",
        (sender_wallet["id"], "Debit", payload.amount, "Completed", f"Sent to {payload.recipient_email}{desc_suffix}")
    )
    await db_execute(
        "INSERT INTO wallet_transactions (wallet_id, type, amount, status, description) VALUES (%s, %s, %s, %s, %s)",
        (recipient_wallet["id"], "Credit", payload.amount, "Completed", f"Received from {account['email']}{desc_suffix}")
    )
    
    return {"ok": True, "message": "Money sent successfully.", "balance": new_sender_balance}


# ---------- Fintech Collection & Wallet System ----------

async def ensure_virtual_account(account_id: str) -> dict:
    va = await db_fetchone("SELECT * FROM virtual_accounts WHERE account_id = %s", (account_id,))
    if va:
        if va.get("created_at"): va["created_at"] = str(va["created_at"])
        if va.get("expires_at"): va["expires_at"] = str(va["expires_at"])
        return va
    
    user_info = await db_fetchone(
        "SELECT a.email, COALESCE(ud.name, a.email) as name FROM accounts a LEFT JOIN user_details ud ON a.user_id = ud.user_id WHERE a.account_id = %s",
        (account_id,)
    )
    user_name = user_info["name"] if user_info else "Homeowner"
    clean_name = "".join(e for e in user_name if e.isalnum())[:10].upper() or "HOMEOWNER"
    
    import random
    numeric_suffix = str(random.randint(10000000, 99999999))
    va_num = f"VA{numeric_suffix}"
    va_ifsc = "ICIC000001"
    va_vpa = f"nestora.{clean_name.lower()}@icici"
    ref_num = f"REF{numeric_suffix}"
    va_id = f"va_{uuid.uuid4().hex[:12]}"
    
    new_id = await db_execute(
        """INSERT INTO virtual_accounts 
           (account_id, virtual_account_id, virtual_account_number, virtual_ifsc, virtual_vpa, beneficiary_name, reference_number, status)
           VALUES (%s, %s, %s, %s, %s, %s, %s, %s)""",
        (account_id, va_id, va_num, va_ifsc, va_vpa, user_name, ref_num, 'active')
    )
    
    res = await db_fetchone("SELECT * FROM virtual_accounts WHERE id = %s OR virtual_account_number = %s", (new_id, va_num))
    if res.get("created_at"): res["created_at"] = str(res["created_at"])
    if res.get("expires_at"): res["expires_at"] = str(res["expires_at"])
    return res

@api_router.get("/wallet/virtual-account")
async def get_my_virtual_account(account: dict = Depends(get_current_account)):
    va = await ensure_virtual_account(account["account_id"])
    return {"ok": True, "data": va}

class WalletTopupIn(BaseModel):
    amount: float
    payment_method: Optional[str] = "Virtual Account"

@api_router.post("/wallet/topup")
async def create_wallet_topup(payload: WalletTopupIn, account: dict = Depends(get_current_account)):
    if payload.amount <= 0:
        raise HTTPException(status_code=400, detail="Amount must be positive.")
    va = await ensure_virtual_account(account["account_id"])
    import random
    ref_num = f"TOPUP{random.randint(100000, 999999)}"
    topup_id = await db_execute(
        "INSERT INTO wallet_topups (account_id, amount, reference_number, payment_method, status) VALUES (%s, %s, %s, %s, 'PENDING')",
        (account["account_id"], payload.amount, ref_num, payload.payment_method)
    )
    return {
        "ok": True,
        "message": "Topup request created.",
        "data": {
            "id": topup_id,
            "reference_number": ref_num,
            "amount": payload.amount,
            "virtual_account_number": va["virtual_account_number"],
            "virtual_ifsc": va["virtual_ifsc"],
            "virtual_vpa": va["virtual_vpa"]
        }
    }

class UPICollectIn(BaseModel):
    amount: float
    vpa: str
    customer_name: Optional[str] = None
    customer_email: Optional[str] = None
    customer_phone: Optional[str] = None

@api_router.post("/wallet/upi-collect")
async def create_upi_collect(payload: UPICollectIn, account: dict = Depends(get_current_account)):
    if payload.amount <= 0:
        raise HTTPException(status_code=400, detail="Amount must be positive.")
    if "@" not in payload.vpa:
        raise HTTPException(status_code=400, detail="Invalid UPI VPA address.")
        
    import random
    req_id = f"REQ{random.randint(10000000, 99999999)}"
    txnid = f"TXN{random.randint(10000000, 99999999)}"
    
    upi_id = await db_execute(
        """INSERT INTO upi_collections 
           (account_id, amount, vpa, customer_name, customer_email, customer_phone, request_id, txnid, status)
           VALUES (%s, %s, %s, %s, %s, %s, %s, %s, 'PENDING')""",
        (account["account_id"], payload.amount, payload.vpa, payload.customer_name or account.get("email"), payload.customer_email or account.get("email"), payload.customer_phone or "", req_id, txnid)
    )
    
    return {
        "ok": True,
        "message": f"UPI collect request sent to {payload.vpa}",
        "data": {
            "id": upi_id,
            "request_id": req_id,
            "txnid": txnid,
            "amount": payload.amount,
            "vpa": payload.vpa,
            "status": "PENDING"
        }
    }

@api_router.get("/wallet/ledger")
async def get_wallet_ledger(account: dict = Depends(get_current_account)):
    wallet = await db_fetchone("SELECT id FROM wallets WHERE account_id = %s", (account["account_id"],))
    if not wallet:
        return {"ok": True, "data": []}
    
    ledger_entries = await db_fetchall(
        """SELECT id, transaction_type, amount, balance_after, reference_id, payment_id, utr, description, status, created_at
           FROM wallet_ledger 
           WHERE account_id = %s ORDER BY created_at DESC""",
        (account["account_id"],)
    )
    
    for item in ledger_entries:
        item["amount"] = float(item["amount"])
        item["balance_after"] = float(item["balance_after"])
        if item.get("created_at"):
            item["created_at"] = str(item["created_at"])
            
    return {"ok": True, "data": ledger_entries}

@api_router.post("/webhooks/bank/{provider}")
async def handle_bank_webhook(provider: str, request: Request):
    try:
        payload = await request.json()
    except Exception:
        payload = {}
        
    import json, random
    
    event_type = payload.get("event_type") or payload.get("event") or "payment.received"
    tx_id = payload.get("transaction_id") or payload.get("payment_id") or f"PAY_{random.randint(100000, 999999)}"
    va_num = payload.get("virtual_account_number") or payload.get("account_number") or payload.get("virtual_account_id") or ""
    utr_val = payload.get("utr") or payload.get("reference_number") or ""
    amount_val = float(payload.get("amount") or payload.get("amount_paid") or 0.0)
    payment_id = payload.get("payment_id") or tx_id
    
    hook_id = await db_execute(
        """INSERT INTO payment_webhooks 
           (provider, event_type, transaction_id, virtual_account_id, payment_id, utr, amount, status, signature_verified, payload, processed)
           VALUES (%s, %s, %s, %s, %s, %s, %s, %s, 1, %s, 0)""",
        (provider, event_type, tx_id, va_num, payment_id, utr_val, amount_val, payload.get("status", "SUCCESS"), json.dumps(payload))
    )
    
    va = None
    if va_num:
        va = await db_fetchone("SELECT * FROM virtual_accounts WHERE virtual_account_number = %s OR virtual_account_id = %s OR virtual_vpa = %s", (va_num, va_num, va_num))
    
    if not va or not utr_val or amount_val <= 0:
        discrepancy_reason = "Virtual account not found" if not va else "Missing UTR or zero amount"
        rec_id = await db_execute(
            """INSERT INTO bank_reconciliations 
               (payment_id, utr, virtual_account_id, account_id, amount, bank_amount, wallet_amount, status, discrepancy_reason, supporting_evidence)
               VALUES (%s, %s, %s, %s, %s, %s, 0, 'DISCREPANCY', %s, %s)""",
            (payment_id, utr_val or "MISSING_UTR", va_num, va["account_id"] if va else None, amount_val, amount_val, discrepancy_reason, f"Provider: {provider}")
        )
        await db_execute("UPDATE payment_webhooks SET processed = 1, error_log = %s WHERE id = %s", (discrepancy_reason, hook_id))
        return {"ok": True, "status": "DISCREPANCY", "message": discrepancy_reason}
        
    existing_rec = await db_fetchone("SELECT id, status FROM bank_reconciliations WHERE payment_id = %s OR utr = %s", (payment_id, utr_val))
    if existing_rec:
        reason = "Duplicate webhook or UTR already recorded"
        rec_id = await db_execute(
            """INSERT INTO bank_reconciliations 
               (payment_id, utr, virtual_account_id, account_id, amount, bank_amount, wallet_amount, status, discrepancy_reason, supporting_evidence)
               VALUES (%s, %s, %s, %s, %s, %s, 0, 'DISCREPANCY', %s, %s)""",
            (f"{payment_id}_dup", utr_val, va_num, va["account_id"], amount_val, amount_val, reason, f"Matches existing reconciliation {existing_rec['id']}")
        )
        await db_execute("UPDATE payment_webhooks SET processed = 1, error_log = %s WHERE id = %s", (reason, hook_id))
        return {"ok": True, "status": "DISCREPANCY", "message": reason}
        
    wallet = await db_fetchone("SELECT id, balance FROM wallets WHERE account_id = %s", (va["account_id"],))
    if not wallet:
        w_id = await db_execute("INSERT INTO wallets (account_id, balance, reward_points) VALUES (%s, 0, 0)", (va["account_id"],))
        wallet = {"id": w_id, "balance": 0.0}
        
    curr_balance = float(wallet["balance"])
    new_balance = curr_balance + amount_val
    
    await db_execute("UPDATE wallets SET balance = %s WHERE id = %s", (new_balance, wallet["id"]))
    
    ledger_id = await db_execute(
        """INSERT INTO wallet_ledger 
           (wallet_id, account_id, transaction_type, amount, balance_after, reference_id, payment_id, utr, description, status)
           VALUES (%s, %s, 'Credit', %s, %s, %s, %s, %s, 'Virtual Account Topup', 'SUCCESS')""",
        (wallet["id"], va["account_id"], amount_val, new_balance, hook_id, payment_id, utr_val)
    )
    
    rec_id = await db_execute(
        """INSERT INTO bank_reconciliations 
           (payment_id, utr, virtual_account_id, account_id, amount, bank_amount, wallet_amount, status, reconciled_at)
           VALUES (%s, %s, %s, %s, %s, %s, %s, 'MATCHED', NOW())""",
        (payment_id, utr_val, va["virtual_account_number"], va["account_id"], amount_val, amount_val, amount_val)
    )
    
    await db_execute("UPDATE wallet_topups SET status = 'SUCCESS' WHERE account_id = %s AND status = 'PENDING' AND amount = %s LIMIT 1", (va["account_id"], amount_val))
    await db_execute("UPDATE payment_webhooks SET processed = 1 WHERE id = %s", (hook_id,))
    
    return {"ok": True, "status": "MATCHED", "message": "Payment credited and reconciled successfully.", "balance": new_balance}

class BankIntegrationIn(BaseModel):
    association_id: Optional[str] = None
    provider_name: str
    client_id: Optional[str] = None
    client_secret: Optional[str] = None
    api_key: Optional[str] = None
    api_secret: Optional[str] = None
    webhook_url: Optional[str] = None
    webhook_secret: Optional[str] = None
    base_url: Optional[str] = None

@api_router.get("/admin/bank-integrations")
async def list_bank_integrations(account: dict = Depends(get_current_account)):
    items = await db_fetchall("SELECT * FROM bank_integrations ORDER BY created_at DESC")
    for i in items:
        if i.get("created_at"): i["created_at"] = str(i["created_at"])
    return {"ok": True, "data": items}

@api_router.post("/admin/bank-integrations")
async def create_bank_integration(payload: BankIntegrationIn, account: dict = Depends(get_current_account)):
    new_id = await db_execute(
        """INSERT INTO bank_integrations (association_id, provider_name, client_id, client_secret, api_key, api_secret, webhook_url, webhook_secret, base_url)
           VALUES (%s, %s, %s, %s, %s, %s, %s, %s, %s)""",
        (payload.association_id, payload.provider_name, payload.client_id, payload.client_secret, payload.api_key, payload.api_secret, payload.webhook_url, payload.webhook_secret, payload.base_url)
    )
    return {"ok": True, "message": "Bank integration saved successfully.", "id": new_id}

class CollectionAccountIn(BaseModel):
    association_id: Optional[str] = None
    bank_name: str
    account_number: str
    ifsc: str
    account_holder_name: str
    account_type: Optional[str] = "Current"
    vpa: Optional[str] = None

@api_router.get("/admin/collection-accounts")
async def list_collection_accounts(account: dict = Depends(get_current_account)):
    items = await db_fetchall("SELECT * FROM collection_accounts ORDER BY created_at DESC")
    for i in items:
        if i.get("created_at"): i["created_at"] = str(i["created_at"])
    return {"ok": True, "data": items}

@api_router.post("/admin/collection-accounts")
async def create_collection_account(payload: CollectionAccountIn, account: dict = Depends(get_current_account)):
    new_id = await db_execute(
        """INSERT INTO collection_accounts (association_id, bank_name, account_number, ifsc, account_holder_name, account_type, vpa)
           VALUES (%s, %s, %s, %s, %s, %s, %s)""",
        (payload.association_id, payload.bank_name, payload.account_number, payload.ifsc, payload.account_holder_name, payload.account_type, payload.vpa)
    )
    return {"ok": True, "message": "Collection account created successfully.", "id": new_id}

@api_router.get("/admin/reconciliation")
async def list_reconciliations(status: Optional[str] = None, account: dict = Depends(require_permission("Reconciliation", "can_view"))):
    query = """
        SELECT r.*, a.email as user_email, COALESCE(ud.name, a.email) as user_name
        FROM bank_reconciliations r
        LEFT JOIN accounts a ON r.account_id = a.account_id
        LEFT JOIN user_details ud ON a.user_id = ud.user_id
    """
    params = []
    if status:
        query += " WHERE r.status = %s"
        params.append(status)
        
    query += " ORDER BY r.created_at DESC"
    rows = await db_fetchall(query, tuple(params))
    for r in rows:
        r["amount"] = float(r["amount"])
        if r.get("bank_amount") is not None: r["bank_amount"] = float(r["bank_amount"])
        if r.get("wallet_amount") is not None: r["wallet_amount"] = float(r["wallet_amount"])
        if r.get("created_at"): r["created_at"] = str(r["created_at"])
        if r.get("reconciled_at"): r["reconciled_at"] = str(r["reconciled_at"])
        
    return {"ok": True, "data": rows}

class ReconciliationOverrideIn(BaseModel):
    reconciliation_id: str
    action: str
    reason: str
    comments: str
    supporting_ref: Optional[str] = None

@api_router.post("/admin/reconciliation/override")
async def override_reconciliation(payload: ReconciliationOverrideIn, account: dict = Depends(require_permission("Reconciliation", "can_view"))):
    if not payload.reason.strip() or not payload.comments.strip():
        raise HTTPException(status_code=400, detail="Override Reason and Comments are mandatory.")
        
    rec = await db_fetchone("SELECT * FROM bank_reconciliations WHERE id = %s", (payload.reconciliation_id,))
    if not rec:
        raise HTTPException(status_code=404, detail="Reconciliation record not found.")
        
    old_status = rec["status"]
    if payload.action == "OVERRIDE_APPROVE":
        new_status = "MATCHED"
    elif payload.action == "OVERRIDE_REJECT":
        new_status = "REJECTED"
    elif payload.action == "REQUEST_INVESTIGATION":
        new_status = "UNDER_INVESTIGATION"
    else:
        raise HTTPException(status_code=400, detail="Invalid action specified.")
        
    if payload.action == "OVERRIDE_APPROVE" and rec.get("account_id"):
        account_id = rec["account_id"]
        wallet = await db_fetchone("SELECT id, balance FROM wallets WHERE account_id = %s", (account_id,))
        if not wallet:
            w_id = await db_execute("INSERT INTO wallets (account_id, balance, reward_points) VALUES (%s, 0, 0)", (account_id,))
            wallet = {"id": w_id, "balance": 0.0}
            
        curr_balance = float(wallet["balance"])
        amount_to_credit = float(rec["amount"])
        new_balance = curr_balance + amount_to_credit
        
        await db_execute("UPDATE wallets SET balance = %s WHERE id = %s", (new_balance, wallet["id"]))
        
        ledger_id = await db_execute(
            """INSERT INTO wallet_ledger 
               (wallet_id, account_id, transaction_type, amount, balance_after, reference_id, payment_id, utr, description, status)
               VALUES (%s, %s, 'Credit', %s, %s, %s, %s, %s, %s, 'SUCCESS')""",
            (wallet["id"], account_id, amount_to_credit, new_balance, payload.supporting_ref or rec["id"], rec["payment_id"], rec["utr"], f"Manual Override: {payload.reason}")
        )
        
        await db_execute("UPDATE bank_reconciliations SET wallet_amount = %s, reconciled_at = NOW() WHERE id = %s", (amount_to_credit, rec["id"]))
        
    await db_execute("UPDATE bank_reconciliations SET status = %s WHERE id = %s", (new_status, rec["id"]))
    
    performed_by = account.get("email") or account.get("account_id")
    audit_id = await db_execute(
        """INSERT INTO reconciliation_audits 
           (reconciliation_id, action, old_status, new_status, reason, comments, supporting_ref, performed_by)
           VALUES (%s, %s, %s, %s, %s, %s, %s, %s)""",
        (rec["id"], payload.action, old_status, new_status, payload.reason, payload.comments, payload.supporting_ref or "", performed_by)
    )
    
    return {"ok": True, "message": f"Reconciliation override successful. Status updated to {new_status}.", "old_status": old_status, "new_status": new_status, "audit_id": audit_id}

@api_router.get("/admin/reconciliation/{rec_id}/audits")
async def get_reconciliation_audits(rec_id: str, account: dict = Depends(require_permission("Reconciliation", "can_view"))):
    audits = await db_fetchall("SELECT * FROM reconciliation_audits WHERE reconciliation_id = %s ORDER BY performed_at DESC", (rec_id,))
    for a in audits:
        if a.get("performed_at"): a["performed_at"] = str(a["performed_at"])
    return {"ok": True, "data": audits}


# ---------- Bank Accounts API ----------

class BankAccountCreateUpdate(BaseModel):
    association_id: str
    account_name: str
    account_holder_name: str
    bank_name: str
    account_number: str
    ifsc_code: str
    branch_name: Optional[str] = None
    account_type: str
    currency: str = "INR"
    upi_id: Optional[str] = None
    qr_code_url: Optional[str] = None
    gateway_provider: Optional[str] = None
    merchant_id: Optional[str] = None
    api_key: Optional[str] = None
    api_secret: Optional[str] = None
    webhook_secret: Optional[str] = None
    is_default: Optional[bool] = False
    status: Optional[str] = "Active"

async def verify_bank_access(account, association_id):
    if not association_id:
        return False
    if account["role"] == "Super admin":
        return True
    res = await db_fetchone("SELECT 1 FROM admin_associations WHERE admin_id=%s AND association_id=%s", (account["account_id"], association_id))
    return bool(res)

@api_router.get("/admin/bank-accounts")
async def get_bank_accounts(request: Request, association_id: Optional[str] = None, account: dict = Depends(require_role(["Super admin", "Admin", "Accountant"]))):
    if account["role"] == "Super admin":
        query = """
            SELECT b.*, a.name as association_name
            FROM association_bank_accounts b
            LEFT JOIN associations a ON b.association_id = a.id
        """
        params = []
        if association_id:
            query += " WHERE b.association_id = %s"
            params.append(association_id)
        query += " ORDER BY b.created_at DESC"
        params = tuple(params)
    else:
        query = """
            SELECT b.*, a.name as association_name
            FROM association_bank_accounts b
            JOIN admin_associations aa ON b.association_id = aa.association_id
            LEFT JOIN associations a ON b.association_id = a.id
            WHERE aa.admin_id = %s
        """
        params = [account["account_id"]]
        if association_id:
            query += " AND b.association_id = %s"
            params.append(association_id)
        query += " ORDER BY b.created_at DESC"
        params = tuple(params)
    
    rows = await db_fetchall(query, params)
    
    # Mask sensitive data before returning to frontend
    for r in rows:
        dec_acc = decrypt_data(r.get("account_number"))
        if dec_acc and len(dec_acc) > 4:
            r["account_number"] = f"********{dec_acc[-4:]}"
        else:
            r["account_number"] = "****"
            
        r["api_key"] = "****" if r.get("api_key") else None
        r["api_secret"] = "****" if r.get("api_secret") else None
        r["webhook_secret"] = "****" if r.get("webhook_secret") else None

        for k, v in list(r.items()):
            if isinstance(v, (date, datetime)):
                r[k] = str(v)
        
    return rows

@api_router.post("/admin/bank-accounts")
async def create_bank_account(request: Request, payload: BankAccountCreateUpdate, account: dict = Depends(require_role(["Super admin", "Admin", "Accountant"]))):
    if not await verify_bank_access(account, payload.association_id):
        raise HTTPException(status_code=403, detail="Unauthorized for this association")
    enc_acc = encrypt_data(payload.account_number)
    enc_api = encrypt_data(payload.api_key)
    enc_sec = encrypt_data(payload.api_secret)
    enc_wh = encrypt_data(payload.webhook_secret)
    
    new_id = await db_execute(
        """INSERT INTO association_bank_accounts (
            association_id, account_name, account_holder_name, bank_name,
            account_number, ifsc_code, branch_name, account_type, currency,
            upi_id, qr_code_url, gateway_provider, merchant_id,
            api_key, api_secret, webhook_secret, is_default, status, created_by
        ) VALUES (%s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s)""",
        (
            payload.association_id, payload.account_name, payload.account_holder_name,
            payload.bank_name, enc_acc, payload.ifsc_code, payload.branch_name, payload.account_type,
            payload.currency, payload.upi_id, payload.qr_code_url, payload.gateway_provider,
            payload.merchant_id, enc_api, enc_sec, enc_wh, payload.is_default, payload.status, account["account_id"]
        )
    )
    
    # Audit log
    await log_audit(
        entity_name="association_bank_accounts",
        entity_id=new_id,
        action="CREATE",
        changes={"account_name": payload.account_name, "bank_name": payload.bank_name},
        performed_by=account["account_id"]
    )
    
    return {"ok": True, "id": new_id, "message": "Bank account created securely."}

@api_router.put("/admin/bank-accounts/{account_id}")
async def update_bank_account(request: Request, account_id: str, payload: BankAccountCreateUpdate, account: dict = Depends(require_role(["Super admin", "Admin", "Accountant"]))):
    existing = await db_fetchone("SELECT * FROM association_bank_accounts WHERE id = %s", (account_id,))
    if not existing:
        raise HTTPException(status_code=404, detail="Bank account not found")
        
    if not await verify_bank_access(account, existing["association_id"]):
        raise HTTPException(status_code=403, detail="Unauthorized to modify this bank account")
        
    if payload.association_id != existing["association_id"] and not await verify_bank_access(account, payload.association_id):
        raise HTTPException(status_code=403, detail="Unauthorized for target association")
        
    # Only update sensitive fields if they are actually provided and not masked
    enc_acc = existing["account_number"]
    if payload.account_number and payload.account_number != "****" and not payload.account_number.startswith("****"):
        enc_acc = encrypt_data(payload.account_number)
        
    enc_api = existing["api_key"]
    if payload.api_key and payload.api_key != "****":
        enc_api = encrypt_data(payload.api_key)
        
    enc_sec = existing["api_secret"]
    if payload.api_secret and payload.api_secret != "****":
        enc_sec = encrypt_data(payload.api_secret)
        
    enc_wh = existing["webhook_secret"]
    if payload.webhook_secret and payload.webhook_secret != "****":
        enc_wh = encrypt_data(payload.webhook_secret)

    await db_execute(
        """UPDATE association_bank_accounts SET
            association_id=%s, account_name=%s, account_holder_name=%s, bank_name=%s,
            account_number=%s, ifsc_code=%s, branch_name=%s, account_type=%s, currency=%s,
            upi_id=%s, qr_code_url=%s, gateway_provider=%s, merchant_id=%s,
            api_key=%s, api_secret=%s, webhook_secret=%s, is_default=%s, status=%s, updated_by=%s
        WHERE id=%s""",
        (
            payload.association_id, payload.account_name, payload.account_holder_name,
            payload.bank_name, enc_acc, payload.ifsc_code, payload.branch_name, payload.account_type,
            payload.currency, payload.upi_id, payload.qr_code_url, payload.gateway_provider,
            payload.merchant_id, enc_api, enc_sec, enc_wh, payload.is_default, payload.status,
            account["account_id"], account_id
        )
    )
    
    await log_audit(
        entity_name="association_bank_accounts",
        entity_id=account_id,
        action="UPDATE",
        changes={"account_name": payload.account_name, "status": payload.status},
        performed_by=account["account_id"]
    )
    
    return {"ok": True, "message": "Bank account updated securely."}

@api_router.delete("/admin/bank-accounts/{account_id}")
async def delete_bank_account(request: Request, account_id: str, account: dict = Depends(require_role(["Super admin", "Admin", "Accountant"]))):
    existing = await db_fetchone("SELECT * FROM association_bank_accounts WHERE id = %s", (account_id,))
    if not existing:
        raise HTTPException(status_code=404, detail="Bank account not found")
        
    if not await verify_bank_access(account, existing["association_id"]):
        raise HTTPException(status_code=403, detail="Unauthorized to delete this bank account")
        
    await db_execute("DELETE FROM association_bank_accounts WHERE id = %s", (account_id,))
    
    await log_audit(
        entity_name="association_bank_accounts",
        entity_id=account_id,
        action="DELETE",
        changes={"account_name": existing["account_name"]},
        performed_by=account["account_id"]
    )
    
    return {"ok": True, "message": "Bank account deleted successfully."}


# ---------- Financial Reports API ----------

class FinancialReportCreate(BaseModel):
    association_id: str
    published_month: str
    report_type: str
    title: str
    file_url: str

@api_router.get("/admin/financial-reports")
async def get_financial_reports(request: Request, account: dict = Depends(require_role(["Super admin"]))):
    # Get all reports and also include association name
    query = """
        SELECT fr.*, a.name as association_name 
        FROM financial_reports fr
        LEFT JOIN associations a ON fr.association_id = a.id
        ORDER BY fr.created_at DESC
    """
    rows = await db_fetchall(query)
    return rows

@api_router.post("/admin/financial-reports")
async def create_financial_report(request: Request, payload: FinancialReportCreate, account: dict = Depends(require_role(["Super admin"]))):
    new_id = await db_execute(
        """INSERT INTO financial_reports (
            association_id, published_month, report_type, title, file_url, uploaded_by
        ) VALUES (%s, %s, %s, %s, %s, %s)""",
        (
            payload.association_id, payload.published_month, 
            payload.report_type, payload.title, payload.file_url, account["account_id"]
        )
    )
    
    return {"ok": True, "id": new_id, "message": "Financial report uploaded successfully."}


# ---------- Chart of Accounts API ----------
import openpyxl
from io import BytesIO
from fastapi import Query
from fastapi.responses import JSONResponse, StreamingResponse

class AssociationCOACreate(BaseModel):
    gl_code: str
    gl_name: str
    structure: str
    grouping: str

@api_router.post("/admin/global-coa/upload")
async def upload_global_coa(request: Request, file: UploadFile = File(...), account: dict = Depends(require_role(["Super admin", "Admin"]))):
    content = await file.read()
    
    # Process Excel file
    try:
        workbook = openpyxl.load_workbook(filename=BytesIO(content), data_only=True)
        sheet = workbook.active
        
        headers = [cell.value.lower().replace(" ", "_") if cell.value else "" for cell in sheet[1]]
        try:
            code_idx = headers.index("gl_code")
        except ValueError:
            code_idx = 0
        try:
            name_idx = headers.index("gl_name")
        except ValueError:
            name_idx = 1
        try:
            structure_idx = headers.index("structure")
        except ValueError:
            structure_idx = 2
        try:
            grouping_idx = headers.index("grouping")
        except ValueError:
            grouping_idx = 3

        count = 0
        for row in sheet.iter_rows(min_row=2, values_only=True):
            if row[code_idx] and row[name_idx]:
                raw_code = str(row[code_idx])
                if raw_code.endswith('.0'):
                    raw_code = raw_code[:-2]
                await db_execute(
                    "INSERT INTO global_chart_of_accounts (gl_code, gl_name, structure, `grouping`) VALUES (%s, %s, %s, %s)",
                    (raw_code, str(row[name_idx]), str(row[structure_idx] or ""), str(row[grouping_idx] or ""))
                )
                count += 1
        
        return {"ok": True, "message": f"Successfully uploaded {count} global chart of accounts."}
    except Exception as e:
        return JSONResponse(status_code=400, content={"detail": f"Error parsing Excel file: {str(e)}"})

@api_router.get("/admin/global-coa")
async def get_global_coa(request: Request, account: dict = Depends(require_role(["Super admin", "Admin", "Accountant"]))):
    rows = await db_fetchall("SELECT * FROM global_chart_of_accounts ORDER BY `grouping`, gl_code")
    return rows

@api_router.post("/accountant/association-coa/map")
async def map_global_coa(request: Request, association_id: str = Query(..., description="Association ID to map"), account: dict = Depends(require_role(["Accountant", "Super admin", "Admin"]))):
    # Check if this association already mapped
    existing = await db_fetchone("SELECT COUNT(*) as count FROM association_chart_of_accounts WHERE association_id = %s", (association_id,))
    if existing and existing["count"] > 0:
        return JSONResponse(status_code=400, content={"detail": "Chart of Accounts is already mapped for this association."})
        
    global_accounts = await db_fetchall("SELECT * FROM global_chart_of_accounts")
    if not global_accounts:
        return JSONResponse(status_code=400, content={"detail": "Global Chart of Accounts is empty. Please contact Super Admin."})
        
    import uuid
    values = []
    params = []
    for g_acc in global_accounts:
        values.append("(%s, %s, %s, %s, %s, %s, %s)")
        params.extend([str(uuid.uuid4()), association_id, g_acc["gl_code"], g_acc["gl_name"], g_acc["structure"], g_acc["grouping"], g_acc["id"]])
        
    if values:
        query = f"INSERT INTO association_chart_of_accounts (id, association_id, gl_code, gl_name, structure, `grouping`, mapped_from_global_id) VALUES {','.join(values)}"
        async with pool.acquire() as conn:
            async with conn.cursor() as cur:
                await cur.execute(query, tuple(params))
        
    return {"ok": True, "message": f"Successfully mapped {len(global_accounts)} accounts."}

@api_router.get("/accountant/association-coa")
async def get_association_coa(request: Request, association_id: str = Query(...), account: dict = Depends(require_role(["Accountant", "Super admin", "Admin"]))):
    rows = await db_fetchall("SELECT * FROM association_chart_of_accounts WHERE association_id = %s ORDER BY `grouping`, gl_code", (association_id,))
    return rows

@api_router.get("/accountant/association-coa/status")
async def get_association_coa_status(request: Request, account: dict = Depends(require_role(["Accountant", "Super admin", "Admin"]))):
    rows = await db_fetchall("SELECT association_id, COUNT(*) as count FROM association_chart_of_accounts GROUP BY association_id")
    status = {row["association_id"]: row["count"] for row in rows}
    return status

@api_router.get("/accountant/association-coa/download")
async def download_association_coa(request: Request, association_id: str = Query(...), account: dict = Depends(require_role(["Accountant", "Super admin", "Admin"]))):
    rows = await db_fetchall("SELECT gl_code, gl_name, structure, `grouping` FROM association_chart_of_accounts WHERE association_id = %s ORDER BY gl_code", (association_id,))
    
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Chart of Accounts"
    ws.append(["GL Code", "GL Name", "Structure", "Grouping"])
    
    for row in rows:
        ws.append([row["gl_code"], row["gl_name"], row["structure"], row["grouping"]])
        
    stream = BytesIO()
    wb.save(stream)
    stream.seek(0)
    
    assoc = await db_fetchone("SELECT name FROM associations WHERE id = %s", (association_id,))
    safe_name = assoc["name"].replace("/", "-").replace("\\", "-") if assoc else "Association"
    filename = f"{safe_name} - Chart of Accounts.xlsx"
    
    return StreamingResponse(
        stream,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'}
    )

@api_router.post("/accountant/association-coa")
async def add_association_coa(request: Request, payload: AssociationCOACreate, association_id: str = Query(...), account: dict = Depends(require_role(["Accountant", "Super admin", "Admin"]))):
    existing = await db_fetchone("SELECT id FROM association_chart_of_accounts WHERE association_id = %s AND gl_code = %s", (association_id, payload.gl_code))
    if existing:
        return JSONResponse(status_code=400, content={"detail": f"Error: GL Code '{payload.gl_code}' already exists for this association."})

    await db_execute(
        "INSERT INTO association_chart_of_accounts (association_id, gl_code, gl_name, structure, `grouping`) VALUES (%s, %s, %s, %s, %s)",
        (association_id, payload.gl_code, payload.gl_name, payload.structure, payload.grouping)
    )
    return {"ok": True, "message": "Chart of Account added successfully."}

@api_router.put("/accountant/association-coa/{coa_id}")
async def update_association_coa(request: Request, coa_id: str, payload: AssociationCOACreate, account: dict = Depends(require_role(["Accountant", "Super admin", "Admin"]))):
    existing = await db_fetchone("SELECT id FROM association_chart_of_accounts WHERE id != %s AND association_id = (SELECT association_id FROM association_chart_of_accounts WHERE id = %s) AND gl_code = %s", (coa_id, coa_id, payload.gl_code))
    if existing:
        return JSONResponse(status_code=400, content={"detail": f"Error: GL Code '{payload.gl_code}' already exists for this association."})

    await db_execute(
        "UPDATE association_chart_of_accounts SET gl_code = %s, gl_name = %s, structure = %s, `grouping` = %s WHERE id = %s",
        (payload.gl_code, payload.gl_name, payload.structure, payload.grouping, coa_id)
    )
    return {"ok": True, "message": "Chart of Account updated successfully."}

@api_router.delete("/accountant/association-coa/{coa_id}")
async def delete_association_coa(request: Request, coa_id: str, account: dict = Depends(require_role(["Accountant", "Super admin", "Admin"]))):
    await db_execute("DELETE FROM association_chart_of_accounts WHERE id = %s", (coa_id,))
    return {"ok": True, "message": "Chart of Account deleted successfully."}

# ---------- Budget Management API ----------
from fastapi import Form, File, UploadFile
import shutil
import json

@api_router.get("/accountant/budgets")
async def get_budgets(request: Request, association_id: str = Query(None), account: dict = Depends(require_role(["Accountant", "Super admin", "Admin"]))):
    if association_id and association_id != "ALL":
        query = """
            SELECT b.*, a.name as association_name 
            FROM association_budgets b
            JOIN associations a ON b.association_id = a.id
            WHERE b.association_id = %s
            ORDER BY b.created_at DESC
        """
        rows = await db_fetchall(query, (association_id,))
    else:
        query = """
            SELECT b.*, a.name as association_name 
            FROM association_budgets b
            JOIN associations a ON b.association_id = a.id
            ORDER BY b.created_at DESC
        """
        rows = await db_fetchall(query)
    
    for row in rows:
        row["financial_year_start"] = row["financial_year_start"].isoformat() if row["financial_year_start"] else None
        row["financial_year_end"] = row["financial_year_end"].isoformat() if row["financial_year_end"] else None
        row["created_at"] = row["created_at"].isoformat() if row["created_at"] else None
        
    return rows

@api_router.post("/accountant/budgets")
async def upload_budget(
    request: Request,
    association_id: str = Form(...),
    budget_type: str = Form(...),
    financial_year_start: str = Form(...),
    financial_year_end: str = Form(...),
    file: UploadFile = File(...),
    account: dict = Depends(require_role(["Accountant", "Super admin", "Admin"]))
):
    existing = await db_fetchone("""
        SELECT id FROM association_budgets 
        WHERE association_id = %s AND budget_type = %s AND financial_year_start = %s AND financial_year_end = %s
    """, (association_id, budget_type, financial_year_start, financial_year_end))
    
    if existing:
        return JSONResponse(status_code=400, content={"detail": f"{budget_type} for this Financial Year already exists."})
        
    file_ext = file.filename.split(".")[-1] if "." in file.filename else "xlsx"
    file_name = f"budget_{uuid.uuid4()}.{file_ext}"
    file_path = UPLOAD_DIR / file_name
    
    UPLOAD_DIR.mkdir(parents=True, exist_ok=True)
    
    with open(file_path, "wb") as buffer:
        shutil.copyfileobj(file.file, buffer)
        
    file_url = f"/uploads/{file_name}"
    
    import uuid as uuid_pkg
    budget_id = str(uuid_pkg.uuid4())
    await db_execute("""
        INSERT INTO association_budgets (id, association_id, budget_type, financial_year_start, financial_year_end, file_url)
        VALUES (%s, %s, %s, %s, %s, %s)
    """, (budget_id, association_id, budget_type, financial_year_start, financial_year_end, file_url))
    
    try:
        import openpyxl
        from io import BytesIO
        
        with open(file_path, "rb") as f:
            content = f.read()
            
        workbook = openpyxl.load_workbook(filename=BytesIO(content), data_only=True)
        sheet = workbook.active
        
        headers = []
        
        for i, row in enumerate(sheet.iter_rows(values_only=True)):
            if i == 0:
                headers = [str(cell) if cell is not None else f"Column_{j}" for j, cell in enumerate(row)]
                continue
                
            if not any(cell is not None and str(cell).strip() != "" for cell in row):
                continue
                
            row_data = {}
            for j, cell in enumerate(row):
                if j < len(headers):
                    val = cell
                    if hasattr(val, 'isoformat'):
                        val = val.isoformat()
                    row_data[headers[j]] = val
            
            item_id = str(uuid_pkg.uuid4())
            await db_execute("""
                INSERT INTO budget_items (id, budget_id, row_data)
                VALUES (%s, %s, %s)
            """, (item_id, budget_id, json.dumps(row_data)))
            
    except Exception as e:
        print(f"Error parsing budget excel: {e}")
        
    return {"ok": True, "message": "Budget uploaded successfully"}


class BoardChatMessageIn(BaseModel):
    message: str
    attachment_url: Optional[str] = None

@api_router.get("/board_chat/{pool_type}/{pool_id}")
async def get_board_chat(pool_type: str, pool_id: str, assoc_id: str, account: dict = Depends(get_current_account)):
    query = '''
        SELECT c.*, acc.email, COALESCE(ud.name, emp.name, acc.email) as sender_name,
               ud.profile_pic_url,
               (c.sender_id = %s) as is_mine
        FROM board_committee_chat c
        JOIN accounts acc ON c.sender_id = acc.account_id
        LEFT JOIN user_details ud ON acc.user_id = ud.user_id
        LEFT JOIN employees emp ON acc.employee_id = emp.employee_id
        WHERE c.association_id = %s AND c.pool_type = %s AND c.pool_id = %s
        ORDER BY c.created_at ASC
    '''
    messages = await db_fetchall(query, (account["account_id"], assoc_id, pool_type, pool_id))
    for msg in messages:
        if msg.get('created_at'): msg['created_at'] = str(msg['created_at'])
    return {"ok": True, "data": messages}

@api_router.post("/board_chat/{pool_type}/{pool_id}")
async def send_board_chat(pool_type: str, pool_id: str, assoc_id: str, payload: BoardChatMessageIn, account: dict = Depends(get_current_account)):
    await db_execute(
        "INSERT INTO board_committee_chat (association_id, pool_type, pool_id, sender_id, message, attachment_url) VALUES (%s, %s, %s, %s, %s, %s)",
        (assoc_id, pool_type, pool_id, account["account_id"], payload.message, payload.attachment_url)
    )
    return {"ok": True}



@api_router.get("/user/committees")
async def get_user_committees(assoc_id: Optional[str] = None, account: dict = Depends(get_current_account)):
    if account["role"] in ["Committee Member", "Board member"]:
        query = '''
            SELECT c.*
            FROM committees c
            JOIN committee_members cm ON c.id = cm.committee_id
            WHERE cm.user_id = %s
        '''
        params = [account["user_id"]]
        if assoc_id and assoc_id != 'ALL':
            query += " AND c.association_id = %s"
            params.append(assoc_id)
            
        committees = await db_fetchall(query, tuple(params))
        return {"ok": True, "data": committees}
    return {"ok": True, "data": []}


class AmenityCreate(BaseModel):
    name: str
    charges: float
    status: bool = True

class AmenityUpdateStatus(BaseModel):
    status: bool

class AmenityBookingCreate(BaseModel):
    booking_date: str
    start_time: str
    end_time: str
    duration_hours: int
    payment_method: str = "wallet"
    pin: str = None

@api_router.get("/admin/associations/{id}/amenities")
async def get_admin_amenities(id: str, account: dict = Depends(require_role(["Super admin", "Admin"]))):
    amenities = await db_fetchall("SELECT * FROM amenities WHERE association_id = %s", (id,))
    return amenities

@api_router.post("/admin/associations/{id}/amenities")
async def create_amenity(id: str, payload: AmenityCreate, account: dict = Depends(require_role(["Super admin", "Admin"]))):
    new_id = await db_execute("INSERT INTO amenities (association_id, name, charges, status) VALUES (%s, %s, %s, %s)",
                     (id, payload.name, payload.charges, payload.status))
    return {"ok": True, "id": new_id, "message": "Amenity added successfully"}

@api_router.put("/admin/associations/{id}/amenities/{amenity_id}")
async def update_amenity_status(id: str, amenity_id: str, payload: AmenityUpdateStatus, account: dict = Depends(require_role(["Super admin", "Admin"]))):
    await db_execute("UPDATE amenities SET status = %s WHERE id = %s AND association_id = %s", (payload.status, amenity_id, id))
    return {"ok": True, "message": "Amenity status updated"}

@api_router.get("/admin/associations/{id}/amenity-bookings")
async def get_admin_amenity_bookings(id: str, account: dict = Depends(require_role(["Super admin", "Admin", "Board Member"]))):
    query = '''
        SELECT b.*, a.name as amenity_name, assoc.name as association_name, ud.contact_number as contact_no, COALESCE(NULLIF(b.homeowner_name, ''), ud.name, acc.email) as computed_homeowner_name, COALESCE(NULLIF(b.unit_number, ''), un.unit_number) as computed_unit_number
        FROM amenity_bookings b 
        JOIN amenities a ON b.amenity_id = a.id
        JOIN associations assoc ON b.association_id = assoc.id
        LEFT JOIN accounts acc ON b.user_id = acc.account_id
        LEFT JOIN user_details ud ON acc.user_id = ud.user_id
        LEFT JOIN units un ON ud.unit_id = un.id
    '''
    if id == "ALL":
        bookings = await db_fetchall(query + ' ORDER BY b.booking_date DESC')
    else:
        bookings = await db_fetchall(query + ' WHERE b.association_id = %s ORDER BY b.booking_date DESC', (id,))
        
    for b in bookings:
        if not b.get("homeowner_name"):
            b["homeowner_name"] = b.get("computed_homeowner_name")
        if not b.get("unit_number"):
            b["unit_number"] = b.get("computed_unit_number")
            
    return bookings

@api_router.get("/associations/{id}/amenities")
async def get_active_amenities(id: str, account: dict = Depends(get_current_account)):
    amenities = await db_fetchall("SELECT * FROM amenities WHERE association_id = %s AND status = 1", (id,))
    return amenities

@api_router.get("/amenities/{amenity_id}/bookings")
async def get_amenity_bookings(amenity_id: str, month: str = None, account: dict = Depends(get_current_account)):
    # Returns all confirmed bookings for this amenity. If month is provided (YYYY-MM), filter by month.
    query = "SELECT DISTINCT booking_date, start_time, end_time, duration_hours, homeowner_name, unit_number FROM amenity_bookings WHERE amenity_id = %s AND payment_status = 'Confirmed'"
    params = [amenity_id]
    if month:
        query += " AND booking_date LIKE %s"
        params.append(f"{month}%")
    bookings = await db_fetchall(query, tuple(params))
    return bookings

@api_router.get("/amenities/my-bookings")
async def get_my_amenity_bookings(account: dict = Depends(get_current_account)):
    user = await db_fetchone('''
        SELECT ud.unit_id, u.unit_number as unit_number, blk.association_id 
        FROM user_details ud 
        LEFT JOIN units u ON ud.unit_id = u.id 
        LEFT JOIN blocks blk ON u.block_id = blk.id
        WHERE ud.user_id = %s
    ''', (account.get("user_id"),))
    
    unit_id = user.get("unit_id") if user and str(user.get("unit_id") or "").strip() else None
    unit_number = user.get("unit_number") if user and str(user.get("unit_number") or "").strip() else None
    assoc_id = user.get("association_id") if user else None
    
    bookings = await db_fetchall('''
        SELECT b.*, a.name as amenity_name, assoc.name as association_name
        FROM amenity_bookings b 
        JOIN amenities a ON b.amenity_id = a.id
        JOIN associations assoc ON b.association_id = assoc.id
        LEFT JOIN accounts acc ON b.user_id = acc.account_id
        LEFT JOIN user_details ud ON acc.user_id = ud.user_id
        WHERE b.user_id = %s 
           OR (ud.unit_id = %s AND %s IS NOT NULL) 
           OR (b.unit_number = %s AND %s IS NOT NULL AND b.association_id = %s AND %s IS NOT NULL)
        ORDER BY b.booking_date DESC, b.start_time DESC
    ''', (account["account_id"], unit_id, unit_id, unit_number, unit_number, assoc_id, assoc_id))
    
    return bookings

@api_router.post("/amenities/{amenity_id}/book")
async def book_amenity(amenity_id: str, payload: AmenityBookingCreate, account: dict = Depends(get_current_account)):
    amenity = await db_fetchone("SELECT * FROM amenities WHERE id = %s", (amenity_id,))
    if not amenity:
        raise HTTPException(status_code=404, detail="Amenity not found")
    if not amenity.get("status"):
        raise HTTPException(status_code=400, detail="Amenity is currently inactive")
        
    assoc_id = amenity.get("association_id")
    user_id = account.get("account_id") or account.get("id")
    
    # Collision check
    existing = await db_fetchone('''
        SELECT id FROM amenity_bookings 
        WHERE amenity_id = %s AND booking_date = %s 
        AND ((start_time < %s AND end_time > %s) OR (start_time >= %s AND start_time < %s))
        AND payment_status = 'Confirmed'
    ''', (amenity_id, payload.booking_date, payload.end_time, payload.start_time, payload.start_time, payload.end_time))
    
    if existing:
        raise HTTPException(status_code=409, detail="This time slot is already booked.")
    
    # Get user unit details
    user = await db_fetchone("SELECT unit_id, name FROM user_details WHERE user_id = %s", (account.get("user_id"),))
    unit_number = ""
    homeowner_name = account.get("name") or (user.get("name", "") if user else "") or account.get("email", "")
    if user and user.get("unit_id"):
        unit_info = await db_fetchone("SELECT unit_number FROM units WHERE id = %s", (user.get("unit_id"),))
        if unit_info:
            unit_number = unit_info.get("unit_number", "")
            
    # Calculate amount: base charges * duration_hours
    base_charge = float(amenity.get("charges") or 0)
    total_amount = base_charge * payload.duration_hours

    wallet = await db_fetchone("SELECT id, balance, security_pin FROM wallets WHERE account_id = %s", (user_id,))
    if not wallet:
        raise HTTPException(status_code=400, detail="Wallet not found for this account.")

    if payload.payment_method == "wallet":
        if not wallet.get("security_pin") or wallet["security_pin"] != payload.pin:
            raise HTTPException(status_code=400, detail="Invalid wallet security PIN.")
            
        if float(wallet.get("balance") or 0) < total_amount:
            raise HTTPException(status_code=400, detail="Insufficient wallet balance.")
        
        # Deduct balance
        await db_execute("UPDATE wallets SET balance = balance - %s WHERE id = %s", (total_amount, wallet["id"]))
        
    # Insert transaction for both Wallet and UPI
    await db_execute(
        "INSERT INTO wallet_transactions (wallet_id, type, amount, status, description) VALUES (%s, %s, %s, %s, %s)",
        (wallet["id"], "Debit" if payload.payment_method == "wallet" else "UPI", total_amount, "Completed", f"Amenity Booking: {amenity.get('name')}")
    )

    booking_id = await db_execute('''
        INSERT INTO amenity_bookings (amenity_id, user_id, association_id, unit_number, homeowner_name, amount, booking_date, payment_status, start_time, end_time, duration_hours)
        VALUES (%s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s)
    ''', (amenity_id, user_id, assoc_id, unit_number, homeowner_name, total_amount, payload.booking_date, "Confirmed", payload.start_time, payload.end_time, payload.duration_hours))
    
    return {"ok": True, "booking_id": booking_id, "message": "Amenity booked successfully"}


app.include_router(api_router)


# Swagger groups are assigned centrally so the existing API paths and handlers
# remain unchanged while /docs is organized by product feature.
OPENAPI_TAGS = [
    {"name": "Authentication", "description": "Login, verification, password, and session endpoints."},
    {"name": "Uploads", "description": "Image, video, and document uploads."},
    {"name": "Profile", "description": "Resident profile, family, vehicle, pet, and education data."},
    {"name": "Associations & Units", "description": "Association onboarding, settings, blocks, units, and members."},
    {"name": "Announcements", "description": "Community announcements, likes, and comments."},
    {"name": "Events & Polls", "description": "Events, RSVPs, polls, votes, and engagement."},
    {"name": "Service Requests", "description": "Service tickets, assignments, status, and messages."},
    {"name": "Board Tasks & Meetings", "description": "Board work, meetings, attendance, minutes, and board chat."},
    {"name": "Committees & Board", "description": "Board members, committees, and committee membership."},
    {"name": "Security & Visitors", "description": "Visitors, pre-approvals, deliveries, and staff attendance."},
    {"name": "Incidents", "description": "Security incident reporting and investigations."},
    {"name": "Marketplace", "description": "Marketplace listings, favourites, reports, and chat."},
    {"name": "Amenities", "description": "Amenity management and bookings."},
    {"name": "Wallet & Payments", "description": "Wallet balances, transfers, dues, UPI, and bank webhooks."},
    {"name": "Financials", "description": "Bank integrations, reconciliation, reports, chart of accounts, and budgets."},
    {"name": "Documents", "description": "Unit and administrative documents."},
    {"name": "Notifications", "description": "User notifications."},
    {"name": "Platform Administration", "description": "Platform users, roles, permissions, features, and subscriptions."},
    {"name": "General", "description": "General-purpose API endpoints."},
]


def swagger_feature_tag(path: str) -> str:
    """Return the Swagger feature group for an existing API path."""
    if path in {"/api/", "/api/login-code", "/api/request-code", "/api/user-details", "/api/create-account", "/api/update-details-request"} or path.startswith("/api/auth/"):
        return "Authentication"
    if path in {"/api/upload", "/api/upload-media-asset", "/api/upload-document-asset"}:
        return "Uploads"
    if path.startswith("/api/profile/"):
        return "Profile"
    if path.startswith("/api/announcements"):
        return "Announcements"
    if path.startswith("/api/events") or path.startswith("/api/polls") or "/announcements/" in path or "/events/" in path or "/polls/" in path:
        return "Events & Polls"
    if path.startswith("/api/service-requests"):
        return "Service Requests"
    if path.startswith("/api/board-tasks") or path.startswith("/api/meetings") or path.startswith("/api/board_chat"):
        return "Board Tasks & Meetings"
    if path.startswith("/api/security/") or path.startswith("/api/resident/visitor") or path.startswith("/api/resident/preapproved") or path.startswith("/api/resident/deliveries"):
        return "Security & Visitors"
    if path.startswith("/api/incidents"):
        return "Incidents"
    if path.startswith("/api/marketplace"):
        return "Marketplace"
    if path.startswith("/api/amenities") or "/amenities" in path:
        return "Amenities"
    if path.startswith("/api/wallet") or path.startswith("/api/webhooks/bank"):
        return "Wallet & Payments"
    if path.startswith("/api/accountant/") or any(item in path for item in ("bank-", "reconciliation", "financial-reports", "global-coa", "budgets")):
        return "Financials"
    if path.startswith("/api/unit-documents") or path.startswith("/api/admin/documents"):
        return "Documents"
    if path.startswith("/api/notifications"):
        return "Notifications"
    if path.startswith("/api/admin/board-members") or path.startswith("/api/admin/committees") or path.startswith("/api/admin/committee-members") or path.startswith("/api/user/committees"):
        return "Committees & Board"
    if path.startswith("/api/associations/") or path.startswith("/api/admin/associations") or path.startswith("/api/admin/blocks") or path.startswith("/api/admin/units") or path.startswith("/api/admin/homeowners"):
        return "Associations & Units"
    if path.startswith("/api/admin/") or path == "/api/features":
        return "Platform Administration"
    return "General"


def custom_openapi():
    if app.openapi_schema:
        return app.openapi_schema

    schema = get_openapi(
        title=app.title,
        version="1.0.0",
        routes=app.routes,
        tags=OPENAPI_TAGS,
    )
    for path, path_item in schema["paths"].items():
        tag = swagger_feature_tag(path)
        for method in ("get", "post", "put", "patch", "delete", "options", "head"):
            if method in path_item:
                path_item[method]["tags"] = [tag]

    app.openapi_schema = schema
    return app.openapi_schema


app.openapi = custom_openapi
